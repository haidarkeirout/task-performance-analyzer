"""ClickUp-only performance analysis built from task and time-in-status data.

This module intentionally has no Jira imports. It evaluates the ClickUp task
snapshot that was selected and collected, while keeping unsupported historical
metrics unavailable instead of inventing them.
"""
from __future__ import annotations

import json
from datetime import date, datetime, time
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from excel_safety import write_excel_cell

from clickup_excel_dashboard import add_clickup_executive_dashboard
from clickup_filters import (
    CANCELLED_STATUSES,
    TERMINAL_STATUSES,
    assignees,
    created_by,
    is_closed,
    location_values,
    milliseconds_to_hours,
    status_name,
    status_type,
    tags,
    task_type,
    timestamp,
    custom_fields,
)


NOT_STARTED_STATUSES = {
    "backlog",
    "to do",
    "todo",
    "not started",
    "open",
    "planning",
    "ready",
}

# The shared KPI contract treats only active execution and review as WIP.
# On Hold and At Risk remain visible as separate workflow states.
WIP_STATUSES = {"in progress", "review"}


def _status_minutes(payload: Any) -> dict[str, float]:
    """Flatten native ClickUp Total time in Status data by status."""
    values: dict[str, float] = {}
    if not isinstance(payload, dict):
        return values
    entries = list(payload.get("status_history") or [])
    current = payload.get("current_status")
    if isinstance(current, dict):
        entries.append(current)
    for entry in entries:
        if not isinstance(entry, dict) or not entry.get("status"):
            continue
        total_time = entry.get("total_time") or {}
        try:
            minutes = float(total_time.get("by_minute"))
        except (TypeError, ValueError, AttributeError):
            continue
        status = str(entry["status"])
        values[status] = max(minutes, values.get(status, 0.0))
    return values


def _duration_hours(start, end):
    if pd.isna(start) or pd.isna(end):
        return None
    return (end - start).total_seconds() / 3600.0


def _timing_values(created, start, end, end_label: str):
    """Return trustworthy task-timing values without inventing negative time.

    ClickUp start dates are frequently date-only values. A same-day start
    timestamp can therefore precede the creation timestamp by a few hours.
    In that case the effective start is the creation time. A start date after
    the completion/cutoff, however, is a chronology conflict and the derived
    timing metrics must remain unavailable.
    """
    if pd.isna(start):
        return None, None, "No start date"
    if pd.notna(end) and start > end:
        return None, None, f"Start date after {end_label}"
    if pd.notna(created) and start.normalize() < created.normalize():
        return None, None, "Start date before creation date"

    effective_start = max(start, created) if pd.notna(created) and start < created else start
    return (
        _duration_hours(effective_start, end),
        _duration_hours(created, effective_start),
        "Available",
    )


def _first_location(task: dict) -> str:
    values = location_values(task)
    return values[-1] if values else "Unavailable"


def _project_space(task: dict) -> str:
    explicit = task.get("employee_space_name")
    if explicit not in (None, ""):
        return str(explicit)
    return _first_location(task)


def _company(task: dict) -> str:
    for key, value in task.items():
        normalized = "".join(character for character in str(key).casefold() if character.isalnum())
        if any(token in normalized for token in ("company", "client", "customer", "organization")) and value not in (None, ""):
            return str(value)
    fields = custom_fields(task)
    for name, value in fields.items():
        if any(token in name.casefold() for token in ("company", "client", "customer", "organization")):
            return value
    return "Company not specified"


def _due_category(completed: bool, open_task: bool, due, variance):
    if pd.isna(due) or variance is None or pd.isna(variance):
        return "No due date"
    if completed:
        if variance > 0:
            return "Completed late"
        if variance < 0:
            return "Completed early"
        return "Completed on time"
    if open_task:
        return "Open overdue" if variance > 0 else "Open not due"
    return "Not applicable"


def _safe_rate(values: pd.Series):
    usable = values.dropna()
    return None if usable.empty else float(usable.mean() * 100.0)


def _mean(frame: pd.DataFrame, column: str, mask=None):
    if frame.empty or column not in frame:
        return None
    values = pd.to_numeric(frame.loc[mask, column] if mask is not None else frame[column], errors="coerce").dropna()
    return None if values.empty else float(values.mean())


def _utc_datetime_series(values):
    """Normalize mixed ClickUp datetime values before pandas comparisons."""
    return pd.to_datetime(values, utc=True, errors="coerce")


def _weekly_flow(task_frame: pd.DataFrame, cutoff=None) -> pd.DataFrame:
    """Build Monday-starting created/completed flow through the evaluation cutoff.

    Empty calendar weeks are retained with zero counts so a dashboard request for
    4, 12, 24, 36, or 48 weeks always represents that actual time window.
    """
    columns = ["Week Starting", "Tasks Created", "Tasks Completed", "Net Flow", "Cumulative Net Flow"]
    if task_frame.empty:
        return pd.DataFrame(columns=columns)

    created = task_frame.loc[task_frame["Created"].notna(), ["Created"]].copy()
    completed = task_frame.loc[task_frame["Completed"].notna(), ["Completed"]].copy()
    created["Created"] = _utc_datetime_series(created["Created"])
    completed["Completed"] = _utc_datetime_series(completed["Completed"])
    cutoff = pd.to_datetime(cutoff, utc=True, errors="coerce")

    if pd.notna(cutoff):
        created = created[created["Created"].le(cutoff)]
        completed = completed[completed["Completed"].le(cutoff)]

    created["Week Starting"] = created["Created"].dt.normalize() - pd.to_timedelta(created["Created"].dt.weekday, unit="D")
    completed["Week Starting"] = completed["Completed"].dt.normalize() - pd.to_timedelta(completed["Completed"].dt.weekday, unit="D")
    created_counts = created.groupby("Week Starting").size().rename("Tasks Created")
    completed_counts = completed.groupby("Week Starting").size().rename("Tasks Completed")
    grouped = pd.concat([created_counts, completed_counts], axis=1).fillna(0)
    if grouped.empty:
        return pd.DataFrame(columns=columns)

    first_week = grouped.index.min()
    last_week = grouped.index.max()
    if pd.notna(cutoff):
        cutoff_week = cutoff.normalize() - pd.to_timedelta(cutoff.weekday(), unit="D")
        if cutoff_week > last_week:
            last_week = cutoff_week

    weeks = pd.date_range(first_week, last_week, freq="7D")
    output = grouped.reindex(weeks).fillna(0)
    output.index.name = "Week Starting"
    output = output.reset_index()
    output[["Tasks Created", "Tasks Completed"]] = output[["Tasks Created", "Tasks Completed"]].astype(int)
    output["Net Flow"] = output["Tasks Created"] - output["Tasks Completed"]
    output["Cumulative Net Flow"] = output["Net Flow"].cumsum()
    return output[columns].reset_index(drop=True)


def _due_status_summary(frame: pd.DataFrame, cutoff) -> pd.DataFrame:
    columns = ["Due Status", "Tasks", "Share of Open Tasks (%)"]
    if frame.empty:
        return pd.DataFrame(columns=columns)
    open_tasks = frame[frame["Open?"]].copy()
    if open_tasks.empty:
        return pd.DataFrame(columns=columns)

    cutoff = pd.to_datetime(cutoff, utc=True, errors="coerce")

    def bucket(row):
        due = pd.to_datetime(row["Due Date"], utc=True, errors="coerce")
        if pd.isna(due):
            return "No due date"
        if pd.notna(cutoff) and due < cutoff:
            return "Overdue"
        if pd.notna(cutoff) and due <= cutoff + pd.Timedelta(days=7):
            return "Due in next 7 days"
        return "Due later"

    output = open_tasks.assign(**{"Due Status": open_tasks.apply(bucket, axis=1)}).groupby("Due Status").size().rename("Tasks").reset_index()
    output["Share of Open Tasks (%)"] = output["Tasks"] / len(open_tasks) * 100.0
    return output.sort_values("Tasks", ascending=False).reset_index(drop=True)


def _assignee_summary(frame: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "Assignee", "Total Tasks", "Completed", "Open", "WIP", "Cancelled",
        "Completion Rate (%)", "On-Time Completion Rate (%)", "Completed Late",
        "Open Overdue", "Average Due Variance (days)",
    ]
    if frame.empty:
        return pd.DataFrame(columns=columns)
    rows = []
    for assignee, group in frame.groupby("Assignee", dropna=False):
        known = group[group["Status Known?"].eq(True)]
        completed = known[known["Completed?"]]
        rows.append({
            "Assignee": assignee or "Unassigned",
            "Total Tasks": int(len(group)),
            "Completed": int(group["Completed?"].sum()),
            "Open": int(group["Open?"].sum()),
            "WIP": int(group["WIP?"].sum()),
            "Cancelled": int(group["Cancelled?"].sum()),
            "Completion Rate (%)": float(completed.shape[0] / len(known) * 100.0) if len(known) else None,
            "On-Time Completion Rate (%)": _safe_rate(completed["On Time?"]) if not completed.empty else None,
            "Completed Late": int((completed["Due Variance (days)"] > 0).sum()) if not completed.empty else 0,
            "Open Overdue": int(((group["Open?"]) & (group["Due Variance (days)"] > 0)).sum()),
            "Average Due Variance (days)": _mean(completed, "Due Variance (days)"),
        })
    return pd.DataFrame(rows, columns=columns).sort_values(["Total Tasks", "Assignee"], ascending=[False, True]).reset_index(drop=True)


def _status_summary(status_frame: pd.DataFrame) -> pd.DataFrame:
    columns = ["Status", "Tasks", "Average Hours", "Median Hours", "Total Hours"]
    if status_frame.empty:
        return pd.DataFrame(columns=columns)
    return (status_frame.groupby("Status", dropna=False)
            .agg(Tasks=("Task ID", "nunique"), Average_Hours=("Hours", "mean"),
                 Median_Hours=("Hours", "median"), Total_Hours=("Hours", "sum"))
            .reset_index()
            .rename(columns={"Average_Hours": "Average Hours", "Median_Hours": "Median Hours", "Total_Hours": "Total Hours"})
            .sort_values("Total Hours", ascending=False)
            .reset_index(drop=True))


def _period_boundary(value: Any, source_zone: ZoneInfo, *, end_of_day: bool = False):
    """Convert a selected local date into an inclusive UTC analysis boundary."""
    if value in (None, ""):
        return pd.NaT
    try:
        parsed = pd.Timestamp(value)
    except (TypeError, ValueError):
        return pd.NaT
    if pd.isna(parsed):
        return pd.NaT
    if parsed.tzinfo is None:
        if end_of_day and parsed.time() == time.min:
            parsed = parsed.normalize() + pd.Timedelta(days=1) - pd.Timedelta(microseconds=1)
        parsed = parsed.tz_localize(source_zone)
    else:
        parsed = parsed.tz_convert("UTC")
    return parsed.tz_convert("UTC")


def _period_label(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value)[:10]


def analyze_clickup(prepared_data, *, period_start=None, period_end=None):
    """Return a ClickUp-native analysis payload for the selected employee period."""
    payload = json.loads(prepared_data.history_json.decode("utf-8"))
    raw_tasks = payload.get("tasks") if isinstance(payload, dict) else []
    time_payloads = payload.get("time_in_status", {}) if isinstance(payload, dict) else {}
    source_timezone = getattr(prepared_data, "source_timezone", "Asia/Damascus")
    source_zone = ZoneInfo(source_timezone)
    collected_at = timestamp(getattr(prepared_data, "collected_at", None))
    collection_cutoff = pd.to_datetime(timestamp(prepared_data.cutoff), utc=True, errors="coerce")
    cutoff = (
        _period_boundary(period_end, source_zone, end_of_day=True)
        if period_end is not None
        else collection_cutoff
    )
    selected_start = (
        _period_boundary(period_start, source_zone)
        if period_start is not None
        else pd.NaT
    )
    historical_snapshot = (
        pd.notna(cutoff)
        and pd.notna(collected_at)
        and cutoff.tz_convert(source_zone).date() < collected_at.tz_convert(source_zone).date()
    )
    rows: list[dict[str, Any]] = []
    status_rows: list[dict[str, Any]] = []

    for task in raw_tasks if isinstance(raw_tasks, list) else []:
        task_id = str(task.get("id", ""))
        status = status_name(task)
        status_key = status.casefold()
        state = status_type(task)
        created = timestamp(task.get("date_created"))
        updated = timestamp(task.get("date_updated"))
        start = timestamp(task.get("start_date"))
        due = timestamp(task.get("due_date"))
        completed = timestamp(task.get("date_closed") or task.get("date_done"))
        if pd.notna(cutoff) and pd.notna(created) and created > cutoff:
            continue
        if pd.notna(selected_start):
            open_at_end = pd.isna(completed) or completed > cutoff
            created_in_period = pd.notna(created) and selected_start <= created <= cutoff
            updated_in_period = pd.notna(updated) and selected_start <= updated <= cutoff
            completed_in_period = pd.notna(completed) and selected_start <= completed <= cutoff
            if not (created_in_period or updated_in_period or completed_in_period or open_at_end):
                continue
        source_cancelled = status_key in CANCELLED_STATUSES
        source_completed = (
            status_key in TERMINAL_STATUSES or is_closed(task) or pd.notna(completed)
        ) and not source_cancelled
        status_known = not historical_snapshot
        status_basis = "Collection snapshot"
        if historical_snapshot:
            if source_completed and pd.notna(completed) and completed <= cutoff:
                status_known = True
                status_basis = "Completion date on or before cutoff"
            elif source_cancelled and pd.notna(completed) and completed <= cutoff:
                status_known = True
                status_basis = "Cancellation close date on or before cutoff"
            else:
                status_known = False
                status_basis = "Historical status unavailable from ClickUp snapshot"

        cancelled = bool(status_known and source_cancelled)
        completed_flag = bool(status_known and source_completed and pd.notna(completed) and completed <= cutoff)
        if status_known and source_completed and pd.isna(completed):
            # The current-day snapshot can confirm the terminal state even when
            # ClickUp omitted its close timestamp; historical snapshots cannot.
            completed_flag = not historical_snapshot
        open_flag = bool(status_known and not completed_flag and not cancelled)
        wip_flag = open_flag and status_key in WIP_STATUSES
        end = completed if pd.notna(completed) else cutoff
        lead_time_hours = _duration_hours(created, end)
        timing_end = completed if completed_flag else cutoff
        timing_label = "completion date" if completed_flag else "evaluation cutoff"
        execution_hours, time_to_start_hours, timing_data_status = _timing_values(
            created, start, timing_end, timing_label
        )
        due_variance_days = None
        due_basis = "Unavailable"
        if pd.notna(due) and pd.notna(completed) and completed_flag:
            due_variance_days = (completed - due).total_seconds() / 86400.0
            due_basis = "Completed date vs due date"
        elif pd.notna(due) and open_flag and pd.notna(cutoff):
            due_variance_days = (cutoff - due).total_seconds() / 86400.0
            due_basis = "Evaluation cutoff vs due date"
        on_time = bool(completed <= due) if completed_flag and pd.notna(completed) and pd.notna(due) else None
        overdue_days = due_variance_days if open_flag and due_variance_days is not None and due_variance_days > 0 else None

        time_payload = time_payloads.get(task_id) if isinstance(time_payloads, dict) else None
        minutes_by_status = _status_minutes(time_payload)
        current_minutes = None
        current_since = None
        if isinstance(time_payload, dict) and isinstance(time_payload.get("current_status"), dict):
            current = time_payload["current_status"]
            current_minutes = (current.get("total_time") or {}).get("by_minute")
            current_since = (current.get("total_time") or {}).get("since")
            try:
                current_minutes = float(current_minutes)
            except (TypeError, ValueError):
                current_minutes = None

        row = {
            "Task ID": task_id,
            "Task Name": task.get("name"),
            "Assignee": ", ".join(assignees(task)),
            "Created By": created_by(task),
            "Priority": (task.get("priority") or {}).get("priority") if isinstance(task.get("priority"), dict) else task.get("priority"),
            "Task Type": task_type(task),
            "Company": _company(task),
            "Project / Space": _project_space(task),
            "Tags": ", ".join(tags(task)) or "No tags",
            "Location/List": task.get("_department_list_name") or _first_location(task),
            "Space": task.get("_department_space_name") or "",
            "List": task.get("_department_list_name") or "",
            "Parent Task ID": str(task.get("parent") or ""),
            "Is Subtask": bool(task.get("parent")),
            "Current Status": status,
            "Status at Cutoff": status if status_known else "Unknown",
            "Status Known?": status_known,
            "Historical Status Basis": status_basis,
            "Status State": state or "Unavailable",
            "Created": created,
            "Updated": updated,
            "Start Date": start,
            "Due Date": due,
            "Completed": completed,
            "Completed?": completed_flag,
            "Cancelled?": cancelled,
            "Open?": open_flag,
            "WIP?": wip_flag,
            "Elapsed Hours": lead_time_hours,
            "Lead Time Hours": lead_time_hours,
            "Execution Hours": execution_hours,
            "Time to Start Hours": time_to_start_hours,
            "Timing Data Status": timing_data_status,
            "On Time?": on_time,
            "Due Variance (days)": due_variance_days,
            "Due Variance Basis": due_basis,
            "Due Variance Category": _due_category(completed_flag, open_flag, due, due_variance_days),
            "Overdue Days": overdue_days,
            "Time Estimate Hours": milliseconds_to_hours(task.get("time_estimate")),
            "Time Tracked Hours": milliseconds_to_hours(task.get("time_spent")),
            "Current Status Time (min)": current_minutes,
            "Current Status Since": timestamp(current_since),
            "Total Time in Status (min)": sum(minutes_by_status.values()) if minutes_by_status else None,
            "Time in Status": json.dumps(minutes_by_status, ensure_ascii=False) if minutes_by_status else None,
        }
        rows.append(row)
        for status_label, minutes in minutes_by_status.items():
            status_rows.append({"Status": status_label, "Task ID": task_id, "Minutes": minutes, "Hours": minutes / 60.0})

    columns = [
        "Task ID", "Task Name", "Assignee", "Created By", "Priority", "Task Type", "Company", "Project / Space", "Tags", "Location/List",
        "Space", "List", "Parent Task ID", "Is Subtask",
        "Current Status", "Status at Cutoff", "Status Known?", "Historical Status Basis",
        "Status State", "Created", "Updated", "Start Date", "Due Date", "Completed",
        "Completed?", "Cancelled?", "Open?", "WIP?", "Elapsed Hours", "Lead Time Hours", "Execution Hours",
        "Time to Start Hours", "Timing Data Status", "On Time?", "Due Variance (days)", "Due Variance Basis", "Due Variance Category",
        "Overdue Days", "Time Estimate Hours", "Time Tracked Hours", "Current Status Time (min)",
        "Current Status Since", "Total Time in Status (min)", "Time in Status",
    ]
    task_frame = pd.DataFrame(rows, columns=columns)
    for date_column in ("Created", "Updated", "Start Date", "Due Date", "Completed", "Current Status Since"):
        task_frame[date_column] = _utc_datetime_series(task_frame[date_column])
    status_frame = pd.DataFrame(status_rows, columns=["Status", "Task ID", "Minutes", "Hours"])
    total = len(task_frame)
    known_status_tasks = task_frame[task_frame["Status Known?"].eq(True)] if total else task_frame
    completed = task_frame[task_frame["Completed?"]] if total else task_frame
    open_tasks = task_frame[task_frame["Open?"]] if total else task_frame
    due_known_completed = completed[completed["Due Variance (days)"].notna()] if total else task_frame
    status_summary = _status_summary(status_frame)
    status_counts = (task_frame["Status at Cutoff"].fillna("Unavailable").value_counts().rename_axis("Status").reset_index(name="Tasks")
                     if total else pd.DataFrame(columns=["Status", "Tasks"]))
    weekly_flow = _weekly_flow(task_frame, cutoff)
    due_status_summary = _due_status_summary(task_frame, cutoff)
    due_variance_summary = (task_frame[task_frame["Due Variance (days)"].notna()]
                            .groupby("Due Variance Category", dropna=False)
                            .agg(Tasks=("Task ID", "count"), Average_Variance_Days=("Due Variance (days)", "mean"),
                                 Median_Variance_Days=("Due Variance (days)", "median"))
                            .reset_index()
                            .rename(columns={"Average_Variance_Days": "Average Variance (days)", "Median_Variance_Days": "Median Variance (days)"})
                            .sort_values("Tasks", ascending=False).reset_index(drop=True)
                            if total else pd.DataFrame(columns=["Due Variance Category", "Tasks", "Average Variance (days)", "Median Variance (days)"]))
    assignee_summary = _assignee_summary(task_frame)
    late_completed = completed[completed["Due Variance (days)"] > 0].copy() if total else task_frame
    overdue_open = open_tasks[open_tasks["Due Variance (days)"] > 0].copy() if total else task_frame
    missing_due = task_frame[task_frame["Due Date"].isna()].copy() if total else task_frame
    status_duration_available = int(task_frame["Total Time in Status (min)"].notna().sum()) if total else 0
    future_start_tasks = (
        int(task_frame["Timing Data Status"].eq("Start date after evaluation cutoff").sum())
        if total else 0
    )
    timing_conflicts = (
        int(task_frame["Timing Data Status"].isin({
            "Start date after completion date", "Start date before creation date",
        }).sum())
        if total else 0
    )

    overall_values = [
        ("Total tasks", total),
        ("Completed tasks", int(task_frame["Completed?"].sum()) if total else 0),
        ("Open tasks", int(task_frame["Open?"].sum()) if total else 0),
        ("WIP tasks", int(task_frame["WIP?"].sum()) if total else 0),
        ("Cancelled tasks", int(task_frame["Cancelled?"].sum()) if total else 0),
        ("Known status tasks", int(len(known_status_tasks))),
        ("Unknown status tasks", int(total - len(known_status_tasks))),
        ("Completion rate (%)", (
            float(known_status_tasks["Completed?"].mean() * 100.0)
            if len(known_status_tasks) else None
        )),
        ("On-time completion rate (%)", _safe_rate(completed["On Time?"]) if total else None),
        ("Open overdue tasks", int(len(overdue_open))),
        ("Completed late tasks", int(len(late_completed))),
        ("Average due variance (days)", _mean(due_known_completed, "Due Variance (days)")),
        ("Average execution hours", _mean(completed, "Execution Hours")),
        ("Average lead time hours", _mean(completed, "Lead Time Hours")),
        ("Average time to start hours", _mean(task_frame, "Time to Start Hours")),
        ("Tasks with Total time in Status", status_duration_available),
    ]
    overall = pd.DataFrame(overall_values, columns=["Metric", "Value"])
    time_status_state = (
        "Available" if status_duration_available
        else ("Unavailable for this scope/account." if getattr(prepared_data, "clickup_time_status_error", "") else "Not requested")
    )
    source_spaces = getattr(prepared_data, "space_names", None) or [prepared_data.space_name]
    analysis_context = pd.DataFrame([
        ("Process Name", prepared_data.space_name),
        ("Department", getattr(prepared_data, "department_name", "")),
        ("Source Spaces", ", ".join(map(str, source_spaces))),
        ("Evaluation Scope", getattr(prepared_data, "filter_summary", "All tasks in the selected ClickUp Space")),
        ("Dataset Type", "ClickUp API collection"),
        ("Evaluation Period Start", _period_label(period_start)),
        ("Evaluation Period End", _period_label(period_end)),
        ("Evaluation Cutoff", cutoff.isoformat() if pd.notna(cutoff) else prepared_data.cutoff),
        ("Source Timezone", prepared_data.source_timezone),
        ("Task Count", total),
        ("Duplicate task records removed", getattr(prepared_data, "duplicate_count", 0)),
        ("Activity History", "Not collected; browser collector is disabled."),
        ("Total time in Status", time_status_state),
    ], columns=["Field", "Value"])
    quality = pd.DataFrame([
        ("Task rows collected", total, "OK" if total else "Unavailable"),
        ("Duplicate task records removed", getattr(prepared_data, "duplicate_count", 0),
         "OK" if getattr(prepared_data, "duplicate_count", 0) == 0 else "Deduplicated"),
        ("Tasks with status-duration data", status_duration_available, "OK" if status_duration_available else "Unavailable"),
        ("Tasks missing created date", int(task_frame["Created"].isna().sum()) if total else 0,
         "Review" if total and task_frame["Created"].isna().any() else "OK"),
        ("Tasks missing due date", int(task_frame["Due Date"].isna().sum()) if total else 0,
         "Review" if total and task_frame["Due Date"].isna().any() else "OK"),
        ("Completed tasks missing completion date", int((task_frame["Completed?"] & task_frame["Completed"].isna()).sum()) if total else 0,
         "Review" if total and (task_frame["Completed?"] & task_frame["Completed"].isna()).any() else "OK"),
        ("Tasks scheduled after the evaluation cutoff", future_start_tasks,
         "Info" if future_start_tasks else "OK"),
        ("Tasks with timing chronology conflicts", timing_conflicts,
         "Review" if timing_conflicts else "OK"),
        ("Tasks with historical status unavailable",
         int((~task_frame["Status Known?"]).sum()) if total else 0,
         "Unavailable" if total and (~task_frame["Status Known?"]).any() else "OK"),
    ], columns=["Check", "Value", "Status"])
    metric_definitions = pd.DataFrame([
        ("Due Variance (days)", "Completed task: Completed date minus Due date. Open task: Evaluation cutoff minus Due date. Positive is late/overdue; negative is early/time remaining."),
        ("On-time completion rate", "Share of completed tasks with both a due date and a completion date where completion was on or before the due date."),
        ("Execution and time to start", "Calculated only when the recorded start date is on or before the completion date or evaluation cutoff. Date-only starts on the creation day use creation time as the effective start. Chronology conflicts remain unavailable."),
        ("WIP tasks", "Open tasks whose status is In Progress or Review. On Hold and At Risk are reported separately."),
        ("Status-duration data", "Read only from ClickUp Total time in Status when the ClickApp/API exposes it. Missing values stay unavailable."),
        ("Historical ClickUp status", "A current ClickUp snapshot is never presented as a past status. If no chronological evidence establishes the cutoff state, status-dependent metrics remain unavailable."),
        ("Individual assignment", "Uses the assignee snapshot returned by ClickUp. It does not establish individual contribution or ownership at completion."),
    ], columns=["Metric", "Definition"])
    findings = pd.DataFrame([
        ("Completed late tasks", int(len(late_completed)), "Review deadlines, task estimates, and dependencies for these completed tasks."),
        ("Open overdue tasks", int(len(overdue_open)), "Prioritize current blockers and agree a recovery plan for open overdue work."),
        ("Tasks without due dates", int(len(missing_due)), "Add due dates where delivery timeliness is expected to be evaluated."),
        ("Tasks with timing chronology conflicts", timing_conflicts, "Review the task start and completion dates. Conflicting timing values are excluded from duration averages."),
        ("Tasks without status-duration data", total - status_duration_available, "Enable or verify Total time in Status only when the ClickUp plan and permissions support it."),
    ], columns=["Finding", "Tasks", "Recommended Follow-up"])
    return {
        "tasks": task_frame,
        "status_detail": status_frame,
        "status_summary": status_summary,
        "status_counts": status_counts,
        "weekly_flow": weekly_flow,
        "due_status_summary": due_status_summary,
        "due_variance_summary": due_variance_summary,
        "assignee_summary": assignee_summary,
        "late_completed_tasks": late_completed,
        "overdue_tasks": overdue_open,
        "missing_due_tasks": missing_due,
        "overall": overall,
        "analysis_context": analysis_context,
        "quality": quality,
        "metric_definitions": metric_definitions,
        "findings": findings,
        "cutoff": cutoff.isoformat() if pd.notna(cutoff) else prepared_data.cutoff,
        "period_start": _period_label(period_start),
        "period_end": _period_label(period_end),
        "source_timezone": prepared_data.source_timezone,
        "space_name": prepared_data.space_name,
        "space_names": source_spaces,
        "department_name": getattr(prepared_data, "department_name", ""),
        "filter_summary": getattr(prepared_data, "filter_summary", "All tasks in the selected ClickUp Space"),
    }


def recalculate_clickup_analysis(result: dict, task_frame: pd.DataFrame) -> dict:
    """Rebuild ClickUp summaries from an already analysed filtered snapshot.

    Employee post-analysis filters must not trigger another API call.  This
    function deliberately reuses the task-level values already calculated by
    :func:`analyze_clickup` and only rebuilds aggregates and dashboard tables.
    """
    recalculated = dict(result)
    tasks = task_frame.copy().reset_index(drop=True)
    task_ids = set(tasks["Task ID"].astype(str)) if "Task ID" in tasks else set()
    original_status_detail = result.get("status_detail", pd.DataFrame())
    if original_status_detail.empty or "Task ID" not in original_status_detail:
        status_frame = original_status_detail.copy()
    else:
        status_frame = original_status_detail[
            original_status_detail["Task ID"].astype(str).isin(task_ids)
        ].copy()

    cutoff = timestamp(result.get("cutoff"))
    total = len(tasks)
    known_status_tasks = (
        tasks[tasks["Status Known?"].eq(True)]
        if total and "Status Known?" in tasks
        else tasks
    )
    completed = tasks[tasks["Completed?"]] if total else tasks
    open_tasks = tasks[tasks["Open?"]] if total else tasks
    due_known_completed = completed[completed["Due Variance (days)"].notna()] if total else tasks
    status_summary = _status_summary(status_frame)
    status_counts = (
        tasks["Status at Cutoff"].fillna("Unavailable").value_counts()
        .rename_axis("Status").reset_index(name="Tasks")
        if total else pd.DataFrame(columns=["Status", "Tasks"])
    )
    weekly_flow = _weekly_flow(tasks, cutoff)
    due_status_summary = _due_status_summary(tasks, cutoff)
    due_variance_summary = (
        tasks[tasks["Due Variance (days)"].notna()]
        .groupby("Due Variance Category", dropna=False)
        .agg(Tasks=("Task ID", "count"), Average_Variance_Days=("Due Variance (days)", "mean"),
             Median_Variance_Days=("Due Variance (days)", "median"))
        .reset_index()
        .rename(columns={"Average_Variance_Days": "Average Variance (days)",
                         "Median_Variance_Days": "Median Variance (days)"})
        .sort_values("Tasks", ascending=False).reset_index(drop=True)
        if total else pd.DataFrame(columns=["Due Variance Category", "Tasks", "Average Variance (days)", "Median Variance (days)"])
    )
    assignee_summary = _assignee_summary(tasks)
    late_completed = completed[completed["Due Variance (days)"] > 0].copy() if total else tasks
    overdue_open = open_tasks[open_tasks["Due Variance (days)"] > 0].copy() if total else tasks
    missing_due = tasks[tasks["Due Date"].isna()].copy() if total else tasks
    status_duration_available = int(tasks["Total Time in Status (min)"].notna().sum()) if total else 0
    future_start_tasks = int(tasks["Timing Data Status"].eq("Start date after evaluation cutoff").sum()) if total else 0
    timing_conflicts = int(tasks["Timing Data Status"].isin({
        "Start date after completion date", "Start date before creation date",
    }).sum()) if total else 0

    overall = pd.DataFrame([
        ("Total tasks", total),
        ("Completed tasks", int(tasks["Completed?"].sum()) if total else 0),
        ("Open tasks", int(tasks["Open?"].sum()) if total else 0),
        ("WIP tasks", int(tasks["WIP?"].sum()) if total else 0),
        ("Cancelled tasks", int(tasks["Cancelled?"].sum()) if total else 0),
        ("Known status tasks", int(len(known_status_tasks))),
        ("Unknown status tasks", int(total - len(known_status_tasks))),
        ("Completion rate (%)", float(known_status_tasks["Completed?"].mean() * 100.0) if len(known_status_tasks) else None),
        ("On-time completion rate (%)", _safe_rate(completed["On Time?"]) if total else None),
        ("Open overdue tasks", int(len(overdue_open))),
        ("Completed late tasks", int(len(late_completed))),
        ("Average due variance (days)", _mean(due_known_completed, "Due Variance (days)")),
        ("Average execution hours", _mean(completed, "Execution Hours")),
        ("Average lead time hours", _mean(completed, "Lead Time Hours")),
        ("Average time to start hours", _mean(tasks, "Time to Start Hours")),
        ("Tasks with Total time in Status", status_duration_available),
    ], columns=["Metric", "Value"])

    context = result.get("analysis_context", pd.DataFrame()).copy()
    if not context.empty and "Field" in context and "Value" in context:
        context.loc[context["Field"].eq("Task Count"), "Value"] = total
    quality = pd.DataFrame([
        ("Task rows collected", total, "OK" if total else "Unavailable"),
        ("Tasks with status-duration data", status_duration_available, "OK" if status_duration_available else "Unavailable"),
        ("Tasks missing created date", int(tasks["Created"].isna().sum()) if total else 0,
         "Review" if total and tasks["Created"].isna().any() else "OK"),
        ("Tasks missing due date", int(tasks["Due Date"].isna().sum()) if total else 0,
         "Review" if total and tasks["Due Date"].isna().any() else "OK"),
        ("Completed tasks missing completion date", int((tasks["Completed?"] & tasks["Completed"].isna()).sum()) if total else 0,
         "Review" if total and (tasks["Completed?"] & tasks["Completed"].isna()).any() else "OK"),
        ("Tasks scheduled after the evaluation cutoff", future_start_tasks, "Info" if future_start_tasks else "OK"),
        ("Tasks with timing chronology conflicts", timing_conflicts, "Review" if timing_conflicts else "OK"),
        ("Tasks with historical status unavailable", int((~tasks["Status Known?"]).sum()) if total else 0,
         "Unavailable" if total and (~tasks["Status Known?"]).any() else "OK"),
    ], columns=["Check", "Value", "Status"])
    findings = pd.DataFrame([
        ("Completed late tasks", int(len(late_completed)), "Review deadlines, task estimates, and dependencies for these completed tasks."),
        ("Open overdue tasks", int(len(overdue_open)), "Prioritize current blockers and agree a recovery plan for open overdue work."),
        ("Tasks without due dates", int(len(missing_due)), "Add due dates where delivery timeliness is expected to be evaluated."),
        ("Tasks with timing chronology conflicts", timing_conflicts, "Review the task start and completion dates. Conflicting timing values are excluded from duration averages."),
        ("Tasks without status-duration data", total - status_duration_available, "Enable or verify Total time in Status only when the ClickUp plan and permissions support it."),
    ], columns=["Finding", "Tasks", "Recommended Follow-up"])

    recalculated.update({
        "tasks": tasks,
        "status_detail": status_frame,
        "status_summary": status_summary,
        "status_counts": status_counts,
        "weekly_flow": weekly_flow,
        "due_status_summary": due_status_summary,
        "due_variance_summary": due_variance_summary,
        "assignee_summary": assignee_summary,
        "late_completed_tasks": late_completed,
        "overdue_tasks": overdue_open,
        "missing_due_tasks": missing_due,
        "overall": overall,
        "analysis_context": context,
        "quality": quality,
        "findings": findings,
        "filter_summary": f"Filtered employee snapshot: {total} task(s)",
    })
    return recalculated


def _write_sheet(workbook, name: str, frame: pd.DataFrame):
    sheet = workbook.create_sheet(name)
    for column, value in enumerate(frame.columns, 1):
        cell = write_excel_cell(sheet, 1, column, value)
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(color="FFFFFF", bold=True)
    for row_number, row in enumerate(frame.itertuples(index=False, name=None), 2):
        for column, value in enumerate(row, 1):
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            if isinstance(value, pd.Timestamp):
                value = value.isoformat()
            write_excel_cell(sheet, row_number, column, value)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 48)
        sheet.column_dimensions[column_cells[0].column_letter].width = width


def analysis_excel(result) -> bytes:
    """Create an English ClickUp analysis workbook with the ClickUp executive dashboard and analysis tables."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [
        ("Executive_Summary", result["overall"]),
        ("Task_Evaluation", result["tasks"]),
        ("Assignee_Summary", result["assignee_summary"]),
        ("Due_Variance", result["due_variance_summary"]),
        ("Open_Due_Status", result["due_status_summary"]),
        ("Weekly_Flow", result["weekly_flow"]),
        ("Status_Summary", result["status_summary"]),
        ("Status_Detail", result["status_detail"]),
        ("Findings", result["findings"]),
        ("Data_Quality", result["quality"]),
        ("Metric_Definitions", result["metric_definitions"]),
        ("Analysis_Context", result["analysis_context"]),
    ]
    for name, frame in sheets:
        _write_sheet(workbook, name, frame)
    add_clickup_executive_dashboard(workbook, result)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
