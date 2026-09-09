"""ClickUp-only Excel Executive Dashboard builder.

This module consumes only the ClickUp analysis result dictionary.
"""
from __future__ import annotations

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DARK = "17324D"
MID = "2F5D7E"
LIGHT = "EAF1F6"
WHITE = "FFFFFF"
TEXT = "1F2937"
BORDER = Side(style="thin", color="D7E1E8")
PANEL_RANGES = {
    "A13": "A13:H26",
    "I13": "I13:P26",
    "A29": "A29:H42",
    "I29": "I29:P42",
    "A45": "A45:H58",
    "I45": "I45:P58",
}


def _metric(result, name, default=None):
    frame = result.get("overall")
    if frame is None or frame.empty:
        return default
    matches = frame.loc[frame["Metric"].eq(name), "Value"]
    if matches.empty or pd.isna(matches.iloc[0]):
        return default
    return matches.iloc[0]


def _fmt_pct(value):
    return "Unavailable" if value is None or pd.isna(value) else f"{float(value):.1f}%"


def _fmt_avg(value, suffix):
    return "Unavailable" if value is None or pd.isna(value) else f"{float(value):.1f} {suffix}"


def _card(sheet, start_col, label, value, row):
    end_col = start_col + 1
    sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    sheet.merge_cells(start_row=row + 1, start_column=start_col, end_row=row + 2, end_column=end_col)
    sheet.cell(row, start_col, label).fill = PatternFill("solid", fgColor=DARK)
    sheet.cell(row, start_col).font = Font(color=WHITE, bold=True, size=9)
    sheet.cell(row + 1, start_col, value).fill = PatternFill("solid", fgColor=LIGHT)
    sheet.cell(row + 1, start_col).font = Font(color=TEXT, bold=True, size=15)
    for r in range(row, row + 3):
        for c in range(start_col, end_col + 1):
            cell = sheet.cell(r, c)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
    sheet.row_dimensions[row].height = 21
    sheet.row_dimensions[row + 1].height = 25
    sheet.row_dimensions[row + 2].height = 14


def _write_table(sheet, start_row, start_col, headers, rows):
    for offset, value in enumerate(headers):
        cell = sheet.cell(start_row, start_col + offset, value)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(color=WHITE, bold=True)
    for row_offset, row in enumerate(rows, start=1):
        for col_offset, value in enumerate(row):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime().replace(tzinfo=None) if value.tzinfo else value.to_pydatetime()
            sheet.cell(start_row + row_offset, start_col + col_offset, value)
    return start_row + len(rows)


def _empty_panel(sheet, anchor: str, title: str):
    panel = PANEL_RANGES.get(anchor)
    if not panel:
        return
    sheet.merge_cells(panel)
    cell = sheet[anchor]
    cell.value = f"{title}\n\nNo data available for this chart."
    cell.fill = PatternFill("solid", fgColor="F8FAFC")
    cell.font = Font(color=MID, bold=True, size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)


def _configure_chart(chart, title: str, show_legend: bool):
    chart.title = title
    chart.height = 6.5
    chart.width = 12.0
    chart.x_axis.title = ""
    chart.y_axis.title = ""
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.tickLblPos = "low"
    chart.y_axis.tickLblPos = "low"
    chart.display_blanks = "zero"
    if show_legend:
        chart.legend.position = "b"
    else:
        chart.legend = None


def _line(sheet, title, start, end, cat_col, first_value_col, last_value_col, anchor):
    if end <= start:
        _empty_panel(sheet, anchor, title)
        return
    chart = LineChart()
    chart.style = 13
    _configure_chart(chart, title, True)
    chart.add_data(
        Reference(sheet, min_col=first_value_col, max_col=last_value_col, min_row=start, max_row=end),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=cat_col, min_row=start + 1, max_row=end))
    sheet.add_chart(chart, anchor)


def _bar(
    sheet,
    title,
    start,
    end,
    cat_col,
    first_value_col,
    last_value_col,
    anchor,
    *,
    stacked=False,
    show_legend=False,
):
    if end <= start:
        _empty_panel(sheet, anchor, title)
        return
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.gapWidth = 60 if stacked else 90
    if stacked:
        chart.grouping = "stacked"
        chart.overlap = 100
    _configure_chart(chart, title, show_legend)
    chart.add_data(
        Reference(sheet, min_col=first_value_col, max_col=last_value_col, min_row=start, max_row=end),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=cat_col, min_row=start + 1, max_row=end))
    sheet.add_chart(chart, anchor)


def add_clickup_executive_dashboard(workbook, result) -> None:
    """Create a ClickUp-only Executive_Dashboard worksheet matching the Streamlit view."""
    if "Executive_Dashboard" in workbook.sheetnames:
        del workbook["Executive_Dashboard"]
    sheet = workbook.create_sheet("Executive_Dashboard", 0)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A12"
    sheet.sheet_properties.tabColor = DARK
    for col in range(1, 17):
        sheet.column_dimensions[get_column_letter(col)].width = 11

    sheet.merge_cells("A1:P1")
    sheet["A1"] = "ClickUp Executive Dashboard"
    sheet["A1"].fill = PatternFill("solid", fgColor=DARK)
    sheet["A1"].font = Font(color=WHITE, bold=True, size=20)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:P2")
    sheet["A2"] = f"Source: ClickUp | Space: {result.get('space_name', 'Unavailable')} | Evaluation cutoff: {result.get('cutoff', 'Unavailable')}"
    sheet["A2"].font = Font(color=MID, italic=True, size=10)
    sheet["A2"].alignment = Alignment(horizontal="center")

    kpis = [
        ("Total Tasks", int(_metric(result, "Total tasks", 0))),
        ("Completed", int(_metric(result, "Completed tasks", 0))),
        ("Completion Rate", _fmt_pct(_metric(result, "Completion rate (%)"))),
        ("On-Time Rate", _fmt_pct(_metric(result, "On-time completion rate (%)"))),
        ("Open Overdue", int(_metric(result, "Open overdue tasks", 0))),
        ("WIP Tasks", int(_metric(result, "WIP tasks", 0))),
    ]
    for index, (label, value) in enumerate(kpis):
        _card(sheet, 1 + index * 2, label, value, 3)

    late = result.get("late_completed_tasks", pd.DataFrame())
    overdue = result.get("overdue_tasks", pd.DataFrame())
    late_mean = None if late.empty else pd.to_numeric(late.get("Due Variance (days)"), errors="coerce").mean()
    overdue_mean = None if overdue.empty else pd.to_numeric(overdue.get("Overdue Days"), errors="coerce").mean()
    averages = [
        ("Avg Execution Time", _fmt_avg(_metric(result, "Average execution hours"), "h")),
        ("Avg Lead Time", _fmt_avg(_metric(result, "Average lead time hours"), "h")),
        ("Avg Time to Start", _fmt_avg(_metric(result, "Average time to start hours"), "h")),
        ("Avg Late Completion", _fmt_avg(late_mean, "days")),
        ("Avg Open Overdue", _fmt_avg(overdue_mean, "days")),
        ("Avg Due Variance", _fmt_avg(_metric(result, "Average due variance (days)"), "days")),
    ]
    for index, (label, value) in enumerate(averages):
        _card(sheet, 1 + index * 2, label, value, 7)

    sheet["A11"] = f"Completed late: {int(_metric(result, 'Completed late tasks', 0))}"
    sheet["D11"] = f"Tasks missing due date: {len(result.get('missing_due_tasks', pd.DataFrame()))}"
    for coordinate in ("A11", "D11"):
        sheet[coordinate].font = Font(bold=True, color=TEXT)

    c = 1
    r = 75

    weekly = result.get("weekly_flow", pd.DataFrame())
    rows = [] if weekly.empty else weekly[["Week Starting", "Tasks Created", "Tasks Completed"]].values.tolist()
    end = _write_table(sheet, r, c, ["Week Starting", "Tasks Created", "Tasks Completed"], rows)
    _line(sheet, "Weekly Task Flow", r, end, c, c + 1, c + 2, "A13")
    r = end + 3

    status = result.get("status_counts", pd.DataFrame())
    rows = [] if status.empty else status[["Status", "Tasks"]].values.tolist()
    end = _write_table(sheet, r, c, ["Status", "Tasks"], rows)
    _bar(sheet, "Task Distribution by Status", r, end, c, c + 1, c + 1, "I13", show_legend=False)
    r = end + 3

    due = result.get("due_status_summary", pd.DataFrame())
    rows = [] if due.empty else due[["Due Status", "Tasks"]].values.tolist()
    end = _write_table(sheet, r, c, ["Due Status", "Tasks"], rows)
    _bar(sheet, "Open Tasks by Due Status", r, end, c, c + 1, c + 1, "A29", show_legend=False)
    r = end + 3

    assignee = result.get("assignee_summary", pd.DataFrame())
    wanted = [col for col in ["Assignee", "Completed", "Open", "WIP"] if col in assignee.columns]
    rows = [] if assignee.empty or len(wanted) < 4 else assignee[wanted].values.tolist()
    end = _write_table(sheet, r, c, ["Assignee", "Completed", "Open", "WIP"], rows)
    _bar(
        sheet,
        "Work Distribution by Assignee",
        r,
        end,
        c,
        c + 1,
        c + 3,
        "I29",
        stacked=True,
        show_legend=True,
    )
    r = end + 3

    variance = result.get("due_variance_summary", pd.DataFrame())
    rows = [] if variance.empty else variance[["Due Variance Category", "Tasks"]].values.tolist()
    end = _write_table(sheet, r, c, ["Due Variance Category", "Tasks"], rows)
    _bar(sheet, "Due Variance Distribution", r, end, c, c + 1, c + 1, "A45", show_legend=False)
    r = end + 3

    status_duration = result.get("status_summary", pd.DataFrame())
    rows = [] if status_duration.empty else status_duration[["Status", "Total Hours"]].values.tolist()
    end = _write_table(sheet, r, c, ["Status", "Total Hours"], rows)
    _bar(sheet, "Total Time in Status", r, end, c, c + 1, c + 1, "I45", show_legend=False)

    sheet["A72"] = "Dashboard chart source data (not included in print area)"
    sheet["A72"].font = Font(color=MID, italic=True, size=9)
    sheet.print_area = "A1:P60"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
