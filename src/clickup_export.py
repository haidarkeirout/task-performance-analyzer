from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

@dataclass
class ClickUpPreparedData:
    """Prepared ClickUp snapshot; deliberately independent of Jira export types."""
    xlsx: bytes = field(repr=False)
    history_json: bytes = field(repr=False)
    cutoff: str
    collected_at: str
    query: str
    fingerprint: str
    count: int
    filename: str
    space_name: str
    source_timezone: str
    filter_summary: str = "All tasks in the selected ClickUp Space"
    filter_criteria: dict = field(default_factory=dict, repr=False)


def _sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(color="FFFFFF", bold=True)
    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _status_minutes(payload):
    """Flatten ClickUp's native status-duration payload for filtering/export."""
    values = {}
    if not isinstance(payload, dict):
        return values
    entries = list(payload.get("status_history") or [])
    current = payload.get("current_status")
    if isinstance(current, dict):
        entries.append(current)
    for entry in entries:
        if not isinstance(entry, dict) or not entry.get("status"):
            continue
        total_time = entry.get("total_time") or {}
        minutes = total_time.get("by_minute") if isinstance(total_time, dict) else None
        try:
            minutes = float(minutes)
        except (TypeError, ValueError):
            continue
        # If a status appears twice, preserve the larger cumulative value.
        status = str(entry["status"])
        values[status] = max(minutes, values.get(status, 0.0))
    return values


def _current_status_info(payload):
    current = payload.get("current_status") if isinstance(payload, dict) else None
    if not isinstance(current, dict):
        return None, None
    total_time = current.get("total_time") or {}
    minutes = total_time.get("by_minute") if isinstance(total_time, dict) else None
    try:
        minutes = float(minutes)
    except (TypeError, ValueError):
        minutes = None
    return minutes, total_time.get("since") if isinstance(total_time, dict) else None


def collect_data(gateway, tasks, space_name, fingerprint, source_timezone="Asia/Damascus", progress=None,
                 space_id=None, time_status_data=None, time_status_error="", filter_summary="",
                 filter_criteria=None):
    """Prepare ClickUp tasks and explicitly supplied status-duration data.

    This function never calls a status/activity endpoint itself.  The UI only
    obtains Total time in Status when the user explicitly selects that More
    filter, then passes the returned map here. Jira's collector and workbook
    are untouched.
    """
    cutoff = datetime.now(timezone.utc).isoformat()
    time_status = {}
    time_status_error = str(time_status_error or "")
    activity_notes = []
    if time_status_data is not None:
        time_status = dict(time_status_data)
    if time_status_error and not time_status:
        activity_notes.append(["", "Total time in Status unavailable", time_status_error])

    flattened = {task_id: _status_minutes(payload) for task_id, payload in time_status.items()}
    status_names = sorted({status for values in flattened.values() for status in values})
    rows = []
    activity_rows = []
    raw = []
    missing_time_status = 0
    for index, task in enumerate(tasks, 1):
        task_id = str(task.get("id", ""))
        if progress:
            progress(f"Preparing ClickUp task {index} of {len(tasks)}...")
        payload = time_status.get(task_id)
        values = flattened.get(task_id, {})
        if payload is None:
            missing_time_status += 1
        status = task.get("status") or {}
        priority = task.get("priority") or {}
        assignees = task.get("assignees") or []
        current_minutes, current_since = _current_status_info(payload)
        history_ids = ""
        rows.append([
            space_id or task.get("space_id"), task_id, task.get("name"),
            ", ".join(str(a.get("username") or a.get("email") or a.get("id")) for a in assignees) or "Unassigned",
            priority.get("priority") if isinstance(priority, dict) else priority,
            status.get("status") if isinstance(status, dict) else status,
            task.get("date_created"), task.get("due_date"), task.get("date_closed"),
            task.get("time_estimate"), task.get("time_spent"),
            json.dumps(payload, ensure_ascii=False) if payload is not None else None,
            current_minutes, current_since, history_ids,
            *[values.get(name) for name in status_names],
        ])
        raw.append([task_id, "task", json.dumps({**task, "selected_space_id": space_id}, ensure_ascii=False)])
        if payload is not None:
            raw.append([task_id, "time_in_status", json.dumps(payload, ensure_ascii=False)])

    wb = Workbook()
    wb.remove(wb.active)
    data_headers = [
        "Space ID", "Task ID", "Task Name", "Assignee", "Priority", "Current Status",
        "Created", "Due Date", "Completed", "Time Estimate (ms)", "Time Spent (ms)",
        "Total Time in Status (JSON)", "Current Status Time (min)", "Current Status Since", "History IDs",
        *[f"Time in Status - {name} (min)" for name in status_names],
    ]
    _sheet(wb, "ClickUp_Data", data_headers, rows)
    _sheet(wb, "Activity", ["Space ID", "Task ID", "History ID", "Timestamp", "User", "Event Type", "Field", "From", "To", "Comment"], activity_rows)
    context_rows = [
        ["Process Name", space_name], ["Space ID", space_id or ""],
        ["Evaluation Scope", filter_summary or "All tasks in the selected ClickUp Space"],
        ["Dataset Type", "ClickUp API collection"],
        ["Evaluation Cutoff Date", cutoff], ["Source Timezone", source_timezone], ["Task Count", len(tasks)],
        ["Activity History", "Disabled; no history collector is used"],
        ["Activity API", "Not collected"],
        ["Total time in Status API", "Available" if time_status else ("Unavailable for this account" if time_status_error else "Not requested")],
        ["Total time in Status Tasks", len(time_status)], ["Missing Total time in Status Tasks", missing_time_status],
        ["History IDs", "Not collected; status-duration data is sourced from Total time in Status."],
    ]
    _sheet(wb, "Process_Context", ["Field", "Value"], context_rows)
    if not activity_notes:
        activity_notes = [["", "Activity history disabled", "Status-duration data is included only when a user selects the Total time in Status filter."]]
    _sheet(wb, "Collection_Notes", ["Task ID", "Issue", "Detail"], activity_notes)
    _sheet(wb, "Raw_JSON", ["Task ID", "Section", "JSON"], raw)

    stream = BytesIO()
    wb.save(stream)
    filename = f"ClickUp_{space_name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    history_payload = {
        "tasks": tasks,
        "time_in_status": time_status,
        "source": "ClickUp API",
        "filter_summary": filter_summary or "All tasks in the selected ClickUp Space",
    }
    prepared = ClickUpPreparedData(
        stream.getvalue(), json.dumps(history_payload, ensure_ascii=False).encode(), cutoff,
        datetime.now(timezone.utc).isoformat(), space_name, fingerprint, len(tasks), filename,
        space_name, source_timezone,
        filter_summary or "All tasks in the selected ClickUp Space", dict(filter_criteria or {}),
    )
    prepared.clickup_activity_available = False
    prepared.clickup_time_status_available = bool(time_status)
    prepared.clickup_time_status_error = time_status_error
    prepared.clickup_activity_notes = activity_notes
    return prepared
