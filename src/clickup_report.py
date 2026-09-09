"""English Word report builder for the isolated ClickUp analysis path."""
from __future__ import annotations

from io import BytesIO
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
HEADER_FILL = "F2F4F7"
USABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def _is_missing(value: Any) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _format_value(value: Any) -> str:
    if _is_missing(value):
        return "Unavailable"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d %H:%M UTC")
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def _format_cutoff(value: Any) -> str:
    parsed = pd.to_datetime(value, utc=True, errors="coerce")
    if pd.isna(parsed):
        return _format_value(value)
    return parsed.strftime("%Y-%m-%d %H:%M UTC")


def _set_run_font(run, *, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_paragraph_spacing(paragraph, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def _set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    grid_columns = list(grid.iterchildren())
    for index, width in enumerate(widths):
        if index < len(grid_columns):
            grid_columns[index].set(qn("w:w"), str(width))
    for row in table.rows:
        keep = row._tr.get_or_add_trPr()
        if keep.find(qn("w:cantSplit")) is None:
            keep.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[index])
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _configure_document(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in tokens.items():
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_spacing(header, after=0)
    run = header.add_run("ClickUp Performance Evaluation")
    _set_run_font(run, size=8, color="6B7280")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_spacing(footer, after=0)
    footer_run = footer.add_run("Generated by Task Performance Intelligence")
    _set_run_font(footer_run, size=8, color="6B7280")


def _add_title(document: Document, title: str, subtitle: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_paragraph_spacing(paragraph, before=0, after=3, line=1.1)
    run = paragraph.add_run(title)
    _set_run_font(run, size=20, bold=True, color=INK)
    sub = document.add_paragraph()
    _set_paragraph_spacing(sub, before=0, after=16, line=1.1)
    sub_run = sub.add_run(subtitle)
    _set_run_font(sub_run, size=10, color="5B6470")


def _add_heading(document: Document, text: str, level=1):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        _set_run_font(
            run,
            size={1: 16, 2: 13, 3: 12}.get(level, 11),
            bold=True,
            color={1: BLUE, 2: BLUE, 3: DARK_BLUE}.get(level, INK),
        )
    return paragraph


def _add_body(document: Document, text: str):
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph)
    run = paragraph.add_run(text)
    _set_run_font(run, size=11)
    return paragraph


def _add_bullet(document: Document, text: str):
    paragraph = document.add_paragraph(style="List Bullet")
    _set_paragraph_spacing(paragraph, after=4)
    run = paragraph.add_run(text)
    _set_run_font(run, size=11)
    return paragraph


def _add_table(document: Document, headers: list[str], rows: list[list[Any]], widths: list[int]):
    table = document.add_table(rows=1, cols=len(headers))
    _set_table_geometry(table, widths)
    _set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_shading(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_paragraph_spacing(paragraph, after=0, line=1.0)
        run = paragraph.add_run(header)
        _set_run_font(run, size=9, bold=True, color=INK)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            paragraph = cells[index].paragraphs[0]
            _set_paragraph_spacing(paragraph, after=0, line=1.0)
            run = paragraph.add_run(_format_value(value))
            _set_run_font(run, size=9)
    _set_table_geometry(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def _frame_rows(frame: pd.DataFrame, columns: list[str], limit=None):
    if frame is None or frame.empty:
        return []
    visible = frame.head(limit) if limit else frame
    return [[row.get(column) for column in columns] for row in visible.to_dict("records")]


def _overall_value(result: dict, metric: str):
    overall = result["overall"]
    matches = overall.loc[overall["Metric"].eq(metric), "Value"]
    return None if matches.empty else matches.iloc[0]


def _safe_float(value):
    if value is None or _is_missing(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _clickup_recommendations(result: dict) -> list[str]:
    recommendations: list[str] = []
    total = int(_overall_value(result, "Total tasks") or 0)
    completed = int(_overall_value(result, "Completed tasks") or 0)
    overdue = int(_overall_value(result, "Open overdue tasks") or 0)
    wip = int(_overall_value(result, "WIP tasks") or 0)
    missing_due = len(result.get("missing_due_tasks", pd.DataFrame()))
    completed_late = int(_overall_value(result, "Completed late tasks") or 0)
    on_time_rate = _safe_float(_overall_value(result, "On-time completion rate (%)"))
    completion_rate = _safe_float(_overall_value(result, "Completion rate (%)"))

    if overdue > 0:
        recommendations.append(
            f"Prioritize the {overdue} open overdue task(s): confirm blockers, owners, and realistic due dates, then review them in the next operating cadence."
        )
    if on_time_rate is not None and on_time_rate < 80:
        recommendations.append(
            f"Improve schedule reliability: the on-time completion rate is {on_time_rate:.1f}%. Review estimation, due-date setting, and early escalation for tasks at risk."
        )
    if completion_rate is not None and completion_rate < 70:
        recommendations.append(
            f"Review backlog conversion: {completed} of {total} tasks are completed ({completion_rate:.1f}%). Prioritize the highest-value open work and remove or re-scope stale items."
        )
    if total and wip / total >= 0.30 and wip > 0:
        recommendations.append(
            f"Control work in progress: {wip} task(s) are WIP ({wip / total * 100:.1f}% of scope). Consider WIP limits and finishing active work before starting additional items."
        )
    if completed_late > 0:
        recommendations.append(
            f"Review the {completed_late} completed-late task(s) to identify recurring estimation, dependency, or prioritization patterns and address them in future planning."
        )
    if missing_due > 0:
        recommendations.append(
            f"Improve due-date coverage for the {missing_due} task(s) without a due date so future timeliness and overdue analysis is more complete."
        )
    if not recommendations:
        recommendations.append(
            "Maintain the current operating controls and continue monitoring completion, timeliness, WIP, due-date coverage, and status-duration evidence at each evaluation cutoff."
        )
    return recommendations


def create_clickup_word_report(result: dict) -> bytes:
    document = Document()
    _configure_document(document)
    _add_title(
        document,
        "ClickUp Task Performance Evaluation Report",
        f"Space: {result['space_name']} | Evaluation cutoff: {_format_cutoff(result['cutoff'])}",
    )

    _add_heading(document, "Executive Summary")
    _add_body(
        document,
        "This report evaluates the selected ClickUp task snapshot. It presents task-level outcomes, overall indicators, due-date adherence, status-duration evidence when available, exceptions, and data-quality limits."
    )
    _add_table(
        document,
        ["Metric", "Value"],
        _frame_rows(result["overall"], ["Metric", "Value"]),
        [2700, 6660],
    )

    _add_heading(document, "Scope and Methodology")
    context = result["analysis_context"].copy()
    context.loc[context["Field"].eq("Evaluation Cutoff"), "Value"] = _format_cutoff(result["cutoff"])
    _add_table(document, ["Field", "Value"], _frame_rows(context, ["Field", "Value"]), [2700, 6660])
    _add_body(
        document,
        "Due Variance is measured in calendar days. For completed tasks it equals Completed date minus Due date. For open tasks it equals the evaluation cutoff minus Due date. Positive values are late or overdue; negative values are early or still have time remaining."
    )

    _add_heading(document, "Due Variance and Timeliness")
    due_average = _overall_value(result, "Average due variance (days)")
    on_time = _overall_value(result, "On-time completion rate (%)")
    _add_body(
        document,
        f"Average completed-task due variance: {_format_value(due_average)} days. On-time completion rate: {_format_value(on_time)}%."
    )
    due_variance_rows = _frame_rows(
        result["due_variance_summary"],
        ["Due Variance Category", "Tasks", "Average Variance (days)", "Median Variance (days)"],
    )
    if due_variance_rows:
        _add_table(
            document,
            ["Due Variance Category", "Tasks", "Average Variance (days)", "Median Variance (days)"],
            due_variance_rows,
            [2640, 1200, 2760, 2760],
        )
    else:
        _add_body(document, "No tasks with due dates were available for Due Variance analysis.")

    _add_heading(document, "Open Work by Due Status", 2)
    due_status_rows = _frame_rows(result["due_status_summary"], ["Due Status", "Tasks", "Share of Open Tasks (%)"])
    if due_status_rows:
        _add_table(
            document,
            ["Due Status", "Tasks", "Share of Open Tasks (%)"],
            due_status_rows,
            [3900, 1500, 3960],
        )
    else:
        _add_body(document, "No open tasks were available for due-date grouping.")

    _add_heading(document, "Workflow and Status Analysis")
    _add_body(
        document,
        "Status-duration values are reported only when ClickUp exposes Total time in Status for the selected account. The system does not reconstruct activity history or replace missing values with zero."
    )
    status_rows = _frame_rows(result["status_summary"], ["Status", "Tasks", "Average Hours", "Median Hours", "Total Hours"])
    if status_rows:
        _add_table(
            document,
            ["Status", "Tasks", "Average Hours", "Median Hours", "Total Hours"],
            status_rows,
            [2160, 960, 2040, 2040, 2160],
        )
    else:
        _add_body(document, "No status-duration values were returned for this selected ClickUp scope.")

    _add_heading(document, "Bottlenecks, Exceptions, and Follow-up")
    _add_table(
        document,
        ["Finding", "Tasks", "Recommended Follow-up"],
        _frame_rows(result["findings"], ["Finding", "Tasks", "Recommended Follow-up"]),
        [2640, 960, 5760],
    )
    if result["late_completed_tasks"].empty and result["overdue_tasks"].empty:
        _add_body(document, "No completed-late or open-overdue tasks were found in the selected scope.")
    else:
        _add_body(document, "The task-level section identifies every task in the selected scope so the listed exceptions can be reviewed in context.")

    _add_heading(document, "Recommendations")
    _add_body(
        document,
        "The following recommendations are generated from the ClickUp results in this report and are limited to the selected scope and evaluation cutoff."
    )
    for recommendation in _clickup_recommendations(result):
        _add_bullet(document, recommendation)

    _add_heading(document, "Individual Achievements and Assignment Summary")
    _add_body(
        document,
        "Assignment values use the assignee snapshot returned by ClickUp. They support workload visibility and do not establish individual contribution or ownership at completion."
    )
    assignee_rows = _frame_rows(
        result["assignee_summary"],
        ["Assignee", "Total Tasks", "Completed", "Open", "On-Time Completion Rate (%)", "Average Due Variance (days)"],
    )
    if assignee_rows:
        _add_table(
            document,
            ["Assignee", "Tasks", "Completed", "Open", "On-Time Rate (%)", "Avg Due Variance (days)"],
            assignee_rows,
            [2400, 900, 1200, 900, 1980, 1980],
        )
    else:
        _add_body(document, "No assignee information was returned for this selected scope.")

    _add_heading(document, "Task-Level Evaluation")
    task_columns = [
        "Current Status", "Assignee", "Priority", "Created", "Start Date", "Due Date", "Completed",
        "On Time?", "Due Variance (days)", "Due Variance Basis", "Lead Time Hours", "Execution Hours",
        "Overdue Days", "Total Time in Status (min)",
    ]
    for task in result["tasks"].to_dict("records"):
        _add_heading(document, f"{_format_value(task.get('Task ID'))} - {_format_value(task.get('Task Name'))}", 2)
        _add_table(
            document,
            ["Field", "Value"],
            [[label, task.get(label)] for label in task_columns],
            [2700, 6660],
        )

    _add_heading(document, "Data Quality and Limitations")
    _add_table(
        document,
        ["Check", "Value", "Status"],
        _frame_rows(result["quality"], ["Check", "Value", "Status"]),
        [5040, 1500, 2820],
    )
    _add_heading(document, "Metric Definitions")
    _add_table(
        document,
        ["Metric", "Definition"],
        _frame_rows(result["metric_definitions"], ["Metric", "Definition"]),
        [2700, 6660],
    )
    _add_body(
        document,
        "This report is a task-process evaluation based on the selected ClickUp snapshot. It should be interpreted together with the documented data-quality limits and should not be used as a standalone employee-performance assessment."
    )

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
