"""Application-facing orchestration for Company Performance.

The source collectors remain the owners of authentication and API access.
This module only reads their already-prepared payloads, turns them into the
isolated company contract, and reconstructs a selected calendar period.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

from .adapters import (
    CompanyCollection,
    adapt_clickup_prepared,
    adapt_jira_collection,
    combine_company_sources,
)
from .dashboard import CompanyDashboardModel, DashboardFilters, build_company_dashboard
from .kpis import BottleneckCandidate, Recommendation, generate_recommendations, identify_bottleneck_candidates
from .models import ParentClassification, TaskPeriodSnapshot
from .normalization import assignee_group, calendar_date, deduplicate_tasks, normalize_priority, normalize_status
from .workflow import reconstruct_task


@dataclass(frozen=True)
class CompanyAnalysisResult:
    """One complete, source-independent company analysis run."""

    collection: CompanyCollection
    snapshots: tuple[TaskPeriodSnapshot, ...]
    model: CompanyDashboardModel
    bottlenecks: tuple[BottleneckCandidate, ...]
    recommendations: tuple[Recommendation, ...]


@dataclass(frozen=True)
class CompanyPreparedItem:
    """Prepared source data plus the company mapping metadata for one Space."""

    prepared: Any
    source_tool: str
    source_space: str
    project_name: str
    company_id: str | None = None
    company_name: str | None = None
    source_id: str | None = None
    source_kind: str | None = None


@dataclass(frozen=True)
class CompanyPreviewRow:
    """One source-neutral task row shown before Company Performance analysis."""

    source_tool: str
    space: str | None
    task_name: str | None
    original_status: str | None
    final_status: str
    assignee: str
    priority: str
    due_date: date | None
    project_name: str | None = None
    company_name: str | None = None
    department: str | None = None
    department_id: str | None = None
    issue_type: str | None = None
    parent_id: str | None = None
    parent_classification: str | None = None
    epic_name: str | None = None


def _value(row: Mapping[str, Any], name: str) -> Any:
    value = row.get(name)
    return None if value == "" else value


def _row_value(row: Mapping[str, Any], *names: str) -> Any:
    """Read a Jira export column while tolerating collector naming variants."""
    for name in names:
        value = _value(row, name)
        if value is not None:
            return value
    expected = {name.strip().casefold() for name in names}
    for key, value in row.items():
        if str(key).strip().casefold() in expected and value not in (None, ""):
            return value
    return None


def _jira_reference(value: Any) -> Any:
    """Restore a Jira relation stored as a plain value or exported JSON text."""
    if isinstance(value, Mapping):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            return value
        return parsed if isinstance(parsed, Mapping) else value
    return value


def _jira_export_issues(prepared_data: Any) -> list[dict[str, Any]]:
    """Restore the minimal Jira issue payload from the source XLSX.

    Jira's existing ``PreparedData`` deliberately persists a history JSON but
    not a second copy of every API issue object.  Its source workbook is the
    audited collector output, so rebuilding the adapter input from ``Jira_Data``
    keeps Company Performance additive and avoids a new API request.
    """
    try:
        workbook = load_workbook(BytesIO(prepared_data.xlsx), read_only=True, data_only=True)
        if "Jira_Data" not in workbook.sheetnames:
            return []
        sheet = workbook["Jira_Data"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value) if value is not None else "" for value in next(rows)]
    except (AttributeError, KeyError, OSError, StopIteration, TypeError, ValueError):
        return []

    output: list[dict[str, Any]] = []
    for values in rows:
        row = dict(zip(headers, values))
        key = _row_value(row, "Issue key") or _row_value(row, "Issue id")
        if not key:
            continue
        assignee = _row_value(row, "Assignee")
        parent = _jira_reference(_row_value(row, "Parent", "Parent key", "Parent ID"))
        epic_link = _jira_reference(_row_value(
            row, "Epic Link", "Epic link", "Custom field (Epic Link)"
        ))
        output.append({
            "key": str(key),
            "id": str(_row_value(row, "Issue id") or ""),
            "fields": {
                "summary": _row_value(row, "Summary"),
                "issuetype": {"name": _row_value(row, "Issue Type")},
                "project": {"key": _row_value(row, "Project key"), "name": _row_value(row, "Project name")},
                "company_name": _row_value(row, "Company", "Company name", "Custom field (Company)", "Client", "Customer"),
                "status": {"name": _row_value(row, "Status")},
                "priority": {"name": _row_value(row, "Priority")},
                "assignee": {"displayName": assignee} if assignee else None,
                "created": _row_value(row, "Created"),
                "duedate": _row_value(row, "Due date"),
                "start_date": _row_value(row, "Custom field (Start date)"),
                "parent": parent,
                "epic_link": epic_link,
            },
        })
    return output


def _jira_histories(prepared_data: Any) -> dict[str, Mapping[str, Any]]:
    try:
        raw = prepared_data.history_json
        payload = json.loads(raw.decode("utf-8") if isinstance(raw, bytes) else raw)
    except (AttributeError, TypeError, UnicodeDecodeError, json.JSONDecodeError):
        return {}
    return payload if isinstance(payload, dict) else {}


def _prepared_metadata(item: Any) -> tuple[Any, str | None, str | None]:
    """Unwrap automatic Company collection metadata when present."""
    if isinstance(item, CompanyPreparedItem):
        return item.prepared, item.project_name, item.source_space
    return (
        item,
        None,
        getattr(item, "space_name", None),
    )


def _prepared_company_name(item: Any) -> str | None:
    """Return the discovered company label when collection metadata exists."""
    if isinstance(item, CompanyPreparedItem):
        return item.company_name or item.project_name
    return getattr(item, "company_name", None)


def _adapt_jira_prepared(item: Any, unified_project: str):
    prepared_data, _, source_space = _prepared_metadata(item)
    return adapt_jira_collection(
        _jira_export_issues(prepared_data),
        _jira_histories(prepared_data),
        unified_project=unified_project,
        source_space=source_space,
        collection_timestamp=getattr(prepared_data, "collected_at", None),
    )


def _prepared_items(value: Any | None) -> tuple[Any, ...]:
    """Accept the legacy single prepared result or a Company multi-space sequence."""
    if value is None:
        return ()
    if isinstance(value, (list, tuple)):
        return tuple(item for item in value if item is not None)
    return (value,)


def build_company_preview(
    *,
    jira_prepared: Any | None = None,
    clickup_prepared: Any | None = None,
) -> tuple[CompanyPreviewRow, ...]:
    """Build the combined source preview without running historical analysis."""
    sources = []
    prepared_by_scope: dict[tuple[str, str | None, str | None], str | None] = {}
    for prepared in _prepared_items(jira_prepared):
        _, project_name, _ = _prepared_metadata(prepared)
        source_space = _prepared_metadata(prepared)[2]
        prepared_by_scope[("Jira", source_space, project_name)] = _prepared_company_name(prepared)
        sources.append(_adapt_jira_prepared(prepared, project_name or "Preview"))
    for prepared in _prepared_items(clickup_prepared):
        prepared_data, project_name, source_space = _prepared_metadata(prepared)
        prepared_by_scope[("ClickUp", source_space, project_name)] = _prepared_company_name(prepared)
        sources.append(
            adapt_clickup_prepared(
                prepared_data,
                unified_project=project_name or "Preview",
                source_space=source_space,
            )
        )

    if not sources:
        return ()
    records = tuple(
        record
        for record in deduplicate_tasks(combine_company_sources(*sources).records)
        if record.parent_classification is not ParentClassification.CONTAINER
    )
    return tuple(
        CompanyPreviewRow(
            source_tool=task.source_tool,
            space=task.source_space,
            task_name=task.task_name,
            original_status=task.raw_status,
            final_status=normalize_status(task.source_tool, task.raw_status).value,
            assignee=assignee_group(task),
            priority=normalize_priority(task.source_tool, task.priority),
            due_date=task.due_date,
            project_name=task.unified_project,
            company_name=prepared_by_scope.get(
                (task.source_tool, task.source_space, task.unified_project)
            ),
            department=task.department,
            department_id=task.department_id,
            issue_type=task.issue_type or (
                "Sub-task" if task.parent_classification.value == "Subtask" else "Task"
            ),
            parent_id=task.parent_id,
            parent_classification=task.parent_classification.value,
            epic_name=task.epic_name,
        )
        for task in records
    )


def filter_company_preview(
    rows: Iterable[CompanyPreviewRow],
    *,
    source_tools: Sequence[str] = (),
    spaces: Sequence[str] = (),
    task_name: str = "",
    original_statuses: Sequence[str] = (),
    final_statuses: Sequence[str] = (),
    assignees: Sequence[str] = (),
    priorities: Sequence[str] = (),
    due_date_start: date | None = None,
    due_date_end: date | None = None,
) -> tuple[CompanyPreviewRow, ...]:
    """Filter the combined preview using the approved pre-analysis fields."""
    sources = {value.casefold() for value in source_tools}
    selected_spaces = set(spaces)
    selected_originals = set(original_statuses)
    selected_finals = set(final_statuses)
    selected_assignees = set(assignees)
    selected_priorities = set(priorities)
    task_name_filter = task_name.strip().casefold()
    selected: list[CompanyPreviewRow] = []

    for row in rows:
        if sources and row.source_tool.casefold() not in sources:
            continue
        if selected_spaces and (row.space or "") not in selected_spaces:
            continue
        if task_name_filter and task_name_filter not in (row.task_name or "").casefold():
            continue
        if selected_originals and (row.original_status or "") not in selected_originals:
            continue
        if selected_finals and row.final_status not in selected_finals:
            continue
        if selected_assignees and row.assignee not in selected_assignees:
            continue
        if selected_priorities and row.priority not in selected_priorities:
            continue
        if due_date_start is not None and (row.due_date is None or row.due_date < due_date_start):
            continue
        if due_date_end is not None and (row.due_date is None or row.due_date > due_date_end):
            continue
        selected.append(row)
    return tuple(selected)


def build_company_analysis(
    *,
    period_start: date | None,
    period_end: date | None,
    jira_prepared: Any | None = None,
    jira_project: str | None = None,
    clickup_prepared: Any | None = None,
    clickup_project: str | None = None,
    unified_project: str | None = None,
    filters: DashboardFilters | None = None,
) -> CompanyAnalysisResult:
    """Build a company run from one or more existing prepared source results.

    ``jira_project`` and ``clickup_project`` remain supported for the existing
    single-source callers.  Company multi-space selection can instead provide
    one explicit ``unified_project`` for the selected Jira and ClickUp spaces.
    """
    jira_items = _prepared_items(jira_prepared)
    clickup_items = _prepared_items(clickup_prepared)
    shared_project = (unified_project or "").strip()
    sources = []

    if jira_items:
        for item in jira_items:
            _, item_project, _ = _prepared_metadata(item)
            project = shared_project or item_project or (jira_project or "").strip()
            if not project:
                raise ValueError("A project mapping is missing for one Jira Space.")
            sources.append(_adapt_jira_prepared(item, project))

    if clickup_items:
        for item in clickup_items:
            prepared_data, item_project, source_space = _prepared_metadata(item)
            project = shared_project or item_project or (clickup_project or "").strip()
            if not project:
                raise ValueError("A project mapping is missing for one ClickUp Space.")
            sources.append(
                adapt_clickup_prepared(
                    prepared_data,
                    unified_project=project,
                    source_space=source_space,
                )
            )

    if not sources:
        raise ValueError("Collect at least one Jira or ClickUp space before running Company Performance.")

    collection = combine_company_sources(*sources)
    records = tuple(
        record
        for record in deduplicate_tasks(collection.records)
        if record.parent_classification is not ParentClassification.CONTAINER
    )
    coverage_dates = [
        calendar_date(record.history_through or record.collection_timestamp)
        for record in records
    ]
    usable_coverage = [value for value in coverage_dates if value is not None]
    if period_start is None or period_end is None:
        automatic_end = min(usable_coverage) if usable_coverage else None
        if automatic_end is None:
            raise ValueError("The selected sources do not contain a usable collection cutoff.")
        created_dates = [record.created_date for record in records if record.created_date is not None]
        automatic_start = min(created_dates) if created_dates else automatic_end
        if automatic_start > automatic_end:
            automatic_start = automatic_end
        period_start = period_start or automatic_start
        period_end = period_end or automatic_end
    if period_end < period_start:
        raise ValueError("Analysis period end must not be before its start.")
    if usable_coverage and period_end > min(usable_coverage):
        raise ValueError(
            "Analysis Period To cannot be later than the collected source coverage. "
            f"Choose {min(usable_coverage).isoformat()} or earlier, or collect the sources again."
        )
    snapshots = tuple(
        reconstruct_task(
            record,
            period_start,
            period_end,
            collection_date=calendar_date(record.collection_timestamp),
        )
        for record in records
    )
    if not snapshots:
        raise ValueError("The selected sources contain no usable task records.")
    model = build_company_dashboard(snapshots, coverages=collection.coverages, filters=filters)
    analysis_snapshots = tuple(snapshot for snapshot in snapshots if snapshot.in_scope)
    return CompanyAnalysisResult(
        collection=collection,
        snapshots=snapshots,
        model=model,
        bottlenecks=identify_bottleneck_candidates(analysis_snapshots),
        recommendations=generate_recommendations(analysis_snapshots),
    )
