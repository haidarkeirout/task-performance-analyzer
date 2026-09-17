"""Department-specific presentation and exports built on normalized ClickUp results."""
from __future__ import annotations

from io import BytesIO
from datetime import date, datetime, time
from typing import Any
from zoneinfo import ZoneInfo

import pandas as pd
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _metric(result: dict, name: str, default=None):
    matches = result["overall"].loc[result["overall"]["Metric"].eq(name), "Value"]
    return default if matches.empty else matches.iloc[0]


def _rate(numerator: int, denominator: int):
    return None if not denominator else numerator / denominator * 100.0


def filter_jira_department_period(
    task_metrics: pd.DataFrame,
    histories: dict,
    period_start,
    cutoff,
    source_timezone: str,
) -> pd.DataFrame:
    """Keep Jira work that was active or changed during the period.

    The Jira search intentionally collects a conservative candidate set.  Full
    histories then decide exact inclusion so an older task that stayed open or
    reopened during the selected period is not lost, while work closed before
    the period with no in-period activity is removed.
    """
    if task_metrics is None or task_metrics.empty or period_start in (None, ""):
        return task_metrics

    zone = ZoneInfo(source_timezone)
    if isinstance(period_start, str):
        start_day = date.fromisoformat(period_start[:10])
    elif isinstance(period_start, datetime):
        start_day = period_start.date()
    else:
        start_day = period_start
    start = pd.Timestamp(datetime.combine(start_day, time.min, tzinfo=zone)).tz_convert("UTC")
    end = pd.to_datetime(cutoff, utc=True, errors="coerce")
    if pd.isna(end):
        raise ValueError("The Department analysis cutoff is invalid.")

    created = pd.to_datetime(task_metrics.get("created_at"), utc=True, errors="coerce")
    completed = pd.to_datetime(task_metrics.get("completed_at"), utc=True, errors="coerce")
    open_at_end = task_metrics.get(
        "is_open", pd.Series(False, index=task_metrics.index)
    ).fillna(False).eq(True)
    include = created.ge(start) | open_at_end | completed.ge(start)

    changed_in_period: set[str] = set()
    for key, history in (histories or {}).items():
        for event in (history or {}).get("status_events") or ():
            changed = pd.to_datetime(event.get("changed_at"), utc=True, errors="coerce")
            if pd.notna(changed) and start <= changed <= end:
                changed_in_period.add(str(key))
                break
    issue_keys = task_metrics.get(
        "issue_key", pd.Series("", index=task_metrics.index)
    ).fillna("").astype(str)
    include |= issue_keys.isin(changed_in_period)

    result = task_metrics.loc[include].copy()
    result.attrs.update(task_metrics.attrs)
    result.attrs["department_period_start"] = start.isoformat()
    return result


def build_department_result(clickup_result: dict, prepared_data) -> dict:
    """Create the isolated department view without changing legacy metrics."""
    tasks = clickup_result["tasks"].copy()
    total = len(tasks)
    is_subtask = tasks.get("Is Subtask", pd.Series(False, index=tasks.index)).fillna(False).astype(bool)
    parent_tasks = tasks.loc[~is_subtask].copy()
    known_parent_tasks = parent_tasks.loc[
        parent_tasks.get("Status Known?", pd.Series(True, index=parent_tasks.index)).fillna(False).eq(True)
    ].copy()
    cancelled = int(parent_tasks["Cancelled?"].sum()) if not parent_tasks.empty else 0
    completed = int(parent_tasks["Completed?"].sum()) if not parent_tasks.empty else 0
    all_parent_statuses_known = len(known_parent_tasks) == len(parent_tasks)
    completion_denominator = len(parent_tasks) if all_parent_statuses_known else 0
    open_tasks = int(tasks["Open?"].sum()) if total else 0
    open_due_tasks = int((tasks["Open?"] & tasks["Due Date"].notna()).sum()) if total else 0
    overdue = int((
        tasks["Open?"]
        & pd.to_numeric(tasks["Overdue Days"], errors="coerce").gt(0)
    ).sum()) if total else 0
    due_known = (
        parent_tasks["Completed?"] & parent_tasks["Due Date"].notna() & parent_tasks["Completed"].notna()
        if not parent_tasks.empty else pd.Series(dtype=bool)
    )
    on_time = int((due_known & parent_tasks["On Time?"].eq(True)).sum()) if not parent_tasks.empty else 0

    kpis = pd.DataFrame([
        ("Total Tasks", total, total, total),
        ("Task Completion Rate (%)", _rate(completed, completion_denominator), completed, completion_denominator),
        ("On-Time Completion Rate (%)", _rate(on_time, int(due_known.sum())), on_time, int(due_known.sum())),
        ("Open Overdue Tasks", overdue, overdue, open_due_tasks),
        ("Open Overdue Rate (%)", _rate(overdue, open_due_tasks), overdue, open_due_tasks),
        ("Average Lead Time (hours)", pd.to_numeric(tasks.loc[tasks["Completed?"], "Lead Time Hours"], errors="coerce").mean(), completed, completed),
        ("Workflow Exception Rate (%)", None, 0, 0),
        ("Cancelled/Rejected Tasks", cancelled, cancelled, total),
        ("Cancellation/Rejected Rate (%)", _rate(cancelled, total), cancelled, total),
    ], columns=["KPI", "Value", "Numerator", "Denominator"])

    employees = clickup_result["assignee_summary"].copy()
    if not employees.empty:
        employees = employees.rename(columns={"Total Tasks": "Total Assigned"})
        overdue_by_employee = (tasks.assign(_overdue=tasks["Open?"] & pd.to_numeric(tasks["Overdue Days"], errors="coerce").gt(0))
                               .groupby("Assignee", dropna=False)["_overdue"].sum())
        employees["Overdue"] = employees["Assignee"].map(overdue_by_employee).fillna(0).astype(int)
        employees["Workflow Exceptions"] = None

    attention_mask = pd.Series(False, index=tasks.index)
    if total:
        priority = tasks["Priority"].fillna("").astype(str).str.casefold()
        attention_mask = (
            (tasks["Open?"] & tasks["Overdue Days"].notna())
            | tasks["Assignee"].fillna("Unassigned").eq("Unassigned")
            | tasks["Due Date"].isna()
            | (tasks["Open?"] & priority.isin({"urgent", "high"}) & tasks["Overdue Days"].notna())
        )
    attention = tasks.loc[attention_mask].copy()
    if not attention.empty:
        def issue(row):
            issues = []
            if row["Open?"] and pd.notna(pd.to_numeric(row["Overdue Days"], errors="coerce")) and pd.to_numeric(row["Overdue Days"], errors="coerce") > 0: issues.append("Open overdue")
            if str(row.get("Priority", "")).casefold() in {"urgent", "high"} and row["Open?"] and pd.notna(pd.to_numeric(row["Overdue Days"], errors="coerce")) and pd.to_numeric(row["Overdue Days"], errors="coerce") > 0: issues.append("High-priority overdue")
            if row.get("Assignee") in (None, "", "Unassigned"): issues.append("Unassigned")
            if pd.isna(row.get("Due Date")): issues.append("Missing due date")
            return "; ".join(issues)
        attention["Issue"] = attention.apply(issue, axis=1)

    bottlenecks = clickup_result["status_summary"].copy()
    if not bottlenecks.empty:
        bottlenecks = bottlenecks.sort_values(["Total Hours", "Tasks"], ascending=False).head(5)
        bottlenecks["Interpretation"] = "Bottleneck candidate; review workload and blockers before concluding cause."

    quality_rows = []
    for row in clickup_result["quality"].to_dict("records"):
        if row.get("Status") != "OK":
            quality_rows.append({"Issue Type": row["Check"], "Count": row["Value"], "Analysis Impact": row["Status"]})
    quality = pd.DataFrame(quality_rows, columns=["Issue Type", "Count", "Analysis Impact"])
    return {
        **clickup_result,
        "analysis_type": "Department Performance",
        "department_name": prepared_data.department_name,
        "department_id": prepared_data.department_id,
        "kpis": kpis,
        "employee_breakdown": employees,
        "attention": attention,
        "bottlenecks": bottlenecks,
        "department_quality": quality,
        "space_names": getattr(prepared_data, "space_names", [prepared_data.space_name]),
        "data_source": "ClickUp",
    }


def build_jira_department_result(task_metrics: pd.DataFrame, process_data: dict, space_name: str, cutoff) -> dict:
    """Adapt verified Jira metrics to the same department presentation contract."""
    source = task_metrics.copy()
    issue_types = source.get("issue_type", pd.Series("", index=source.index)).fillna("").astype(str)
    is_subtask = issue_types.str.casefold().str.replace("-", "", regex=False).str.replace(" ", "", regex=False).eq("subtask")
    tasks = pd.DataFrame({
        "Task ID": source.get("issue_key"), "Task Name": source.get("task_name"),
        "Assignee": source.get("assignee_name", pd.Series("Unassigned", index=source.index)).fillna("Unassigned"),
        "Priority": source.get("priority"), "Task Type": source.get("issue_type"),
        "Current Status": source.get("status_at_cutoff"), "Created": source.get("created_at"),
        "Start Date": source.get("actual_start_at"), "Due Date": source.get("due_date"),
        "Completed": source.get("completed_at"), "Completed?": source.get("is_completed", False),
        "Cancelled?": source.get("is_rejected", False), "Open?": source.get("is_open", False),
        "WIP?": source.get("is_wip", False), "Lead Time Hours": source.get("lead_time_business_hours"),
        "Execution Hours": source.get("execution_business_hours"),
        "On Time?": source.get("on_time_completion"), "Overdue Days": source.get("overdue_days"),
        "Is Subtask": is_subtask,
    })
    parent_tasks = tasks.loc[~tasks["Is Subtask"]].copy()
    total = len(tasks); cancelled = int(parent_tasks["Cancelled?"].eq(True).sum())
    completed = int(parent_tasks["Completed?"].eq(True).sum()); open_count = int(tasks["Open?"].eq(True).sum())
    completion_denominator = len(parent_tasks)
    open_due_tasks = int((tasks["Open?"].eq(True) & tasks["Due Date"].notna()).sum())
    overdue = int((tasks["Open?"].eq(True) & pd.to_numeric(tasks["Overdue Days"], errors="coerce").gt(0)).sum())
    due_known = parent_tasks["Completed?"].eq(True) & parent_tasks["Due Date"].notna() & parent_tasks["Completed"].notna()
    on_time = int((due_known & parent_tasks["On Time?"].eq(True)).sum())
    history_valid = source.get("history_complete", pd.Series(False, index=source.index)).eq(True)
    exception = (
        source.get("rework_count", pd.Series(0, index=source.index)).fillna(0).gt(0)
        | source.get("replanning_count", pd.Series(0, index=source.index)).fillna(0).gt(0)
        | source.get("re_evaluation_count", pd.Series(0, index=source.index)).fillna(0).gt(0)
    )
    exception_count = int((history_valid & exception).sum()); exception_denominator = int(history_valid.sum())
    kpis = pd.DataFrame([
        ("Total Tasks", total, total, total),
        ("Task Completion Rate (%)", _rate(completed, completion_denominator), completed, completion_denominator),
        ("On-Time Completion Rate (%)", _rate(on_time, int(due_known.sum())), on_time, int(due_known.sum())),
        ("Open Overdue Tasks", overdue, overdue, open_due_tasks),
        ("Open Overdue Rate (%)", _rate(overdue, open_due_tasks), overdue, open_due_tasks),
        ("Average Lead Time (hours)", pd.to_numeric(parent_tasks.loc[parent_tasks["Completed?"].eq(True), "Lead Time Hours"], errors="coerce").mean(), completed, completed),
        ("Workflow Exception Rate (%)", _rate(exception_count, exception_denominator), exception_count, exception_denominator),
        ("Cancelled/Rejected Tasks", cancelled, cancelled, total),
        ("Cancellation/Rejected Rate (%)", _rate(cancelled, total), cancelled, total),
    ], columns=["KPI", "Value", "Numerator", "Denominator"])
    employees = (tasks.groupby("Assignee", dropna=False)
                 .agg(**{"Total Assigned": ("Task ID", "count"), "Completed": ("Completed?", "sum"),
                         "Open": ("Open?", "sum"), "WIP": ("WIP?", "sum")}).reset_index())
    overdue_map = tasks.assign(_overdue=tasks["Open?"].eq(True) & pd.to_numeric(tasks["Overdue Days"], errors="coerce").gt(0)).groupby("Assignee")["_overdue"].sum()
    employees["Overdue"] = employees["Assignee"].map(overdue_map).fillna(0).astype(int)
    employees["Workflow Exceptions"] = source.assign(_assignee=tasks["Assignee"], _exception=exception).groupby("_assignee")["_exception"].sum().reindex(employees["Assignee"]).to_numpy()
    attention_mask = tasks["Open?"].eq(True) & pd.to_numeric(tasks["Overdue Days"], errors="coerce").gt(0)
    attention_mask |= tasks["Assignee"].eq("Unassigned") | tasks["Due Date"].isna() | exception
    attention = tasks.loc[attention_mask].copy()
    if not attention.empty:
        attention["Issue"] = ["; ".join(filter(None, [
            "Open overdue" if bool(row["Open?"]) and pd.notna(row["Overdue Days"]) and row["Overdue Days"] > 0 else "",
            "Unassigned" if row["Assignee"] == "Unassigned" else "",
            "Missing due date" if pd.isna(row["Due Date"]) else "",
            "Workflow exception" if bool(exception.loc[index]) else "",
        ])) for index, row in attention.iterrows()]
    stage = process_data.get("stage_summary", pd.DataFrame()).copy()
    bottlenecks = stage.sort_values("total_elapsed_hours", ascending=False).head(5) if "total_elapsed_hours" in stage else stage.head(5)
    if not bottlenecks.empty: bottlenecks["Interpretation"] = "Bottleneck candidate; validate cause with the department."
    status_counts = tasks["Current Status"].fillna("Unavailable").value_counts().rename_axis("Status").reset_index(name="Tasks")
    delivery = pd.DataFrame({"Due Variance Category": ["Completed on time", "Completed late", "Open overdue"],
                             "Tasks": [on_time, int((due_known & tasks["On Time?"].eq(False)).sum()), overdue]})
    quality = process_data.get("data_quality", pd.DataFrame()).copy()
    if not quality.empty:
        quality = quality.rename(
            columns={"issue_key": "Issue Type", "finding": "Analysis Impact"}
        )
        quality["Count"] = 1
        quality = quality[["Issue Type", "Count", "Analysis Impact"]]
    else:
        quality = pd.DataFrame(
            columns=["Issue Type", "Count", "Analysis Impact"]
        )
    return {
        "tasks": tasks, "kpis": kpis, "employee_breakdown": employees, "attention": attention,
        "bottlenecks": bottlenecks, "department_quality": quality, "status_counts": status_counts,
        "due_variance_summary": delivery, "weekly_flow": process_data.get("weekly_flow", pd.DataFrame()).rename(
            columns={"week_start":"Week Starting", "tasks_opened":"Tasks Created", "tasks_completed":"Tasks Completed"}),
        "department_name": "Tech Development", "department_id": "jira-tech-development",
        "space_name": space_name, "cutoff": cutoff, "data_source": "Jira",
    }


def _write_frame(ws, frame: pd.DataFrame, start_row=1, title=None):
    row = start_row
    if title:
        ws.cell(row, 1, title).font = Font(bold=True, size=14, color="17324D")
        row += 2
    for column, value in enumerate(frame.columns, 1):
        cell = ws.cell(row, column, value)
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(color="FFFFFF", bold=True)
    for values in frame.itertuples(index=False, name=None):
        row += 1
        for column, value in enumerate(values, 1):
            if isinstance(value, pd.Timestamp): value = value.isoformat()
            if pd.isna(value) if not isinstance(value, (list, dict)) else False: value = None
            ws.cell(row, column, value)
    return row


def _department_task_frame(result: dict) -> pd.DataFrame:
    tasks = result["tasks"].copy()
    if "Space" not in tasks.columns:
        tasks.insert(0, "Space", result.get("space_name", "Unknown"))
    return tasks


def _department_status_column(tasks: pd.DataFrame) -> str:
    return "Current Status" if "Current Status" in tasks.columns else "Status"


def _department_status_rows(tasks: pd.DataFrame) -> pd.DataFrame:
    column = _department_status_column(tasks)
    if column not in tasks:
        return pd.DataFrame(columns=["Status", "Tasks"])
    return tasks[column].fillna("Unknown").value_counts().rename_axis("Status").reset_index(name="Tasks")


def _cutoff_timestamp(value):
    """Return a timezone-consistent cutoff for Department date comparisons."""
    return pd.to_datetime(value, utc=True, errors="coerce")


def _department_due_rows(tasks: pd.DataFrame, cutoff) -> pd.DataFrame:
    cutoff_timestamp = _cutoff_timestamp(cutoff)
    if tasks.empty or "Open?" not in tasks:
        return pd.DataFrame([("Open overdue", 0), ("Open within due date", 0), ("Open without due date", 0)],
                            columns=["Due Status", "Tasks"])
    open_mask = tasks["Open?"].eq(True)
    due = tasks["Due Date"] if "Due Date" in tasks else pd.Series(pd.NaT, index=tasks.index)
    due_timestamps = pd.to_datetime(due, utc=True, errors="coerce")
    overdue_days = pd.to_numeric(tasks.get("Overdue Days", pd.Series(index=tasks.index)), errors="coerce")
    overdue_mask = open_mask & (
        overdue_days.gt(0)
        | (due_timestamps.notna() & due_timestamps.lt(cutoff_timestamp))
    )
    return pd.DataFrame([
        ("Open overdue", int(overdue_mask.sum())),
        ("Open within due date", int((open_mask & due_timestamps.notna() & ~overdue_mask).sum())),
        ("Open without due date", int((open_mask & due_timestamps.isna()).sum())),
    ], columns=["Due Status", "Tasks"])


def _department_space_breakdown(tasks: pd.DataFrame, cutoff=None) -> pd.DataFrame:
    if tasks.empty:
        return pd.DataFrame(columns=["Space", "Total Tasks", "Completed", "Open", "Open Overdue"])
    cutoff_timestamp = _cutoff_timestamp(cutoff)
    completed = tasks.get("Completed?", pd.Series(False, index=tasks.index)).eq(True)
    open_mask = tasks.get("Open?", pd.Series(False, index=tasks.index)).eq(True)
    due = tasks.get("Due Date", pd.Series(pd.NaT, index=tasks.index))
    due_timestamps = pd.to_datetime(due, utc=True, errors="coerce")
    overdue_days = pd.to_numeric(tasks.get("Overdue Days", pd.Series(index=tasks.index)), errors="coerce")
    overdue = open_mask & (
        overdue_days.gt(0)
        | (due_timestamps.notna() & due_timestamps.lt(cutoff_timestamp))
    )
    grouped = tasks.assign(
        _completed=completed, _open=open_mask, _overdue=overdue,
    ).groupby("Space", dropna=False).agg(
        **{"Total Tasks": ("Task ID", "count"), "Completed": ("_completed", "sum"),
           "Open": ("_open", "sum"), "Open Overdue": ("_overdue", "sum")}
    ).reset_index()
    grouped["Space"] = grouped["Space"].fillna("Unknown")
    return grouped


def _department_weekly_frame(result: dict) -> pd.DataFrame:
    weekly = result.get("weekly_flow", pd.DataFrame()).copy()
    if weekly.empty:
        return pd.DataFrame([{"Week Starting": result.get("cutoff"), "Tasks Created": 0, "Tasks Completed": 0}],
                            columns=["Week Starting", "Tasks Created", "Tasks Completed"])
    weekly = weekly.rename(columns={
        "week_start": "Week Starting", "tasks_opened": "Tasks Created",
        "tasks_completed": "Tasks Completed",
    })
    for name in ("Week Starting", "Tasks Created", "Tasks Completed"):
        if name not in weekly:
            weekly[name] = 0
    return weekly[["Week Starting", "Tasks Created", "Tasks Completed"]]


def _department_card(ws, index: int, title: str, value: Any) -> None:
    row = 3 + (index // 5) * 4
    start_col = 1 + (index % 5) * 2
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=start_col + 1)
    ws.merge_cells(start_row=row + 1, start_column=start_col, end_row=row + 2, end_column=start_col + 1)
    header = ws.cell(row, start_col, title)
    header.fill = PatternFill("solid", fgColor="17324D")
    header.font = Font(color="FFFFFF", bold=True)
    header.alignment = Alignment(horizontal="center")
    value_cell = ws.cell(row + 1, start_col, "N/A" if value is None else value)
    value_cell.font = Font(size=15, bold=True, color="1F2937")
    value_cell.alignment = Alignment(horizontal="center", vertical="center")


def _department_dashboard_chart(ws, chart_type, title, anchor, start_row, headers, rows):
    _write_frame(ws, pd.DataFrame(rows, columns=headers), start_row=start_row)
    if chart_type == "line":
        chart = LineChart()
        chart.add_data(Reference(ws, min_col=2, max_col=len(headers), min_row=start_row, max_row=start_row + len(rows)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=start_row + 1, max_row=start_row + len(rows)))
    else:
        chart = BarChart()
        chart.type = "col"
        chart.add_data(Reference(ws, min_col=2, max_col=len(headers), min_row=start_row, max_row=start_row + len(rows)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=start_row + 1, max_row=start_row + len(rows)))
    chart.title = title
    chart.style = 10
    chart.height = 7
    chart.width = 13
    if len(headers) > 2:
        chart.legend.position = "b"
    ws.add_chart(chart, anchor)


def department_excel(result: dict) -> bytes:
    """Return the exact Department Performance workbook requested by the user."""
    tasks = _department_task_frame(result)
    kpi_values = result["kpis"].set_index("KPI")["Value"].to_dict()
    total = len(tasks)
    wip = int(tasks.get("WIP?", pd.Series(False, index=tasks.index)).eq(True).sum())
    source_spaces = ", ".join(map(str, result.get("space_names", [result.get("space_name", "Unknown")])))
    cutoff = result.get("cutoff")
    workbook = Workbook()
    workbook.remove(workbook.active)

    dashboard = workbook.create_sheet("Department_Executive_Dashboard")
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:J1")
    dashboard["A1"] = "Department Performance"
    dashboard["A1"].fill = PatternFill("solid", fgColor="17324D")
    dashboard["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    dashboard["A1"].alignment = Alignment(horizontal="center")
    dashboard.merge_cells("A2:J2")
    dashboard["A2"] = f"Department: {result['department_name']} | Source: {result.get('data_source', 'N/A')} | Cutoff: {cutoff}"
    dashboard["A2"].font = Font(color="6B7280", italic=True)
    dashboard["A2"].alignment = Alignment(horizontal="center")
    for column in "ABCDEFGHIJ":
        dashboard.column_dimensions[column].width = 16

    cards = [
        ("Total Tasks", total),
        ("Completion Rate", None if pd.isna(kpi_values.get("Task Completion Rate (%)")) else f"{float(kpi_values['Task Completion Rate (%)']):.1f}%"),
        ("On-Time Rate", None if pd.isna(kpi_values.get("On-Time Completion Rate (%)")) else f"{float(kpi_values['On-Time Completion Rate (%)']):.1f}%"),
        ("Open Overdue", kpi_values.get("Open Overdue Tasks")),
        ("WIP", wip),
        ("Average Lead Time", None if pd.isna(kpi_values.get("Average Lead Time (hours)")) else f"{float(kpi_values['Average Lead Time (hours)']):.1f} h"),
        ("Workflow Exception Rate", None if pd.isna(kpi_values.get("Workflow Exception Rate (%)")) else f"{float(kpi_values['Workflow Exception Rate (%)']):.1f}%"),
    ]
    for index, (title, value) in enumerate(cards):
        _department_card(dashboard, index, title, value)

    weekly = _department_weekly_frame(result)
    statuses = _department_status_rows(tasks)
    due = _department_due_rows(tasks, cutoff)
    employees = result.get("employee_breakdown", pd.DataFrame()).copy()
    employee_rows = [
        [row.get("Assignee", "Unknown"), row.get("Total Assigned", 0),
         row.get("Completed", 0), row.get("Open", 0)]
        for row in employees.to_dict("records")
    ] or [["No data", 0, 0, 0]]
    spaces = _department_space_breakdown(tasks, cutoff)
    space_rows = spaces[["Space", "Total Tasks"]].values.tolist() if not spaces.empty else [["No data", 0]]
    _department_dashboard_chart(dashboard, "line", "Weekly Task Flow", "A11", 45,
                                ["Week Starting", "Tasks Created", "Tasks Completed"], weekly.values.tolist())
    _department_dashboard_chart(dashboard, "bar", "Task Distribution by Status", "I11", 45 + len(weekly) + 3,
                                ["Status", "Tasks"], statuses.values.tolist() or [["No data", 0]])
    _department_dashboard_chart(dashboard, "bar", "Open Tasks by Due Status", "A27", 45 + len(weekly) + len(statuses) + 6,
                                ["Due Status", "Tasks"], due.values.tolist())
    _department_dashboard_chart(dashboard, "bar", "Work Distribution by Employee", "I27", 45 + len(weekly) + len(statuses) + len(due) + 9,
                                ["Employee", "Total Assigned", "Completed", "Open"], employee_rows)
    _department_dashboard_chart(dashboard, "bar", "Work Distribution by Space", "A43", 45 + len(weekly) + len(statuses) + len(due) + len(employee_rows) + 12,
                                ["Space", "Tasks"], space_rows)
    _fit_department_sheets = [dashboard]

    summary = workbook.create_sheet("Department Summary")
    info = pd.DataFrame([
        ("Department", result["department_name"]),
        ("Data Source", result.get("data_source", "N/A")),
        ("Spaces", source_spaces),
        ("Analysis Cutoff", cutoff),
        ("Total Tasks", total),
    ], columns=["Field", "Value"])
    end = _write_frame(summary, info, title="Department Summary")
    _write_frame(summary, result["kpis"], start_row=end + 2, title="Department KPI Summary")

    _write_frame(workbook.create_sheet("Employee Breakdown"), employees)
    _write_frame(workbook.create_sheet("Space Breakdown"), spaces)
    _write_frame(workbook.create_sheet("Task Details"), tasks)
    _write_frame(workbook.create_sheet("Status Summary"), statuses)
    _write_frame(workbook.create_sheet("Weekly Flow"), weekly)

    overdue_mask = tasks.get("Open?", pd.Series(False, index=tasks.index)).eq(True) & pd.to_numeric(
        tasks.get("Overdue Days", pd.Series(index=tasks.index)), errors="coerce"
    ).gt(0)
    _write_frame(workbook.create_sheet("Overdue Tasks"), tasks.loc[overdue_mask])
    _write_frame(workbook.create_sheet("Tasks Requiring Attention"), result.get("attention", pd.DataFrame()))
    _write_frame(workbook.create_sheet("Bottlenecks"), result.get("bottlenecks", pd.DataFrame()))
    exceptions = result.get("attention", pd.DataFrame()).copy()
    if "Issue" in exceptions:
        exceptions = exceptions[exceptions["Issue"].fillna("").astype(str).str.contains("Workflow", case=False)]
    _write_frame(workbook.create_sheet("Workflow Exceptions"), exceptions)
    _write_frame(workbook.create_sheet("Data Quality"), result.get("department_quality", pd.DataFrame()))

    _write_frame(workbook.create_sheet("Analysis Context"), pd.DataFrame([
        ("Analysis Level", "Department"),
        ("Department", result["department_name"]),
        ("Data Source", result.get("data_source", "N/A")),
        ("Spaces", source_spaces),
        ("Analysis Cutoff", cutoff),
        ("Task Scope", "All collected tasks in the department Spaces."),
        ("Subtask Treatment", "Visible in Task Details and total count; excluded from completion-rate numerator and denominator."),
    ], columns=["Field", "Value"]))
    _write_frame(workbook.create_sheet("Metric Definitions"), pd.DataFrame([
        ("Total Tasks", "All tasks collected from the department Spaces."),
        ("Completion Rate", "Completed parent/standalone tasks divided by parent/standalone tasks."),
        ("On-Time Rate", "Completed tasks on or before due date divided by completed tasks with a known due date."),
        ("Open Overdue", "Open tasks with a due date before the analysis cutoff."),
        ("WIP", "Open tasks currently in the department's in-progress workflow statuses."),
        ("Average Lead Time", "Average completion date minus creation date for completed tasks."),
        ("Workflow Exception Rate", "Tasks with recorded workflow exceptions divided by tasks with valid workflow coverage."),
    ], columns=["Metric", "Definition"]))

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 45)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(12, width)
    return _save_workbook(workbook)


def _save_workbook(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _display(value: Any) -> str:
    """Render report values safely for Word tables and narrative text."""
    if value is None:
        return "N/A"
    if isinstance(value, (list, tuple, dict)):
        return str(value)
    try:
        if pd.isna(value):
            return "N/A"
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    return str(value)


def _bottleneck_word_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Normalize ClickUp and Jira bottleneck columns for the Word report."""
    columns = [
        "Status", "Tasks", "Average Hours", "Total Hours",
        "Open Tasks", "Interpretation",
    ]
    if frame is None or frame.empty:
        return pd.DataFrame(columns=columns)

    aliases = {
        "Status": ("Status", "status"),
        "Tasks": ("Tasks", "tasks", "tasks_visited", "task_count"),
        "Average Hours": (
            "Average Hours", "average_hours", "elapsed_mean_hours",
        ),
        "Total Hours": (
            "Total Hours", "total_hours", "elapsed_total_hours",
        ),
        "Open Tasks": (
            "Open Tasks", "open_tasks", "open_tasks_currently_here",
        ),
        "Interpretation": ("Interpretation", "interpretation"),
    }
    output = pd.DataFrame(index=frame.index)
    for target, candidates in aliases.items():
        for candidate in candidates:
            if candidate in frame.columns:
                output[target] = frame[candidate]
                break
        if target not in output:
            output[target] = None
    return output[columns]


def _word_table(document: Document, frame: pd.DataFrame, columns: list[str]) -> None:
    visible = frame[[column for column in columns if column in frame.columns]].copy()
    if visible.empty or visible.shape[1] == 0:
        document.add_paragraph("No items were identified in this section.")
        return
    table = document.add_table(rows=1, cols=len(visible.columns))
    table.style = "Table Grid"
    for index, name in enumerate(visible.columns):
        table.rows[0].cells[index].text = name
    for values in visible.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = _display(value)


def department_word(result: dict) -> bytes:
    """Create the exact Department Performance Word report requested by the user."""
    tasks = _department_task_frame(result)
    kpis = result["kpis"].set_index("KPI")["Value"].to_dict()
    source_spaces = ", ".join(map(str, result.get("space_names", [result.get("space_name", "Unknown")])))
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    document.add_heading("Department Performance Report", 0)
    document.add_paragraph(
        f"Department: {result['department_name']} | Source: {result.get('data_source', 'N/A')} | "
        f"Spaces: {source_spaces} | Analysis cutoff: {result.get('cutoff', 'N/A')}"
    )

    document.add_heading("Executive Summary", level=1)
    completion = kpis.get("Task Completion Rate (%)")
    completion_text = "N/A" if completion is None or pd.isna(completion) else f"{float(completion):.1f}%"
    overdue = kpis.get("Open Overdue Tasks", 0)
    document.add_paragraph(
        f"The department contains {len(tasks)} task(s) across the listed Spaces. "
        f"Completion rate is {completion_text}; open overdue work is {_display(overdue)} task(s)."
    )

    document.add_heading("Department KPI Summary", level=1)
    _word_table(document, result["kpis"], ["KPI", "Value", "Numerator", "Denominator"])

    employees = result.get("employee_breakdown", pd.DataFrame()).copy()
    document.add_heading("Employee Work Distribution and Workload Summary", level=1)
    _word_table(document, employees, ["Assignee", "Total Assigned", "Completed", "Open", "WIP", "Overdue", "Workflow Exceptions"])

    document.add_heading("Space / Project Comparison", level=1)
    _word_table(document, _department_space_breakdown(tasks, result.get("cutoff")), ["Space", "Total Tasks", "Completed", "Open", "Open Overdue"])

    document.add_heading("Department Bottlenecks", level=1)
    bottleneck_frame = _bottleneck_word_frame(result.get("bottlenecks", pd.DataFrame()))
    _word_table(document, bottleneck_frame, [
        "Status", "Tasks", "Average Hours", "Total Hours", "Open Tasks", "Interpretation",
    ])

    document.add_heading("Tasks Requiring Attention", level=1)
    _word_table(document, result.get("attention", pd.DataFrame()), [
        "Task ID", "Task Name", "Assignee", "Current Status", "Priority", "Due Date", "Issue",
    ])

    overdue_mask = tasks.get("Open?", pd.Series(False, index=tasks.index)).eq(True) & pd.to_numeric(
        tasks.get("Overdue Days", pd.Series(index=tasks.index)), errors="coerce"
    ).gt(0)
    late_mask = tasks.get("Completed?", pd.Series(False, index=tasks.index)).eq(True)
    if "Due Variance (days)" in tasks:
        late_mask &= pd.to_numeric(tasks["Due Variance (days)"], errors="coerce").gt(0)
    elif "On Time?" in tasks:
        late_mask &= tasks["On Time?"].eq(False)
    document.add_heading("Overdue and Late Tasks", level=1)
    document.add_paragraph("Open Overdue Tasks")
    _word_table(document, tasks.loc[overdue_mask], ["Task ID", "Task Name", "Space", "Assignee", "Priority", "Due Date", "Current Status"])
    document.add_paragraph("Late Completed Tasks")
    _word_table(document, tasks.loc[late_mask], ["Task ID", "Task Name", "Space", "Assignee", "Priority", "Due Date", "Completed"])

    document.add_heading("Recommendations", level=1)
    if overdue_mask.any():
        document.add_paragraph("Prioritize open overdue tasks and confirm an owner and recovery date.", style="List Bullet")
    if not employees.empty:
        document.add_paragraph("Review employee workload distribution together with task priority and due dates.", style="List Bullet")
    if result.get("bottlenecks") is not None and not result["bottlenecks"].empty:
        document.add_paragraph("Validate bottleneck candidates with the responsible process owners.", style="List Bullet")
    if not overdue_mask.any() and employees.empty and (result.get("bottlenecks") is None or result["bottlenecks"].empty):
        document.add_paragraph("No additional recommendations were generated from the available evidence.", style="List Bullet")

    document.add_heading("Data Quality and Limitations", level=1)
    _word_table(document, result.get("department_quality", pd.DataFrame()), ["Issue Type", "Count", "Analysis Impact"])
    document.add_paragraph(
        "The Department report aggregates all selected Spaces. Full task detail remains in Excel; "
        "Word lists only tasks requiring operational follow-up, overdue tasks, and late completed tasks."
    )

    document.add_heading("Metric Definitions", level=1)
    definitions = pd.DataFrame([
        ("Total Tasks", "All tasks collected from the department Spaces."),
        ("Completion Rate", "Completed parent/standalone tasks divided by parent/standalone tasks."),
        ("On-Time Rate", "Completed tasks on or before due date divided by completed tasks with a known due date."),
        ("Open Overdue", "Open tasks with a due date before the analysis cutoff."),
        ("WIP", "Open tasks currently in the department's in-progress workflow statuses."),
        ("Average Lead Time", "Average completion date minus creation date for completed tasks."),
        ("Workflow Exception Rate", "Tasks with recorded workflow exceptions divided by tasks with valid workflow coverage."),
    ], columns=["Metric", "Definition"])
    _word_table(document, definitions, ["Metric", "Definition"])

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
