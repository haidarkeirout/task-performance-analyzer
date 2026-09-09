"""Jira-only Excel Executive Dashboard builder.

This module consumes only the Jira task-metrics frame and Jira process tables.
It intentionally has no ClickUp imports or ClickUp field assumptions.
"""
from __future__ import annotations

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DARK = "17324D"
MID = "2F5D7E"
LIGHT = "EAF1F6"
WHITE = "FFFFFF"
TEXT = "1F2937"
BORDER = Side(style="thin", color="D7E1E8")
PANEL_RANGES = {
    "A13": "A13:H26",
    "I13": "I13:P26",
    "A29": "A29:H42",
    "I29": "I29:P42",
}


def _safe_number(value, default=None):
    try:
        if value is None or pd.isna(value):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _mean(frame: pd.DataFrame, column: str, mask=None):
    if frame is None or frame.empty or column not in frame.columns:
        return None
    values = frame.loc[mask, column] if mask is not None else frame[column]
    values = pd.to_numeric(values, errors="coerce").dropna()
    return None if values.empty else float(values.mean())


def _format_percent(numerator, denominator):
    numerator = _safe_number(numerator, 0)
    denominator = _safe_number(denominator, 0)
    if not denominator:
        return "Unavailable"
    return f"{numerator / denominator * 100:.1f}%"


def _format_average(value, suffix):
    value = _safe_number(value)
    return "Unavailable" if value is None else f"{value:.1f} {suffix}"


def _card(sheet, start_col: int, label: str, value, row: int):
    end_col = start_col + 1
    sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    sheet.merge_cells(start_row=row + 1, start_column=start_col, end_row=row + 2, end_column=end_col)
    label_cell = sheet.cell(row, start_col, label)
    value_cell = sheet.cell(row + 1, start_col, value)
    label_cell.fill = PatternFill("solid", fgColor=DARK)
    label_cell.font = Font(color=WHITE, bold=True, size=9)
    value_cell.fill = PatternFill("solid", fgColor=LIGHT)
    value_cell.font = Font(color=TEXT, bold=True, size=15)
    for r in range(row, row + 3):
        for c in range(start_col, end_col + 1):
            cell = sheet.cell(r, c)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)
    sheet.row_dimensions[row].height = 21
    sheet.row_dimensions[row + 1].height = 25
    sheet.row_dimensions[row + 2].height = 14


def _write_support_table(sheet, start_row: int, start_col: int, headers, rows):
    for offset, value in enumerate(headers):
        cell = sheet.cell(start_row, start_col + offset, value)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(color=WHITE, bold=True)
    for row_offset, row in enumerate(rows, start=1):
        for col_offset, value in enumerate(row):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime().replace(tzinfo=None) if value.tzinfo else value.to_pydatetime()
            sheet.cell(start_row + row_offset, start_col + col_offset, value)
    return start_row + len(rows)


def _empty_panel(sheet, anchor: str, title: str):
    panel = PANEL_RANGES.get(anchor)
    if not panel:
        return
    sheet.merge_cells(panel)
    cell = sheet[anchor]
    cell.value = f"{title}\n\nNo data available for this chart."
    cell.fill = PatternFill("solid", fgColor="F8FAFC")
    cell.font = Font(color=MID, bold=True, size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(left=BORDER, right=BORDER, top=BORDER, bottom=BORDER)


def _configure_chart(chart, title: str, y_title: str):
    chart.title = title
    chart.height = 6.5
    chart.width = 12.0
    chart.legend.position = "b"
    chart.y_axis.title = y_title
    chart.x_axis.title = ""
    chart.display_blanks = "zero"


def _add_line_chart(sheet, title, data_start_row, data_end_row, category_col, value_cols, anchor):
    if data_end_row <= data_start_row:
        _empty_panel(sheet, anchor, title)
        return
    chart = LineChart()
    chart.style = 13
    _configure_chart(chart, title, "Tasks")
    data = Reference(sheet, min_col=value_cols[0], max_col=value_cols[-1], min_row=data_start_row, max_row=data_end_row)
    categories = Reference(sheet, min_col=category_col, min_row=data_start_row + 1, max_row=data_end_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, anchor)


def _add_bar_chart(sheet, title, data_start_row, data_end_row, category_col, value_cols, anchor, y_title="Tasks"):
    if data_end_row <= data_start_row:
        _empty_panel(sheet, anchor, title)
        return
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    _configure_chart(chart, title, y_title)
    data = Reference(sheet, min_col=value_cols[0], max_col=value_cols[-1], min_row=data_start_row, max_row=data_end_row)
    categories = Reference(sheet, min_col=category_col, min_row=data_start_row + 1, max_row=data_end_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, anchor)


def add_jira_executive_dashboard(workbook, frame: pd.DataFrame, tables: dict) -> None:
    """Create a Jira-only Executive_Dashboard worksheet with KPI cards and charts."""
    if "Executive_Dashboard" in workbook.sheetnames:
        del workbook["Executive_Dashboard"]
    sheet = workbook.create_sheet("Executive_Dashboard", 0)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A12"
    sheet.sheet_properties.tabColor = DARK

    for col in range(1, 17):
        sheet.column_dimensions[get_column_letter(col)].width = 11

    sheet.merge_cells("A1:P1")
    sheet["A1"] = "Jira Executive Dashboard"
    sheet["A1"].fill = PatternFill("solid", fgColor=DARK)
    sheet["A1"].font = Font(color=WHITE, bold=True, size=20)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 32

    cutoff = frame.iloc[0].get("evaluation_cutoff") if frame is not None and not frame.empty else "Unavailable"
    process_name = "Jira"
    if tables:
        context = tables.get("process_context")
        if context is not None and not context.empty and {"Field", "Value"}.issubset(context.columns):
            matches = context.loc[context["Field"].eq("Process Name"), "Value"]
            if not matches.empty:
                process_name = str(matches.iloc[0])
    sheet.merge_cells("A2:P2")
    sheet["A2"] = f"Source: Jira | Space: {process_name} | Evaluation cutoff: {cutoff}"
    sheet["A2"].font = Font(color=MID, italic=True, size=10)
    sheet["A2"].alignment = Alignment(horizontal="center")

    summary = tables.get("overall_summary") if tables else None
    summary_row = summary.iloc[0] if summary is not None and not summary.empty else {}
    total = summary_row.get("total_tasks", len(frame)) if hasattr(summary_row, "get") else len(frame)
    completed = summary_row.get("completed_tasks", 0) if hasattr(summary_row, "get") else 0
    on_time = summary_row.get("on_time_tasks", 0) if hasattr(summary_row, "get") else 0
    on_time_valid = summary_row.get("on_time_valid_tasks", 0) if hasattr(summary_row, "get") else 0
    overdue = summary_row.get("overdue_open_tasks", 0) if hasattr(summary_row, "get") else 0
    wip = summary_row.get("wip_tasks", 0) if hasattr(summary_row, "get") else 0

    kpis = [
        ("Total Tasks", int(_safe_number(total, 0))),
        ("Completed", int(_safe_number(completed, 0))),
        ("Completion Rate", _format_percent(completed, total)),
        ("On-Time Rate", _format_percent(on_time, on_time_valid)),
        ("Open Overdue", int(_safe_number(overdue, 0))),
        ("WIP Tasks", int(_safe_number(wip, 0))),
    ]
    for index, (label, value) in enumerate(kpis):
        _card(sheet, 1 + index * 2, label, value, 3)

    completed_mask = frame["is_completed"].eq(True) if "is_completed" in frame else pd.Series(False, index=frame.index)
    completed_late_mask = completed_mask & frame["schedule_variance_days"].gt(0) if "schedule_variance_days" in frame else completed_mask & False
    open_overdue_mask = (frame["is_open"].eq(True) & frame["overdue_days"].gt(0)) if {"is_open", "overdue_days"}.issubset(frame.columns) else pd.Series(False, index=frame.index)
    started_mask = frame["time_to_start_business_hours"].notna() if "time_to_start_business_hours" in frame else pd.Series(False, index=frame.index)
    start_variance_mask = frame["start_schedule_variance_days"].notna() if "start_schedule_variance_days" in frame else pd.Series(False, index=frame.index)

    averages = [
        ("Avg Execution Time", _format_average(_mean(frame, "execution_business_hours", completed_mask), "h")),
        ("Avg Lead Time", _format_average(_mean(frame, "lead_time_business_hours", completed_mask), "h")),
        ("Avg Time to Start", _format_average(_mean(frame, "time_to_start_business_hours", started_mask), "h")),
        ("Avg Late Completion", _format_average(_mean(frame, "schedule_variance_days", completed_late_mask), "days")),
        ("Avg Open Overdue", _format_average(_mean(frame, "overdue_days", open_overdue_mask), "days")),
        ("Avg Start Variance", _format_average(_mean(frame, "start_schedule_variance_days", start_variance_mask), "days")),
    ]
    for index, (label, value) in enumerate(averages):
        _card(sheet, 1 + index * 2, label, value, 7)

    completed_late = int(completed_late_mask.sum())
    status_unavailable = int(frame["status_known"].eq(False).sum()) if "status_known" in frame else 0
    sheet["A11"] = f"Completed late: {completed_late}"
    sheet["D11"] = f"Status unavailable: {status_unavailable}"
    for coordinate in ("A11", "D11"):
        sheet[coordinate].font = Font(bold=True, color=TEXT)

    # Keep chart source data on ordinary visible cells below the dashboard.
    # Hidden source columns can render as blank charts in some Excel clients.
    support_col = 1
    support_row = 70

    weekly = tables.get("weekly_flow", pd.DataFrame()) if tables else pd.DataFrame()
    weekly_rows = [] if weekly is None or weekly.empty else weekly[["week_start", "tasks_opened", "tasks_completed"]].values.tolist()
    weekly_end = _write_support_table(sheet, support_row, support_col, ["Week Starting", "Tasks Opened", "Tasks Completed"], weekly_rows)
    _add_line_chart(sheet, "Weekly Task Flow", support_row, weekly_end, support_col, (support_col + 1, support_col + 2), "A13")
    support_row = weekly_end + 3

    status_counts = (
        frame["status_at_cutoff"].fillna("Unavailable").value_counts().rename_axis("Status").reset_index(name="Tasks")
        if "status_at_cutoff" in frame else pd.DataFrame(columns=["Status", "Tasks"])
    )
    status_end = _write_support_table(sheet, support_row, support_col, ["Status", "Tasks"], status_counts.values.tolist())
    _add_bar_chart(sheet, "Task Distribution by Status", support_row, status_end, support_col, (support_col + 1, support_col + 1), "I13")
    support_row = status_end + 3

    due = tables.get("deadline_summary", pd.DataFrame()) if tables else pd.DataFrame()
    due_rows = [] if due is None or due.empty else due[["due_status", "task_count"]].values.tolist()
    due_end = _write_support_table(sheet, support_row, support_col, ["Due Status", "Tasks"], due_rows)
    _add_bar_chart(sheet, "Open Tasks by Due Status", support_row, due_end, support_col, (support_col + 1, support_col + 1), "A29")
    support_row = due_end + 3

    if "assignee_name" in frame:
        assignee = (
            frame.assign(assignee_name=frame["assignee_name"].fillna("Unassigned"))
            .groupby("assignee_name", dropna=False)
            .agg(Completed=("is_completed", "sum"), Open=("is_open", "sum"), Rejected=("is_rejected", "sum"))
            .reset_index()
            .rename(columns={"assignee_name": "Assignee"})
        )
    else:
        assignee = pd.DataFrame(columns=["Assignee", "Completed", "Open", "Rejected"])
    assignee_end = _write_support_table(sheet, support_row, support_col, ["Assignee", "Completed", "Open", "Rejected"], assignee.values.tolist())
    _add_bar_chart(sheet, "Work Distribution by Assignee", support_row, assignee_end, support_col, (support_col + 1, support_col + 3), "I29")

    sheet["A67"] = "Dashboard chart source data (not included in print area)"
    sheet["A67"].font = Font(color=MID, italic=True, size=9)
    sheet.print_area = "A1:P44"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
