from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
HEADER_FILL = "F2F4F7"
TABLE_INDENT_DXA = 120


def is_true(value: Any) -> bool:
    return value is True or str(value).strip().lower() == "true"


def format_value(value: Any) -> str:
    if value is None:
        return "Unavailable"
    if isinstance(value, float):
        if pd.isna(value):
            return "Unavailable"
        return f"{value:.2f}"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d %H:%M UTC")
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null"}:
        return "Unavailable"
    return text


def format_percent(numerator: int, denominator: int) -> str:
    if denominator == 0:
        return "Unavailable"
    return f"{numerator / denominator * 100:.1f}%"


def parse_labels(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item) for item in value]
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null", "[]"}:
        return []
    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return [str(item) for item in parsed]
    except json.JSONDecodeError:
        pass
    return [text]


def _set_run_font(run, *, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_paragraph_spacing(paragraph, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def _set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    grid_columns = list(grid.iterchildren())
    for index, width in enumerate(widths):
        if index < len(grid_columns):
            grid_columns[index].set(qn("w:w"), str(width))
    for row in table.rows:
        keep = row._tr.get_or_add_trPr()
        if keep.find(qn("w:cantSplit")) is None:
            keep.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[index])
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_document_style(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in tokens.items():
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_spacing(header, after=0)
    run = header.add_run("Jira Performance Evaluation")
    _set_run_font(run, size=8, color="6B7280")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_paragraph_spacing(footer, after=0)
    footer_run = footer.add_run("Generated by Task Performance Intelligence")
    _set_run_font(footer_run, size=8, color="6B7280")


def add_title(document: Document, title: str, subtitle: str = "") -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_paragraph_spacing(paragraph, before=0, after=3, line=1.1)
    run = paragraph.add_run(title)
    _set_run_font(run, size=20, bold=True, color=INK)
    if subtitle:
        sub = document.add_paragraph()
        _set_paragraph_spacing(sub, before=0, after=16, line=1.1)
        sub_run = sub.add_run(subtitle)
        _set_run_font(sub_run, size=10, color="5B6470")


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        _set_run_font(
            run,
            size={1: 16, 2: 13, 3: 12}.get(level, 11),
            bold=True,
            color={1: BLUE, 2: BLUE, 3: DARK_BLUE}.get(level, INK),
        )


def add_paragraph(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = paragraph.add_run(bold_prefix)
        _set_run_font(prefix_run, size=11, bold=True)
        run = paragraph.add_run(text[len(bold_prefix):])
        _set_run_font(run, size=11)
    else:
        run = paragraph.add_run(text)
        _set_run_font(run, size=11)


def _add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    _set_paragraph_spacing(paragraph, after=4)
    for run in paragraph.runs:
        _set_run_font(run, size=11)
    if not paragraph.runs:
        run = paragraph.add_run(text)
        _set_run_font(run, size=11)
    else:
        paragraph.runs[0].text = text


def add_key_value_table(document: Document, values: dict[str, Any]) -> None:
    table = document.add_table(rows=1, cols=2)
    _set_repeat_table_header(table.rows[0])
    headers = ["Metric", "Value"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_shading(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        _set_paragraph_spacing(paragraph, after=0, line=1.0)
        run = paragraph.add_run(header)
        _set_run_font(run, size=9, bold=True, color=INK)
    for key, value in values.items():
        cells = table.add_row().cells
        for index, content in enumerate((key, format_value(value))):
            paragraph = cells[index].paragraphs[0]
            _set_paragraph_spacing(paragraph, after=0, line=1.0)
            run = paragraph.add_run(str(content))
            _set_run_font(run, size=9, bold=(index == 0))
    _set_table_geometry(table, [2700, 6660])
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def _jira_recommendations(row, task_metrics: pd.DataFrame) -> list[str]:
    recommendations: list[str] = []
    total = int(row.get("total_tasks", 0) or 0)
    completed = int(row.get("completed_tasks", 0) or 0)
    overdue = int(row.get("overdue_open_tasks", 0) or 0)
    wip = int(row.get("wip_tasks", 0) or 0)
    on_time_valid = int(row.get("on_time_valid_tasks", 0) or 0)
    on_time = int(row.get("on_time_tasks", 0) or 0)
    review_exceptions = int(row.get("review_exception_tasks", 0) or 0)
    excluded_histories = int(row.get("history_excluded_tasks", 0) or 0)

    completion_rate = completed / total if total else None
    on_time_rate = on_time / on_time_valid if on_time_valid else None
    wip_share = wip / total if total else 0

    if overdue > 0:
        recommendations.append(
            f"Prioritize the {overdue} open overdue task(s): confirm blockers, owners, and realistic due dates, then review them in the next operating cadence."
        )
    if on_time_rate is not None and on_time_rate < 0.80:
        recommendations.append(
            f"Improve schedule reliability: the on-time completion rate is {on_time_rate * 100:.1f}%. Review estimation, due-date setting, and early escalation for tasks at risk."
        )
    if completion_rate is not None and completion_rate < 0.70:
        recommendations.append(
            f"Review backlog conversion: {completed} of {total} tasks are completed. Prioritize the highest-value open work and remove or re-scope stale items."
        )
    if wip_share >= 0.30 and wip > 0:
        recommendations.append(
            f"Control work in progress: {wip} task(s) are currently WIP ({wip_share * 100:.1f}% of scope). Consider WIP limits and finishing active work before starting additional items."
        )
    if review_exceptions > 0:
        recommendations.append(
            f"Inspect the {review_exceptions} task(s) with rework, replanning, or re-evaluation evidence and address recurring causes at workflow or requirements level."
        )
    if excluded_histories > 0:
        recommendations.append(
            f"Improve Jira history coverage for the {excluded_histories} task(s) excluded from history-dependent metrics so future cycle-time and exception analysis is more complete."
        )
    if not recommendations:
        recommendations.append(
            "Maintain the current operating controls and continue monitoring completion, timeliness, WIP, and review exceptions at each evaluation cutoff."
        )
    return recommendations


def _process_name(tables: dict) -> str:
    context = tables.get("process_context")
    if context is not None and not context.empty and {"Field", "Value"}.issubset(context.columns):
        matches = context.loc[context["Field"].eq("Process Name"), "Value"]
        if not matches.empty:
            return str(matches.iloc[0])
    return "Jira"


def _format_cutoff(value: Any) -> str:
    parsed = pd.to_datetime(value, utc=True, errors="coerce")
    if pd.isna(parsed):
        return format_value(value)
    return parsed.strftime("%Y-%m-%d %H:%M UTC")


def build_report(
    task_metrics, output_path, *, title, evaluation_cutoff,
    timezone_name, work_days, work_window, process_data=None,
):
    from process_analysis import process_tables
    from metrics_engine import aggregate

    tables = process_data if process_data is not None else process_tables(task_metrics)
    document = Document()
    set_document_style(document)
    process_name = _process_name(tables)
    add_title(
        document,
        title,
        f"Space: {process_name} | Evaluation cutoff: {_format_cutoff(evaluation_cutoff)}",
    )

    add_heading(document, "Executive Summary")
    add_paragraph(
        document,
        "This report evaluates the selected Jira task scope through the stated cutoff. It summarizes task outcomes, timing, workflow exceptions, assignment distribution, data-quality coverage, and task-level evidence."
    )

    add_heading(document, "Process Context and Scope")
    add_key_value_table(document, dict(tables["process_context"].itertuples(index=False, name=None)))

    add_heading(document, "Overall Performance Indicators")
    row = tables["overall_summary"].iloc[0]
    names = {
        "Total tasks": "total_tasks",
        "Completed tasks": "completed_tasks",
        "Rejected tasks": "rejected_tasks",
        "Open tasks": "open_tasks",
        "WIP tasks": "wip_tasks",
        "Status unavailable": "unknown_status_tasks",
        "Complete histories": "history_complete_tasks",
        "Histories excluded": "history_excluded_tasks",
        "Reviewed eligible tasks": "reviewed_valid_tasks",
    }
    add_key_value_table(document, {name: int(row[key]) for name, key in names.items()})
    for label, count, denom in [
        ("Completion", "completed_tasks", "total_tasks"),
        ("On-time completion", "on_time_tasks", "on_time_valid_tasks"),
        ("Open overdue", "overdue_open_tasks", "overdue_valid_tasks"),
        ("Rework", "tasks_with_rework", "reviewed_valid_tasks"),
        ("Replanning", "tasks_with_replanning", "reviewed_valid_tasks"),
        ("Re-evaluation", "tasks_with_re_evaluation", "reviewed_valid_tasks"),
        ("Any review exception", "review_exception_tasks", "reviewed_valid_tasks"),
    ]:
        add_paragraph(
            document,
            f"{label}: {format_percent(int(row[count]), int(row[denom]))} ({int(row[count])}/{int(row[denom])} tasks).",
        )
    add_key_value_table(
        document,
        {
            kind.replace("_", " ").title() + " events": row["total_" + kind + "_count"]
            for kind in ["rework", "replanning", "re_evaluation"]
        },
    )

    add_heading(document, "Process Timing")
    for metric in ["execution", "lead_time", "time_to_start"]:
        add_paragraph(
            document,
            f"{metric.replace('_', ' ').title()}: mean {format_value(row['mean_' + metric + '_elapsed_hours'])} elapsed hours / {format_value(row['mean_' + metric + '_business_hours'])} business hours; valid tasks: {int(row[metric + '_elapsed_hours_valid_tasks'])}."
        )
    add_paragraph(
        document,
        "Business hours are process residence within the configured calendar, not recorded employee effort. Durations include waiting and repeated visits. Date adherence uses uploaded schedule values, not a reconstructed historical baseline."
    )

    add_heading(document, "Stage Residence and Open Work")
    for stage in tables["stage_summary"].to_dict("records"):
        add_heading(document, str(stage["status"]), 2)
        add_key_value_table(
            document,
            {
                "Tasks visited": stage["tasks_visited"],
                "Mean elapsed / business hours": f"{stage['elapsed_mean_hours']:.2f} / {stage['business_mean_hours']:.2f}",
                "Median elapsed / business hours": f"{stage['elapsed_median_hours']:.2f} / {stage['business_median_hours']:.2f}",
                "Total elapsed / business hours": f"{stage['elapsed_total_hours']:.2f} / {stage['business_total_hours']:.2f}",
                "Open tasks currently here": stage["open_tasks_currently_here"],
            },
        )
    add_paragraph(
        document,
        "Stage statistics sum all visits per task through cutoff, including unfinished visits. Done and Rejected residence is excluded. High residence identifies a review candidate; it does not establish a cause."
    )

    add_heading(document, "Bottlenecks, Exceptions, and Follow-up")
    findings = tables["process_findings"]
    if findings.empty:
        add_paragraph(document, "No supported exception findings are available in this scope.")
    for finding in findings.to_dict("records"):
        add_paragraph(document, f"{finding['issue_key']}: {finding['observation']} {finding['follow_up']}")

    add_heading(document, "Recommendations")
    add_paragraph(
        document,
        "The following recommendations are generated from the Jira results in this report and are limited to the selected scope and evaluation cutoff."
    )
    for recommendation in _jira_recommendations(row, task_metrics):
        _add_bullet(document, recommendation)

    add_heading(document, "Individual Achievements and Assignment Summary")
    add_paragraph(
        document,
        "Assignment uses the assignee recorded in the Jira snapshot. Unassigned is a separate group; these values support workload visibility and do not establish individual contribution or ownership at completion."
    )
    for group in aggregate(task_metrics, ["assignee_name"]).to_dict("records"):
        add_paragraph(
            document,
            f"{group['assignee_name']}: {group['total_tasks']} tasks; {group['completed_tasks']} completed, {group['open_tasks']} open, {group['rejected_tasks']} rejected, {group['unknown_status_tasks']} status unavailable."
        )

    add_heading(document, "Task-Level Evaluation")
    for task in task_metrics.to_dict("records"):
        add_heading(document, f"{task['issue_key']} - {task['task_name']}", 2)
        add_key_value_table(
            document,
            {
                "Status at cutoff": task["status_at_cutoff"],
                "Actual start / completion (UTC)": f"{format_value(task['actual_start_at'])} / {format_value(task['completed_at'])}",
                "Due date (uploaded snapshot)": task["due_date"],
                "On-time completion": task["on_time_completion"],
                "Execution elapsed / business hours": f"{format_value(task['execution_elapsed_hours'])} / {format_value(task['execution_business_hours'])}",
                "Open task age elapsed / business hours": f"{format_value(task['task_age_elapsed_hours'])} / {format_value(task['task_age_business_hours'])}",
                "Current status age elapsed / business hours": f"{format_value(task['current_status_age_elapsed_hours'])} / {format_value(task['current_status_age_business_hours'])}",
                "Overdue days": task["overdue_days"],
                "Rework / replanning / re-evaluation events": " / ".join(
                    format_value(task[k]) for k in ["rework_count", "replanning_count", "re_evaluation_count"]
                ),
                "History complete": task["history_complete"],
            },
        )

    add_heading(document, "Data Quality and Limitations")
    if tables["data_quality"].empty:
        add_paragraph(document, "No data-quality findings were recorded by these validation checks.")
    for item in tables["data_quality"].to_dict("records"):
        add_paragraph(document, f"{item['issue_key']}: {item['finding']}")

    add_heading(document, "Metric Definitions")
    for metric, definition in tables["metric_definitions"].itertuples(index=False, name=None):
        add_paragraph(document, f"{metric}: {definition}")
    add_paragraph(
        document,
        "All percentages use a 0-100 scale. A zero denominator is unavailable. The Excel export includes the transition audit trail and detailed stage statistics. Simulation data cannot establish long-term employee performance."
    )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build a Jira process-performance Word report.")
    parser.add_argument("--metrics-workbook", required=True, help="Workbook containing the task_metrics sheet.")
    parser.add_argument("--output", required=True, help="Output DOCX path.")
    parser.add_argument("--title", default="Jira Task Performance Evaluation Report", help="Report title.")
    parser.add_argument("--evaluation-cutoff", required=True, help="Evaluation cutoff shown in the report.")
    parser.add_argument("--timezone", default="Asia/Damascus", help="Report timezone.")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    workbook = pd.read_excel(args.metrics_workbook, sheet_name=None)
    task_metrics = workbook["task_metrics"]
    required = {
        "process_context",
        "overall_summary",
        "stage_summary",
        "process_findings",
        "data_quality",
        "metric_definitions",
    }
    process_data = {name: workbook[name] for name in required} if required.issubset(workbook) else None
    build_report(
        task_metrics,
        Path(args.output),
        title=args.title,
        evaluation_cutoff=args.evaluation_cutoff,
        timezone_name=args.timezone,
        work_days="Sunday-Thursday",
        work_window="09:00-17:00",
        process_data=process_data,
    )
    print(f"Word report written to: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
