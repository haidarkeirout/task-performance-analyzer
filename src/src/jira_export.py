"""Prepare a real XLSX and complete history JSON for the unchanged analysis API."""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from jira_gateway import CollectionError
from export_columns import DIRECT, fixed_headers


@dataclass
class PreparedData:
    xlsx: bytes = field(repr=False)
    history_json: bytes = field(repr=False)
    cutoff: str
    collected_at: str
    query: str
    fingerprint: str
    count: int
    filename: str
    space_name: str
    source_timezone: str


CANONICAL = ["Issue key", "Issue id", "Summary", "Issue Type", "Assignee", "Assignee Id",
             "Priority", "Status", "Resolution", "Created", "Due date", "Custom field (Start date)", "Labels"]
DIRECT_COLUMNS = {"summary": "Summary", "issuetype": "Issue Type", "assignee": "Assignee",
                  "priority": "Priority", "status": "Status", "resolution": "Resolution", "created": "Created",
                  "duedate": "Due date", "labels": "Labels"}


def _json(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def start_field_id(definitions, issues, preferred=""):
    ids = {key for issue in issues for key in issue.get("fields", {})}
    if preferred:
        if not re.fullmatch(r"customfield_\d+|[A-Za-z][A-Za-z0-9_]*", preferred):
            raise CollectionError("The planned start date field configuration is invalid.")
        if preferred not in ids:
            raise CollectionError("The configured planned start date field was not returned by Jira. Ask the administrator to check its configuration.")
        return preferred
    matches = [f["id"] for f in definitions if f.get("id") in ids and f.get("name", "").strip().casefold() == "start date"]
    if len(matches) > 1:
        used = [fid for fid in matches if any(i["fields"].get(fid) is not None for i in issues)]
        if len(used) == 1:
            return used[0]
        raise CollectionError("More than one Start date field is available. Ask the administrator to select the planned start date field.")
    return matches[0] if matches else None


def _named(value):
    return (value.get("name") or value.get("displayName") or value.get("value")) if isinstance(value, dict) else value


def _display(value):
    if isinstance(value, (dict, list)):
        if isinstance(value, dict) and value.get("name") and len(value) < 15:
            return value["name"]
        return _json(value)
    return value


def issue_rows(issues, definitions, preferred_start=""):
    start = start_field_id(definitions, issues, preferred_start)
    catalog = {f["id"]: f for f in definitions if f.get("id")}
    for issue in issues:
        for fid in issue.get("fields", {}):
            if fid not in catalog:
                catalog[fid] = {"id": fid, "name": (issue.get("names") or {}).get(fid, fid)}
    mapping = {**DIRECT_COLUMNS, **DIRECT}
    for fid, definition in catalog.items():
        if definition.get("name", "").casefold() == "sprint":
            mapping[fid] = "Sprint"
    if start:
        mapping[start] = "Custom field (Start date)"
    headers = list(dict.fromkeys([*CANONICAL, *mapping.values()]))
    field_rows, used = [], {h.casefold() for h in headers}
    for fid, definition in catalog.items():
        if fid not in mapping:
            label = definition.get("name") or fid
            header = f"Custom field ({label})" if fid.startswith("customfield_") else label
            if header.casefold() in used:
                header += f" [{fid}]"
            used.add(header.casefold())
            mapping[fid] = header
            headers.append(header)
        field_rows.append([fid, definition.get("name", fid), mapping[fid], _json(definition.get("schema") or {})])
    rows = []
    for issue in issues:
        fields = issue.get("fields", {})
        assignee = fields.get("assignee") or {}
        row = {"Issue key": issue["key"], "Issue id": str(issue.get("id", "")),
               "Summary": fields.get("summary"), "Issue Type": _named(fields.get("issuetype")),
               "Assignee": assignee.get("displayName"), "Assignee Id": assignee.get("accountId"),
               "Priority": _named(fields.get("priority")), "Status": _named(fields.get("status")),
               "Resolution": _named(fields.get("resolution")), "Created": fields.get("created"),
               "Due date": fields.get("duedate"), "Custom field (Start date)": fields.get(start) if start else None,
               "Labels": _json(fields.get("labels") or [])}
        project = fields.get("project") or {}
        lead = project.get("lead") or {}
        row.update({"Project key": project.get("key"), "Project name": project.get("name"),
                    "Project type": project.get("projectTypeKey"), "Project lead": lead.get("displayName"),
                    "Project lead id": lead.get("accountId"), "Project description": project.get("description"),
                    "Status Category": ((fields.get("status") or {}).get("statusCategory") or {}).get("name"),
                    "Votes": (fields.get("votes") or {}).get("votes"),
                    "Watchers": _json([w.get("displayName") for w in fields["watches"]["watchers"]]) if isinstance((fields.get("watches") or {}).get("watchers"), list) else None,
                    "Watchers Id": _json([w.get("accountId") for w in fields["watches"]["watchers"]]) if isinstance((fields.get("watches") or {}).get("watchers"), list) else None})
        for person, title in [("reporter", "Reporter"), ("creator", "Creator")]:
            value = fields.get(person) or {}
            row[title] = value.get("displayName")
            row[title + " Id"] = value.get("accountId")
        for fid, header in mapping.items():
            if header not in row:
                row[header] = _display(fields.get(fid))
        rows.append(row)
    return fixed_headers(mapping, headers), rows, field_rows


def _write_cell(ws, row, col, value):
    if isinstance(value, str):
        # Escape Excel-illegal control characters without losing their identity.
        value = ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\u{ord(m[0]):04x}", value)
        if len(value) > 32767:
            raise CollectionError("A generated Excel cell exceeds its size limit.")
    cell = ws.cell(row, col, value)
    if isinstance(value, str):
        # An issue title beginning '=' remains text and is never an Excel formula.
        cell.data_type = "s"


def _sheet(wb, name, headers, rows):
    if len(headers) > 16384 or len(rows) + 1 > 1048576:
        raise CollectionError("The selection exceeds Excel worksheet limits. Narrow the filters and collect again.")
    ws = wb.create_sheet(name)
    for i, header in enumerate(headers, 1):
        _write_cell(ws, 1, i, header)
        c = ws.cell(1, i)
        c.fill = PatternFill("solid", fgColor="17324D")
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(wrap_text=True)
    for index, values in enumerate(rows, 2):
        for col, value in enumerate(values, 1):
            _write_cell(ws, index, col, value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.iter_cols(min_row=1, max_row=1):
        label = str(column[0].value)
        ws.column_dimensions[column[0].column_letter].width = min(48, max(18, len(label) + 3))
    return ws


def build_workbook(issues, histories, definitions, *, cutoff, collected_at, query, space_name,
                   source_timezone="Asia/Damascus", preferred_start=""):
    headers, rows, field_rows = issue_rows(issues, definitions, preferred_start)
    # Preserve large text in numbered columns; preserve full structured values in Raw_JSON too.
    expanded_headers = list(headers)
    for row in rows:
        for header, value in list(row.items()):
            if isinstance(value, str) and len(value) > 30000:
                parts = [value[i:i+30000] for i in range(0, len(value), 30000)]
                row[header] = parts[0]
                for index, part in enumerate(parts[1:], 2):
                    extra = f"{header} [text part {index}]"
                    if extra not in expanded_headers: expanded_headers.append(extra)
                    row[extra] = part
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Jira_Data", expanded_headers, [[row.get(h) for h in expanded_headers] for row in rows])
    missing = [[h, "No value returned for this selection. This may be an empty, unavailable, or inaccessible field."]
               for h in headers if all(row.get(h) is None for row in rows)]
    _sheet(wb, "Source_Data_Quality", ["Column", "Note"], missing)
    events, coverage, changes, raw = [], [], [], []
    local_tz = ZoneInfo(source_timezone)
    for issue in issues:
        key, history = issue["key"], histories[issue["key"]]
        coverage.append([key, history["history_complete"], history["history_through"], history["initial_status"]])
        for event in history["status_events"]:
            stamp = datetime.fromisoformat(event["changed_at"].replace("Z", "+00:00")).astimezone(local_tz)
            events.append([key, event.get("from_status"), event.get("to_status"), stamp.date(), stamp.time().replace(tzinfo=None),
                           event.get("author_name"), event.get("change_id"), event["changed_at"]])
        for change in history.get("raw_changelog", []):
            for item in change.get("items", []):
                changes.append([key, change.get("id"), change.get("created"), item.get("field"), item.get("fieldId"),
                                item.get("from"), item.get("fromString"), item.get("to"), item.get("toString")])
        for section, value in [("issue", issue), ("history", history)]:
            serialized = _json(value)
            raw.extend([key, section, part//30000 + 1, serialized[part:part+30000]] for part in range(0, len(serialized), 30000))
    _sheet(wb, "Workflow_Events", ["Work Item ID", "From Status", "To Status", "Transition Date", "Transition Time", "Performed By", "Change ID", "Timestamp (UTC offset)"], events)
    _sheet(wb, "History_Coverage", ["Work Item ID", "History Complete", "History Through", "Initial Status"], coverage)
    # Arbitrary changelog strings may also be longer than a cell: Raw_JSON remains lossless.
    for row in changes:
        for i, value in enumerate(row):
            if isinstance(value, str) and len(value) > 30000:
                row[i] = value[:29900] + " [Full value in Raw_JSON]"
    _sheet(wb, "Field_Changes", ["Work Item ID", "Change ID", "Changed At", "Field", "Field ID", "From ID", "From Value", "To ID", "To Value"], changes)
    _sheet(wb, "Field_Catalog", ["Field ID", "Jira Name", "Excel Column", "Schema"], field_rows)
    _sheet(wb, "Raw_JSON", ["Work Item ID", "Section", "Part", "JSON Text"], raw)
    context = {"Process Name": space_name, "Evaluation Scope": "Selected Jira work items",
               "Dataset Type": "Jira API collection", "Evaluation Cutoff Date": cutoff,
               "Collection Completed At": collected_at, "Source Timezone": source_timezone,
               "Selected JQL": query, "Work Item Count": len(issues),
               "History": "All accessible changelog pages collected; full history is not limited by task date filters.",
               "Export": "Real XLSX generated from Jira REST fields. Jira_Data is for analysis; Raw_JSON preserves structured values and full history."}
    context_rows = []
    for key, value in context.items():
        text = str(value)
        for part in range(0, len(text), 30000):
            context_rows.append([key if part == 0 else f"{key} [part {part//30000 + 1}]", text[part:part+30000]])
    _sheet(wb, "Process_Context", ["Field", "Value"], context_rows)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def collect_data(gateway, space, query, fingerprint, definitions, progress=None, clock=None, checkpoint=None):
    clock = clock or (lambda: datetime.now(timezone.utc))
    checkpoint = {} if checkpoint is None else checkpoint
    if checkpoint and checkpoint.get("fingerprint") != fingerprint:
        raise CollectionError("The selection changed. Start a new collection.")
    checkpoint.setdefault("fingerprint", fingerprint)
    checkpoint.setdefault("cutoff", clock().isoformat())
    checkpoint.setdefault("completed", {})
    cutoff = checkpoint["cutoff"]
    if "issues" not in checkpoint:
        issues = gateway.all_issues(query, progress=progress)
        keys = [item.get("key") for item in issues]
        if not issues:
            raise CollectionError("No work items match these filters. Change the filters and click Done again.")
        if any(not key for key in keys) or len(set(keys)) != len(keys):
            raise CollectionError("The selected work item list is incomplete or contains duplicates.")
        checkpoint["issues"] = issues
    issues = checkpoint["issues"]
    if "project" not in checkpoint:
        if progress: progress("Reading project details...")
        checkpoint["project"] = gateway.project_details(space["id"]) if hasattr(gateway, "project_details") else space
    completed = checkpoint["completed"]
    for number, issue in enumerate(issues, 1):
        if issue["key"] in completed:
            continue
        if progress: progress(f"Completed {len(completed)} of {len(issues)}; reading {issue['key']}...")
        item, history = gateway.complete_issue(issue)
        actual_project = str((item.get("fields", {}).get("project") or {}).get("id", ""))
        if actual_project != str(space["id"]) or item.get("key") != issue["key"]:
            raise CollectionError("A work item changed spaces or keys during collection. Start a new collection.")
        if history.get("history_complete") is not True or not history.get("history_through"):
            raise CollectionError("A task history is incomplete. No partial export was prepared.")
        if datetime.fromisoformat(history["history_through"].replace("Z", "+00:00")) < datetime.fromisoformat(cutoff.replace("Z", "+00:00")):
            raise CollectionError("A task history does not cover the evaluation cutoff.")
        item["fields"]["project"] = {**checkpoint["project"], **(item["fields"].get("project") or {})}
        # Commit the pair only AFTER both fields and history have been verified.
        completed[issue["key"]] = (item, history)
        if progress: progress(f"Completed {len(completed)} of {len(issues)} work items.")
    full = [completed[issue["key"]][0] for issue in issues]
    histories = {issue["key"]: completed[issue["key"]][1] for issue in issues}
    collected_at = clock().isoformat()
    if progress: progress("Preparing your Excel file...")
    data = build_workbook(full, histories, definitions, cutoff=cutoff, collected_at=collected_at,
                          query=query, space_name=space["name"], source_timezone=gateway.settings.source_timezone,
                          preferred_start=gateway.settings.start_date_field)
    filename = f"Jira_{space['key']}_{clock().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return PreparedData(data, _json(histories).encode(), cutoff, collected_at, query, fingerprint,
                        len(full), filename, space["name"], gateway.settings.source_timezone)
