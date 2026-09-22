"""Shared visual system for Employee and Project Word reports."""
from __future__ import annotations

import math
from datetime import date, datetime
from typing import Any, Iterable, Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "6B7280"
HEADER_FILL = "F2F4F7"
BORDER = "D9D9D9"
USABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def _set_run_font(run, *, size=11, bold=False, italic=False, color=INK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def _set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def _set_cell_margins(cell, top=90, bottom=90, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: Sequence[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for index, width in enumerate(widths):
        columns = list(grid.iterchildren())
        if index < len(columns):
            columns[index].set(qn("w:w"), str(width))
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            if index < len(widths):
                _set_cell_width(cell, widths[index])
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def _value_text(value: Any) -> str:
    if value is None:
        return "Unavailable"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float):
        if math.isnan(value):
            return "Unavailable"
        return f"{value:.2f}"
    text = str(value).strip()
    return text if text and text.lower() not in {"nan", "none", "null"} else "Unavailable"


def configure_report_document(
    document: Document,
    *,
    header_label: str = "PROCESS PERFORMANCE REPORT | PERFORMANCE EVALUATION",
    footer_label: str = "Performance Management",
):
    for section in document.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.82)
        section.right_margin = Inches(0.82)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)

        header = section.header.paragraphs[0]
        header.text = ""
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_paragraph_spacing(header, after=0, line=1.0)
        _set_run_font(header.add_run(header_label.upper()), size=8, color=MUTED)

        footer = section.footer.paragraphs[0]
        footer.text = ""
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_paragraph_spacing(footer, after=0, line=1.0)
        _set_run_font(footer.add_run(footer_label), size=8, color=MUTED)
        _set_run_font(footer.add_run("  |  Page "), size=8, color=MUTED)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title_style = document.styles["Title"]
    title_style.font.name = "Calibri"
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(INK)
    title_style.paragraph_format.space_before = Pt(3)
    title_style.paragraph_format.space_after = Pt(8)
    title_style.paragraph_format.line_spacing = 1.0

    tokens = {
        "Heading 1": (17, BLUE, 15, 8),
        "Heading 2": (13, BLUE, 11, 5),
        "Heading 3": (11, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in tokens.items():
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.keep_with_next = True


def add_report_title(
    document: Document,
    title: str,
    subtitle: str = "",
    metadata: Iterable[str] = (),
):
    label = document.add_paragraph()
    _set_paragraph_spacing(label, after=5, line=1.0)
    _set_run_font(label.add_run("PROCESS PERFORMANCE REPORT"), size=10, bold=True, color=INK)

    title_paragraph = document.add_paragraph(style="Title")
    _set_paragraph_spacing(title_paragraph, after=7, line=1.0)
    title_run = title_paragraph.add_run(title)
    _set_run_font(title_run, size=25, bold=True, color=INK)

    if subtitle:
        paragraph = document.add_paragraph()
        _set_paragraph_spacing(paragraph, after=8, line=1.0)
        _set_run_font(paragraph.add_run(subtitle), size=14, italic=True, color=MUTED)

    for line in metadata:
        paragraph = document.add_paragraph()
        _set_paragraph_spacing(paragraph, after=3, line=1.0)
        _set_run_font(paragraph.add_run(str(line)), size=10, color=MUTED)


def add_report_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        _set_run_font(
            run,
            size={1: 17, 2: 13, 3: 11}.get(level, 11),
            bold=True,
            color={1: BLUE, 2: BLUE, 3: DARK_BLUE}.get(level, INK),
        )
    return paragraph


def add_report_paragraph(document: Document, text: str):
    paragraph = document.add_paragraph()
    _set_paragraph_spacing(paragraph)
    _set_run_font(paragraph.add_run(str(text)), size=11, color=INK)
    return paragraph


def add_report_bullet(document: Document, text: str):
    paragraph = document.add_paragraph(style="List Bullet")
    _set_paragraph_spacing(paragraph, after=4)
    _set_run_font(paragraph.add_run(str(text)), size=11, color=INK)
    return paragraph


def add_report_table(
    document: Document,
    headers: Sequence[Any],
    rows: Iterable[Sequence[Any]],
    widths: Sequence[int] | None = None,
):
    headers = list(headers)
    rows = [list(row) for row in rows]
    if not rows:
        rows = [["Unavailable"] + [""] * (len(headers) - 1)]
    if widths is None:
        base = USABLE_WIDTH_DXA // max(1, len(headers))
        widths = [base] * len(headers)
        widths[-1] += USABLE_WIDTH_DXA - sum(widths)
    else:
        widths = list(widths)
        if len(widths) != len(headers):
            base = USABLE_WIDTH_DXA // max(1, len(headers))
            widths = [base] * len(headers)
            widths[-1] += USABLE_WIDTH_DXA - sum(widths)

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _set_table_geometry(table, widths)
    _repeat_table_header(table.rows[0])

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_shading(cell, HEADER_FILL)
        _set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_paragraph_spacing(paragraph, after=0, line=1.0)
        _set_run_font(paragraph.add_run(_value_text(header)), size=9, bold=True, color=INK)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values[: len(headers)]):
            cell = cells[index]
            _set_cell_border(cell)
            paragraph = cell.paragraphs[0]
            _set_paragraph_spacing(paragraph, after=0, line=1.0)
            _set_run_font(paragraph.add_run(_value_text(value)), size=9, color=INK)
        if row_index % 2 == 1:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, widths)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def add_report_key_value_table(document: Document, values: dict[str, Any]):
    return add_report_table(
        document,
        ["Metric", "Value"],
        [(key, value) for key, value in values.items()],
        [2700, 6660],
    )
