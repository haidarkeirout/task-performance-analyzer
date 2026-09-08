"""Offline transport, filter, export, and unchanged-engine regression checks."""
import ast
import copy
import io
import json
import os
from pathlib import Path
import sys
import unittest
from datetime import datetime, timezone
from unittest.mock import Mock
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from automation_auth import Settings, SetupError, read_settings, credentials_match, hash_password
from jira_gateway import JiraGateway, CollectionError
from jira_filters import FilterField, FilterError, field_catalog, field_clause, date_clause, build_query
from jira_export import build_workbook, collect_data, issue_rows
from jira_processor import choose_sheet, normalize_tasks
from metrics_engine import analyze_tasks, aggregate, WorkCalendar
from process_analysis import workbook_histories, process_tables


SETTINGS = Settings("test-admin", "test-password", "", "https://test.atlassian.net", "test@example.invalid", "test-token")
CUTOFF = "2026-09-03T20:59:59+00:00"
FIELDS = [
    {"id": "summary", "name": "Summary", "schema": {"type": "string"}},
    {"id": "customfield_10015", "name": "Start date", "schema": {"type": "date"}},
    {"id": "customfield_12345", "name": "Business value", "schema": {"type": "number"}},
]


def example_issue(key="TEST-1", status="Idea"):
    return {"id": "10001", "key": key, "fields": {
        "summary": "Test task", "project": {"id": "100", "key": "TEST", "name": "Test space"},
        "status": {"name": status}, "issuetype": {"name": "Task"}, "priority": {"name": "High"},
        "assignee": None, "resolution": None, "labels": ["first", "second"],
        "created": "2026-09-01T06:00:00+00:00", "updated": "2026-09-01T06:00:00+00:00",
        "duedate": "2026-09-05", "customfield_10015": "2026-09-01", "customfield_12345": 5,
    }}


def no_events(key="TEST-1"):
    return {"issue_key": key, "history_complete": True, "history_through": CUTOFF,
            "initial_status": "Idea", "snapshot": {"current_status": "Idea"},
            "status_events": [], "assignee_events": [], "raw_changelog": []}


def app_functions():
    source = ast.parse((ROOT / "app.py").read_text())
    names = {"ROOT_DIR", "SRC_DIR", "CONFIG_DIR", "INPUT_CONFIG_PATH", "METRICS_CONFIG_PATH"}
    nodes = [n for n in source.body if isinstance(n, (ast.Import, ast.ImportFrom, ast.FunctionDef)) or
             isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id in names for t in n.targets)]
    namespace = {"__file__": str(ROOT / "app.py")}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), "app.py", "exec"), namespace)
    return namespace


class AuthenticationTests(unittest.TestCase):
    def test_login_fail_closed_and_no_secret_repr(self):
        self.assertFalse(credentials_match("other", "test-password", SETTINGS))
        self.assertFalse(credentials_match("test-admin", "wrong", SETTINGS))
        self.assertTrue(credentials_match("test-admin", "test-password", SETTINGS))
        self.assertNotIn("test-token", repr(SETTINGS))
        self.assertNotIn("test-password", repr(SETTINGS))
        with self.assertRaises(SetupError): read_settings({}, {})

    def test_hashed_password_and_configuration_rotation(self):
        hashed = Settings("test-admin", "", hash_password("pássword", 100_000), SETTINGS.jira_url, SETTINGS.jira_email, SETTINGS.jira_token)
        self.assertTrue(credentials_match("test-admin", "pássword", hashed))
        self.assertFalse(credentials_match("test-admin", "password", hashed))
        self.assertNotEqual(hashed.revision, SETTINGS.revision)


class FilterTests(unittest.TestCase):
    def test_dynamic_custom_fields_and_operators(self):
        reference = {"visibleFieldNames": [
            {"value": "summary", "operators": ["~", "!~", "is"], "types": ["java.lang.String"]},
            {"value": "Start date", "cfid": "cf[10015]", "displayName": "Start date", "operators": ["=", ">", "<"], "types": ["java.util.Date"]},
            {"value": "cf[12345]", "cfid": "cf[12345]", "displayName": "Business value", "operators": ["=", ">="], "types": ["java.lang.Double"]},
            {"value": "Hidden", "searchable": "false", "operators": ["="]},
        ]}
        catalog = field_catalog(reference, FIELDS)
        self.assertIn("cf[12345]", catalog)
        self.assertEqual(catalog["cf[10015]"].kind, "date")
        self.assertEqual(field_clause(catalog["cf[12345]"], ">=", ["3.5"]), "cf[12345] >= 3.5")
        self.assertNotIn("hidden", catalog)
        with self.assertRaises(FilterError): field_clause(catalog["cf[12345]"], "in", ["3"])

    def test_safe_literals_project_scope_and_expression_validation(self):
        f = FilterField("status", "Status", "status", "value", ("in", "is"))
        clause = field_clause(f, "in", ['Done\") OR project = EVIL OR (status = \"Open'])
        query = build_query("TEST", [clause])
        self.assertTrue(query.startswith('project = "TEST" AND (status IN ('))
        self.assertIn('\\"', query)
        with self.assertRaises(FilterError): build_query("TEST", advanced=") OR project = EVIL (")
        with self.assertRaises(FilterError): build_query("TEST", advanced='status = Done ORDER BY created')
        self.assertIn('(status = Done OR priority = High)', build_query("TEST", advanced="status = Done OR priority = High"))
        self.assertEqual(field_clause(f, "is"), "status IS EMPTY")

    def test_date_end_is_inclusive_and_invalid_range_fails(self):
        from datetime import date
        f = FilterField("created", "Created", "created", "date", ("=", ">=", "<=", "is"))
        result = date_clause(f, "Between", date(2026, 8, 1), date(2026, 8, 31))
        self.assertEqual(result, 'created >= "2026-08-01" AND created < "2026-09-01"')
        with self.assertRaises(FilterError): date_clause(f, "Between", date(2026, 9, 1), date(2026, 8, 31))


class GatewayTests(unittest.TestCase):
    def gateway(self, replies):
        gateway = JiraGateway(SETTINGS)
        gateway.request = Mock(side_effect=replies)
        return gateway

    def test_all_spaces_pages_even_when_server_caps_page_size(self):
        g = self.gateway([{"startAt": 0, "total": 2, "maxResults": 1, "values": [{"id": "1"}], "isLast": False},
                          {"startAt": 1, "total": 2, "maxResults": 1, "values": [{"id": "2"}], "isLast": True}])
        self.assertEqual(len(g.spaces()), 2)
        self.assertEqual(g.request.call_args.kwargs["params"]["startAt"], 1)

    def test_search_cursor_pages_and_repeated_cursor_failure(self):
        g = self.gateway([{"queries": [{"errors": []}]},
                          {"issues": [{"key": "T-1"}], "nextPageToken": "page-two", "isLast": False},
                          {"issues": [{"key": "T-2"}], "isLast": True}])
        self.assertEqual(len(g.all_issues('project = "T"')), 2)
        self.assertEqual(g.request.call_args.kwargs["body"]["nextPageToken"], "page-two")
        self.assertEqual(g.request.call_args.kwargs["body"]["fields"], ["*all"])
        g = self.gateway([{"queries": [{}]}, {"issues": [], "isLast": False, "nextPageToken": "again"},
                          {"issues": [], "isLast": False, "nextPageToken": "again"}])
        with self.assertRaises(CollectionError): g.all_issues("project = T")

    def test_incomplete_or_duplicate_history_pages_fail(self):
        for pages in [
            [{"values": [], "startAt": 0, "total": 2, "isLast": True}],
            [{"values": [{"id": "1"}], "total": 2, "isLast": False}, {"values": [{"id": "1"}], "total": 2, "startAt": 1, "isLast": True}],
        ]:
            with self.assertRaises(CollectionError): self.gateway(pages)._paged("/rest/api/3/issue/T-1/changelog")

    def test_equal_time_transitions_keep_jira_order(self):
        item = example_issue(status="Done")
        changes = [
            {"id": "2", "created": "2026-09-02T10:00:00+00:00", "items": [
                {"field": "status", "fromString": "Idea", "toString": "In Progress"}]},
            {"id": "10", "created": "2026-09-02T10:00:00+00:00", "items": [
                {"field": "status", "fromString": "In Progress", "toString": "Done"}]},
        ]
        check = {"fields": {"updated": item["fields"]["updated"], "status": item["fields"]["status"]}}
        _, history = self.gateway([item, {"values": changes, "total": 2}, check]).complete_issue(item)
        self.assertEqual([event["change_id"] for event in history["status_events"]], ["2", "10"])
        self.assertEqual(history["initial_status"], "Idea")

    def test_complete_issue_paginates_nested_fields_and_retries_edit_race(self):
        item = example_issue()
        item["fields"]["comment"] = {"total": 2, "comments": []}
        item["fields"]["worklog"] = {"total": 0, "worklogs": []}
        final_check = {"fields": {"updated": item["fields"]["updated"], "status": item["fields"]["status"]}}
        pages = [item, {"values": [], "total": 0},
                 {"comments": [{"id": "1"}], "total": 2, "startAt": 0},
                 {"comments": [{"id": "2"}], "total": 2, "startAt": 1},
                 {"worklogs": [], "total": 0}, final_check]
        full, history = self.gateway(pages).complete_issue(item)
        self.assertEqual(full["fields"]["comment"]["total"], 2)
        self.assertTrue(history["history_complete"])
        self.assertEqual(history["initial_status"], "Idea")
        simple = example_issue()
        g = self.gateway([simple, {"values": [], "total": 0}, {"fields": {"updated": "changed"}},
                          simple, {"values": [], "total": 0}, final_check])
        self.assertTrue(g.complete_issue(simple)[1]["history_complete"])

    def test_http_retry_and_redacted_errors(self):
        response = Mock(status_code=401)
        response.json.return_value = {"secret": "test-token"}
        session = Mock()
        session.request.return_value = response
        gateway = JiraGateway(SETTINGS, session=session, sleep=lambda delay: None)
        with self.assertRaises(CollectionError) as error: gateway.identity()
        self.assertNotIn("test-token", str(error.exception))
        busy, good = Mock(status_code=429, headers={"Retry-After": "0"}), Mock(status_code=200)
        good.json.return_value = {"accountId": "1"}
        session.request.side_effect = [busy, good]
        self.assertEqual(gateway.identity()["accountId"], "1")
        self.assertFalse(session.request.call_args.kwargs["allow_redirects"])


class ExportTests(unittest.TestCase):
    def build(self, item=None, definitions=None):
        return build_workbook([item or example_issue()], {"TEST-1": no_events()}, definitions or FIELDS,
                               cutoff=CUTOFF, collected_at=CUTOFF, query="project = TEST", space_name="Test")

    def test_real_xlsx_labels_custom_fields_and_raw_json_round_trip(self):
        item = example_issue()
        item["fields"]["summary"] = '=HYPERLINK("https://example.invalid","literal task name")'
        item["fields"]["description"] = {"type": "doc", "content": [{"text": "a" * 40000}]}
        blob = self.build(item)
        self.assertTrue(blob.startswith(b"PK"))
        book = load_workbook(io.BytesIO(blob))
        headers = [c.value for c in book["Jira_Data"][1]]
        cell = book["Jira_Data"].cell(2, headers.index("Summary")+1)
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, item["fields"]["summary"])
        raw = [r for r in book["Raw_JSON"].iter_rows(min_row=2, values_only=True) if r[1] == "issue"]
        restored = json.loads("".join(row[3] for row in raw))
        self.assertEqual(restored, item)
        sheets = pd.read_excel(io.BytesIO(blob), sheet_name=None, dtype=object)
        config = json.loads((ROOT / "configs/jira_input_config.json").read_text())
        name, source = choose_sheet(sheets, config, None)
        self.assertEqual(name, "Jira_Data")
        normalized, _ = normalize_tasks(source, input_config=config, source_timezone=ZoneInfo("Asia/Damascus"), labels_delimiter=None, explicit_date_format=None)
        self.assertEqual(normalized.iloc[0]["labels"], ["first", "second"])
        self.assertEqual(normalized.iloc[0]["planned_start_date"], "2026-09-01")

    def test_empty_selection_and_failed_history_do_not_return_ready(self):
        gateway = Mock(settings=SETTINGS)
        gateway.all_issues.return_value = []
        with self.assertRaises(CollectionError): collect_data(gateway, {"id": "100", "key": "TEST", "name": "Test"}, "project=TEST", "fp", FIELDS)
        gateway.all_issues.return_value = [example_issue()]
        gateway.complete_issue.side_effect = CollectionError("Incomplete history")
        with self.assertRaises(CollectionError): collect_data(gateway, {"id": "100", "key": "TEST", "name": "Test"}, "project=TEST", "fp", FIELDS)

    def test_generated_workbook_runs_unchanged_pipeline_and_both_downloads(self):
        ns = app_functions()
        frame, log, tables = ns["run_analysis"](io.BytesIO(self.build()), io.BytesIO(json.dumps({"TEST-1": no_events()}).encode()), "", "", "", CUTOFF, "Asia/Damascus", history_source="History JSON")
        self.assertEqual(len(frame), 1)
        self.assertTrue(frame.iloc[0]["history_complete"])
        self.assertEqual(frame.iloc[0]["status_at_cutoff"], "Idea")
        self.assertTrue(ns["create_excel_download"](frame, None, None, None, tables).startswith(b"PK"))
        self.assertTrue(ns["create_word_report_download"](frame, CUTOFF, tables).startswith(b"PK"))

    def test_duplicate_start_field_requires_explicit_mapping(self):
        item = example_issue()
        item["fields"]["customfield_20015"] = "2026-09-02"
        fields = FIELDS + [{"id": "customfield_20015", "name": "Start date"}]
        with self.assertRaises(CollectionError): issue_rows([item], fields)
        self.assertEqual(issue_rows([item], fields, "customfield_20015")[1][0]["Custom field (Start date)"], "2026-09-02")


@unittest.skipUnless(os.environ.get("JIRA_TEST_WORKBOOK"), "Set JIRA_TEST_WORKBOOK for original-data regression.")
class OriginalDataRegression(unittest.TestCase):
    def test_api_export_and_original_workbook_produce_identical_task_metrics(self):
        workbook = pd.read_excel(os.environ["JIRA_TEST_WORKBOOK"], sheet_name=None, dtype=object)
        config = json.loads((ROOT / "configs/jira_input_config.json").read_text())
        _, source = choose_sheet(workbook, config, None)
        tasks, _ = normalize_tasks(source, input_config=config, source_timezone=ZoneInfo("Asia/Damascus"), labels_delimiter=None, explicit_date_format=None)
        histories = workbook_histories(workbook, tasks, "Asia/Damascus", CUTOFF, True)
        original = analyze_tasks(tasks, histories, pd.Timestamp(CUTOFF), WorkCalendar())
        issues = []
        for task in tasks.to_dict("records"):
            f = {"summary": task["task_name"], "status": {"name": task["current_status"]},
                 "issuetype": {"name": task["issue_type"]}, "created": task["created_at"], "duedate": task["due_date"],
                 "customfield_10015": task["planned_start_date"], "priority": {"name": task["priority"]},
                 "assignee": {"displayName": task["assignee_name"], "accountId": task["assignee_id"]},
                 "labels": task["labels"], "resolution": {"name": task["resolution"]}}
            issues.append({"key": task["issue_key"], "id": str(task["issue_id"]), "fields": f})
        for history in histories.values():
            history["raw_changelog"] = []
        xlsx = build_workbook(issues, histories, FIELDS, cutoff=CUTOFF, collected_at=CUTOFF, query="project=SCRUM", space_name="Performance Analysis")
        ns = app_functions()
        automated, _, _ = ns["run_analysis"](io.BytesIO(xlsx), io.BytesIO(json.dumps(histories).encode()), "", "", "", CUTOFF, "Asia/Damascus", history_source="History JSON")
        pd.testing.assert_frame_equal(original.reset_index(drop=True), automated.reset_index(drop=True), check_dtype=False)
        pd.testing.assert_frame_equal(aggregate(original, []), aggregate(automated, []), check_dtype=False)


if __name__ == "__main__":
    unittest.main()
