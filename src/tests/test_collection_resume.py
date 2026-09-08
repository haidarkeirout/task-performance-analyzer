import copy
import io
import threading
import unittest
from datetime import datetime, timezone
from unittest.mock import patch
from openpyxl import load_workbook
from test_automation import SETTINGS, FIELDS, example_issue, no_events
from collection_job import CollectionJob
from jira_export import build_workbook, issue_rows
from jira_gateway import CollectionError
from test_automation_ui import FakeJira, button, finish_collection
import test_automation_ui as ui_tests
import automation_ui

SPACE = {"id": "100", "key": "TEST", "name": "Test"}

class ResumeGateway(FakeJira):
    visited = []
    fail_once = True
    def all_issues(self, query, progress=None):
        return [example_issue('TEST-' + str(n)) for n in range(1, 25)]
    def complete_issue(self, item):
        self.visited.append(item['key'])
        if item['key'] == 'TEST-17' and type(self).fail_once:
            type(self).fail_once = False
            raise CollectionError('Simulated temporary interruption')
        history = no_events(item['key'])
        history['history_through'] = datetime.now(timezone.utc).isoformat()
        return copy.deepcopy(item), history

class ResumeTests(unittest.TestCase):
    def test_twenty_four_tasks_resume_without_repeating_completed_tasks(self):
        ResumeGateway.visited, ResumeGateway.fail_once = [], True
        job = CollectionJob(SETTINGS, SPACE, 'project=TEST', 'fp', FIELDS, ResumeGateway)
        job.start(); job.step()
        while job.running: job.step()
        self.assertFalse(job.running)
        self.assertEqual(job.snapshot()['completed'], 16)
        self.assertIsNone(job.result)
        self.assertIn('TEST-17', job.error)
        cutoff = job.checkpoint['cutoff']
        job.start()
        while job.running: job.step()
        self.assertIsNone(job.error)
        self.assertEqual(job.result.count, 24)
        self.assertEqual(job.result.cutoff, cutoff)
        self.assertEqual(ResumeGateway.visited.count('TEST-1'), 1)
        self.assertEqual(ResumeGateway.visited.count('TEST-17'), 2)
        book = load_workbook(io.BytesIO(job.result.xlsx))
        self.assertEqual(book['Jira_Data'].max_row, 25)
        self.assertEqual(book['History_Coverage'].max_row, 25)

    def test_fixed_columns_project_people_and_duplicate_custom_fields(self):
        item = example_issue()
        item['fields']['project'].update(projectTypeKey='software', lead={'displayName':'Lead','accountId':'a1'}, description='Project description')
        item['fields']['reporter'] = {'displayName':'Reporter','accountId':'a2'}
        item['fields']['customfield_1'] = 'first'
        item['fields']['customfield_2'] = 'second'
        fields = FIELDS + [{'id': 'customfield_1','name':'Vulnerability'}, {'id':'customfield_2','name':'Vulnerability'}]
        headers, rows, _ = issue_rows([item], fields)
        self.assertEqual(headers[:5], ['Summary','Issue key','Issue id','Issue Type','Status'])
        self.assertEqual(len(headers),len(set(headers)))
        self.assertIn('Σ Time Spent', headers)
        self.assertIn('Watchers Id', headers)
        self.assertEqual(rows[0]['Project lead id'], 'a1')
        self.assertEqual(rows[0]['Reporter Id'], 'a2')
        self.assertEqual(rows[0]['Custom field (Vulnerability)'], 'first')
        self.assertEqual(rows[0]['Custom field (Vulnerability) [customfield_2]'], 'second')

    def test_worker_survives_interface_rerun(self):
        started, release = threading.Event(), threading.Event()
        class SlowGateway(FakeJira):
            def complete_issue(self, item):
                started.set()
                if not release.wait(10): raise CollectionError('Test gate timeout')
                return super().complete_issue(item)
        helper = ui_tests.UserJourneyTests()
        with patch.object(automation_ui, 'JiraGateway', SlowGateway):
            at = helper.app(); helper.sign_in(at); helper.select_space(at)
            try:
                button(at,'Done').click().run()
                job = at.session_state['collection_job']
                at.run()
                self.assertIs(at.session_state['collection_job'], job)
                self.assertFalse(job.running)
                self.assertNotIn('prepared_data',at.session_state)
                release.set(); finish_collection(at)
                self.assertEqual(len(at.exception),0)
            finally:
                release.set()

    def test_changed_filter_cancels_pending_worker(self):
        started, release = threading.Event(), threading.Event()
        class SlowGateway(FakeJira):
            def complete_issue(self, item):
                started.set(); release.wait(10)
                return super().complete_issue(item)
        helper=ui_tests.UserJourneyTests()
        with patch.object(automation_ui,'JiraGateway',SlowGateway):
            at=helper.app(); helper.sign_in(at); helper.select_space(at)
            try:
                button(at,'Done').click().run(); self.assertTrue(started.wait(2))
                old=at.session_state['collection_job']
                at.text_input(key='search_text').input('different').run()
                self.assertTrue(old.cancelled)
                release.set(); at.run()
                self.assertIsNone(old.result)
                self.assertNotIn('prepared_data',at.session_state)
            finally: release.set()
