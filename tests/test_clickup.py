import io
import json
import unittest
from unittest.mock import Mock

from openpyxl import load_workbook
from docx import Document

from clickup_analysis import analyze_clickup, analysis_excel
from clickup_export import collect_data
from clickup_filters import filter_tasks
from clickup_gateway import ClickUpGateway, ClickUpTimeStatusUnavailable
from clickup_report import create_clickup_word_report


def status_payload(current="IN PROGRESS", current_minutes=120, history=None):
    return {
        "current_status": {
            "status": current,
            "color": "#5f55ee",
            "total_time": {"by_minute": current_minutes, "since": "2026-09-01T10:00:00Z"},
        },
        "status_history": history or [
            {"status": "TO DO", "color": "#ccc", "type": "open",
             "total_time": {"by_minute": 60, "since": "2026-09-01T09:00:00Z"}, "orderindex": 0},
        ],
    }


class ClickUpGatewayTests(unittest.TestCase):
    def test_bulk_time_in_status_uses_public_api_without_browser(self):
        session = Mock()
        response = Mock(status_code=200)
        response.json.return_value = {"a": status_payload(), "b": status_payload("COMPLETE", 30)}
        session.get.return_value = response
        gateway = ClickUpGateway("pk_test", session=session, sleep=lambda _: None)
        result = gateway.time_in_status(["a", "b"])
        self.assertEqual(set(result), {"a", "b"})
        url = session.get.call_args.args[0]
        self.assertTrue(url.endswith("/task/bulk_time_in_status/task_ids"))
        self.assertEqual(session.get.call_args.kwargs["params"], [("task_ids", "a"), ("task_ids", "b")])

    def test_time_in_status_permission_failure_is_explicit(self):
        session = Mock()
        response = Mock(status_code=404)
        response.json.return_value = {}
        session.get.return_value = response
        gateway = ClickUpGateway("pk_test", session=session, sleep=lambda _: None)
        with self.assertRaises(ClickUpTimeStatusUnavailable):
            gateway.time_in_status(["a"])


class ClickUpCollectionAndAnalysisTests(unittest.TestCase):
    def test_collection_uses_time_status_and_not_activity(self):
        tasks = [{
            "id": "a", "name": "Ship release", "status": {"status": "COMPLETE"},
            "priority": {"priority": "high"}, "assignees": [{"username": "Haidar"}],
            "date_created": "1788249600000", "due_date": "1788508800000",
            "date_closed": "1788422400000", "space": {"id": "s"},
        }]
        gateway = Mock()
        gateway.time_in_status.return_value = {"a": status_payload("COMPLETE", 30)}
        gateway.activity.side_effect = AssertionError("Activity collector must not be called")
        prepared = collect_data(gateway, tasks, "Performance Analysis", "fp", time_status_data=gateway.time_in_status.return_value)
        book = load_workbook(io.BytesIO(prepared.xlsx), data_only=True)
        headers = [cell.value for cell in book["ClickUp_Data"][1]]
        self.assertIn("Total Time in Status (JSON)", headers)
        self.assertIn("Time in Status - COMPLETE (min)", headers)
        self.assertEqual(book["Activity"].max_row, 1)
        self.assertFalse(prepared.clickup_activity_available)
        self.assertTrue(prepared.clickup_time_status_available)
        result = analyze_clickup(prepared)
        self.assertEqual(len(result["tasks"]), 1)
        self.assertTrue(bool(result["tasks"].iloc[0]["Completed?"]))
        self.assertAlmostEqual(result["tasks"].iloc[0]["Total Time in Status (min)"], 90.0)
        self.assertIn("COMPLETE", set(result["status_summary"]["Status"]))

    def test_due_variance_filters_and_dashboard_tables_are_clickup_only(self):
        tasks = [
            {
                "id": "late", "name": "Deliver late", "status": {"status": "COMPLETE", "type": "done"},
                "priority": {"priority": "high"}, "assignees": [{"username": "Haidar"}],
                "tags": [{"name": "release"}], "date_created": "2026-09-01T09:00:00Z",
                "due_date": "2026-09-03T09:00:00Z", "date_closed": "2026-09-05T09:00:00Z",
            },
            {
                "id": "early", "name": "Deliver early", "status": {"status": "COMPLETE", "type": "done"},
                "priority": {"priority": "low"}, "assignees": [{"username": "Maya"}],
                "tags": [{"name": "release"}], "date_created": "2026-09-01T09:00:00Z",
                "due_date": "2026-09-05T09:00:00Z", "date_closed": "2026-09-04T09:00:00Z",
            },
        ]
        filtered = filter_tasks(
            tasks,
            {
                "status": ["COMPLETE"],
                "priority": ["high"],
                "extras": {"Tags": {"mode": "Has any of", "values": ["release"]}},
            },
        )
        self.assertEqual([task["id"] for task in filtered], ["late"])

        prepared = collect_data(
            Mock(), tasks, "Performance Analysis", "due-variance", time_status_data={},
            filter_summary="Status: COMPLETE",
        )
        result = analyze_clickup(prepared)
        variance = result["tasks"].set_index("Task ID")["Due Variance (days)"]
        self.assertEqual(variance["late"], 2.0)
        self.assertEqual(variance["early"], -1.0)
        self.assertIn("Completed late", set(result["due_variance_summary"]["Due Variance Category"]))
        self.assertIn("Completed early", set(result["due_variance_summary"]["Due Variance Category"]))
        self.assertFalse(result["assignee_summary"].empty)
        self.assertFalse(result["weekly_flow"].empty)
        workbook = load_workbook(io.BytesIO(analysis_excel(result)), data_only=True)
        self.assertIn("Due_Variance", workbook.sheetnames)
        self.assertIn("Assignee_Summary", workbook.sheetnames)

    def test_word_report_contains_clickup_due_variance_section(self):
        tasks = [{
            "id": "late", "name": "Deliver late", "status": {"status": "COMPLETE", "type": "done"},
            "assignees": [{"username": "Haidar"}], "date_created": "2026-09-01T09:00:00Z",
            "due_date": "2026-09-03T09:00:00Z", "date_closed": "2026-09-05T09:00:00Z",
        }]
        prepared = collect_data(Mock(), tasks, "Performance Analysis", "word-report", time_status_data={})
        report = create_clickup_word_report(analyze_clickup(prepared))
        document = Document(io.BytesIO(report))
        content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("ClickUp Task Performance Evaluation Report", content)
        self.assertIn("Due Variance and Timeliness", content)


if __name__ == "__main__":
    unittest.main()
