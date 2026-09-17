import copy
import io
import unittest
from datetime import datetime, timezone
from openpyxl import load_workbook

from test_automation import SETTINGS, FIELDS, example_issue, no_events
from collection_job import CollectionJob
from jira_export import issue_rows
from jira_gateway import CollectionError


class FakeStore:
    jobs = {}

    @classmethod
    def reset(cls):
        cls.jobs = {}

    def close(self):
        pass

    def begin(self, owner_key, job_id, fingerprint, space, query, definitions, cutoff):
        self.jobs.setdefault(job_id, {"job_id": job_id, "owner_key": owner_key,
            "fingerprint": fingerprint, "space": copy.deepcopy(space), "query": query,
            "definitions": copy.deepcopy(definitions), "cutoff": cutoff, "status": "running",
            "stage": "Ready", "current_issue": None, "last_successful_issue": None,
            "message": "Persistent checkpoint created.", "error_detail": None,
            "result_meta": None, "items": [], "total_count": 0, "completed_count": 0})
        return copy.deepcopy(self.jobs[job_id])

    def seed_items(self, owner_key, job_id, items):
        job = self.jobs[job_id]
        rows = []
        existing = {row["issue_key"]: row for row in job["items"]}
        for position, item in enumerate(items):
            row = existing.get(item["key"], {"position": position, "issue_key": item["key"],
                "seed_item": copy.deepcopy(item), "item_data": None, "history_data": None,
                "completed": False})
            row.update(position=position, seed_item=copy.deepcopy(item))
            rows.append(row)
        job.update(items=rows, total_count=len(rows))
        return {"total_count": len(rows)}

    def save_item(self, owner_key, job_id, issue_key, item, history):
        job = self.jobs[job_id]
        row = next(row for row in job["items"] if row["issue_key"] == issue_key)
        row.update(item_data=copy.deepcopy(item), history_data=copy.deepcopy(history), completed=True)
        job["completed_count"] = sum(row["completed"] for row in job["items"])
        return {"completed_count": job["completed_count"], "total_count": job["total_count"]}

    def update_job(self, owner_key, job_id, status, stage, current_issue, message,
                   error_detail=None, result_meta=None):
        self.jobs[job_id].update(status=status, stage=stage, current_issue=current_issue,
            message=message, error_detail=copy.deepcopy(error_detail))
        if result_meta is not None:
            self.jobs[job_id]["result_meta"] = copy.deepcopy(result_meta)
        return {"status": status}

    def load_job(self, owner_key, job_id):
        job = self.jobs.get(job_id)
        if not job:
            return None
        return copy.deepcopy({key: value for key, value in job.items() if key != "owner_key"})


class FakeJira:
    visited = []
    fail_once = False

    def __init__(self, settings):
        self.progress = None

    def close(self):
        pass

    def all_issues(self, query, progress=None):
        return [example_issue("TEST-" + str(n)) for n in range(1, 25)]

    def project_details(self, project_id):
        return copy.deepcopy(SPACE)

    def complete_issue(self, item):
        self.visited.append(item["key"])
        if item["key"] == "TEST-17" and type(self).fail_once:
            type(self).fail_once = False
            raise CollectionError("Simulated temporary interruption")
        history = no_events(item["key"])
        history["history_through"] = datetime.now(timezone.utc).isoformat()
        return copy.deepcopy(item), history

SPACE = {"id": "100", "key": "TEST", "name": "Test"}
OWNER_KEY = "f" * 64


class ResumeGateway(FakeJira):
    visited = []
    fail_once = True

    def all_issues(self, query, progress=None):
        return [example_issue("TEST-" + str(n)) for n in range(1, 25)]

    def project_details(self, project_id):
        return copy.deepcopy(SPACE)

    def complete_issue(self, item):
        self.visited.append(item["key"])
        if item["key"] == "TEST-17" and type(self).fail_once:
            type(self).fail_once = False
            raise CollectionError("Simulated temporary interruption")
        history = no_events(item["key"])
        history["history_through"] = datetime.now(timezone.utc).isoformat()
        return copy.deepcopy(item), history


def drive(job, limit=100):
    for _ in range(limit):
        if job.result is not None or job.error is not None:
            return
        job.step()
    raise AssertionError("Collection did not reach a terminal state")


class ResumeTests(unittest.TestCase):
    def setUp(self):
        FakeStore.reset()
        ResumeGateway.visited = []
        ResumeGateway.fail_once = True

    def test_twenty_four_tasks_resume_without_repeating_completed_tasks(self):
        store = FakeStore()
        job = CollectionJob(
            SETTINGS, SPACE, "project=TEST", "fp", FIELDS, ResumeGateway,
            store=store, owner_key=OWNER_KEY,
        )
        job.start()
        drive(job)
        self.assertFalse(job.running)
        self.assertEqual(job.snapshot()["completed"], 16)
        self.assertIsNone(job.result)
        self.assertIn("temporary interruption", job.error.lower())
        self.assertEqual(job.snapshot()["current_issue"], "TEST-17")
        cutoff = job.checkpoint["cutoff"]

        job.start()
        drive(job)
        self.assertIsNone(job.error)
        self.assertEqual(job.result.count, 24)
        self.assertEqual(job.result.cutoff, cutoff)
        self.assertEqual(ResumeGateway.visited.count("TEST-1"), 1)
        self.assertEqual(ResumeGateway.visited.count("TEST-17"), 2)

        book = load_workbook(io.BytesIO(job.result.xlsx))
        self.assertEqual(book["Jira_Data"].max_row, 25)
        self.assertEqual(book["History_Coverage"].max_row, 25)

    def test_persisted_checkpoint_rehydrates_after_process_loss(self):
        store = FakeStore()
        first = CollectionJob(
            SETTINGS, SPACE, "project=TEST", "fp", FIELDS, ResumeGateway,
            store=store, owner_key=OWNER_KEY,
        )
        first.start()
        first.step()
        self.assertEqual(first.snapshot()["completed"], 1)
        job_id = first.id

        # Simulate a process/browser loss before the retry boundary is reached.
        # The new process must not inherit transient in-memory retry/failure state.
        ResumeGateway.fail_once = False
        payload = store.load_job(OWNER_KEY, job_id)
        restored = CollectionJob.from_persisted(
            SETTINGS, payload, ResumeGateway, store=store, owner_key=OWNER_KEY,
        )
        self.assertEqual(restored.id, job_id)
        self.assertEqual(restored.snapshot()["completed"], 1)
        restored.start()
        drive(restored)
        self.assertEqual(ResumeGateway.visited.count("TEST-1"), 1)
        self.assertIsNotNone(restored.result)
        self.assertEqual(restored.result.count, 24)

    def test_persistence_acknowledges_item_before_ram_checkpoint(self):
        class RecordingStore(FakeStore):
            saved = []
            def save_item(self, owner_key, job_id, issue_key, item, history):
                type(self).saved.append(issue_key)
                return super().save_item(owner_key, job_id, issue_key, item, history)

        RecordingStore.saved = []
        store = RecordingStore()
        job = CollectionJob(
            SETTINGS, SPACE, "project=TEST", "fp", FIELDS, ResumeGateway,
            store=store, owner_key=OWNER_KEY,
        )
        job.start()
        job.step()
        self.assertEqual(RecordingStore.saved, ["TEST-1"])
        self.assertIn("TEST-1", job.checkpoint["completed"])
        persisted = store.load_job(OWNER_KEY, job.id)
        first_row = persisted["items"][0]
        self.assertTrue(first_row["completed"])
        self.assertEqual(first_row["issue_key"], "TEST-1")

    def test_fixed_columns_project_people_and_duplicate_custom_fields(self):
        item = example_issue()
        item["fields"]["project"].update(
            projectTypeKey="software",
            lead={"displayName": "Lead", "accountId": "a1"},
            description="Project description",
        )
        item["fields"]["reporter"] = {"displayName": "Reporter", "accountId": "a2"}
        item["fields"]["customfield_1"] = "first"
        item["fields"]["customfield_2"] = "second"
        fields = FIELDS + [
            {"id": "customfield_1", "name": "Vulnerability"},
            {"id": "customfield_2", "name": "Vulnerability"},
        ]
        headers, rows, _ = issue_rows([item], fields)
        self.assertEqual(headers[:5], ["Summary", "Issue key", "Issue id", "Issue Type", "Status"])
        self.assertEqual(len(headers), len(set(headers)))
        self.assertIn("Σ Time Spent", headers)
        self.assertIn("Watchers Id", headers)
        self.assertEqual(rows[0]["Project lead id"], "a1")
        self.assertEqual(rows[0]["Reporter Id"], "a2")
        self.assertEqual(rows[0]["Custom field (Vulnerability)"], "first")
        self.assertEqual(rows[0]["Custom field (Vulnerability) [customfield_2]"], "second")


if __name__ == "__main__":
    unittest.main()
