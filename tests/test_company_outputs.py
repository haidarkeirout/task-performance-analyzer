"""Focused contract tests for Company Performance executive outputs."""

from datetime import date, datetime, timezone
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from company_performance.dashboard import build_company_dashboard
from company_performance.kpis import generate_recommendations, identify_bottleneck_candidates
from company_performance.models import StatusTransition, TaskRecord
from company_performance.outputs import COMPANY_SHEET_NAMES, write_company_excel, write_company_raw_data, write_company_word_report
from company_performance.workflow import reconstruct_task


UTC = timezone.utc
START = date(2026, 9, 1)
END = date(2026, 9, 30)


def transition(before, after, timestamp):
    return StatusTransition(datetime.fromisoformat(timestamp).replace(tzinfo=UTC), before, after)


def sample_snapshot():
    record = TaskRecord(
        source_tool="Jira", task_id="CPA-1", task_name="Sample task", source_space="Engineering",
        unified_project="Platform", raw_status="In Review", initial_status="To Do", created_date=START,
        due_date=date(2026, 9, 10), history_complete=True,
        history_through=datetime(2026, 10, 1, tzinfo=UTC),
        collection_timestamp=datetime(2026, 10, 1, tzinfo=UTC),
        workflow_history=(transition("To Do", "In Progress", "2026-09-02T09:00:00"),
                          transition("In Progress", "In Review", "2026-09-04T09:00:00")),
    )
    return reconstruct_task(record, START, END)


class CompanyOutputTests(unittest.TestCase):
    def test_excel_has_exactly_four_approved_sheets(self):
        item = sample_snapshot()
        model = build_company_dashboard([item])
        with tempfile.TemporaryDirectory() as directory:
            path = write_company_excel(model, Path(directory) / "company.xlsx")
            workbook = load_workbook(path)
            self.assertEqual(tuple(workbook.sheetnames), COMPANY_SHEET_NAMES)
            self.assertEqual(workbook["Company_Executive_Dashboard"]["A1"].value, "Company Performance")
            self.assertNotIn("Raw Collected Data", workbook.sheetnames)

    def test_raw_data_is_a_separate_audit_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_company_raw_data([sample_snapshot()], Path(directory) / "raw.xlsx")
            self.assertEqual(load_workbook(path).sheetnames, ["Raw Collected Data"])

    def test_word_contains_all_approved_sections_and_na_handling(self):
        item = sample_snapshot()
        model = build_company_dashboard([item])
        with tempfile.TemporaryDirectory() as directory:
            path = write_company_word_report(
                model, Path(directory) / "company.docx"),
            document = Document(path[0])
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Company-wide Jira and ClickUp analysis", text)
            for title in (
                "Company Executive Summary", "Data Sources and Analysis Scope",
                "Company KPI Summary", "Department Performance Comparison",
                "Company-wide Bottlenecks", "Key Risks and Overdue Tasks",
                "Workflow Exceptions", "Top Departments Requiring Attention",
                "Executive Recommendations", "Data Quality and Coverage Limitations",
                "Metric Definitions",
            ):
                self.assertIn(title, text)
            self.assertIn("N/A", "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells))


if __name__ == "__main__":
    unittest.main()
