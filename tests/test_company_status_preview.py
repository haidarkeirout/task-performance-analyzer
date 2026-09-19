from datetime import date, datetime, timezone
from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
from docx import Document
from openpyxl import Workbook, load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from company_performance.adapters import adapt_clickup_collection
from company_performance.application import CompanyPreparedItem, build_company_analysis, build_company_preview, filter_company_preview
from company_performance.dashboard import DashboardFilters, build_company_dashboard
from company_performance.models import StatusTransition, TaskRecord, UnifiedStatus
from company_performance.normalization import deduplicate_tasks, normalize_status
from company_performance.workflow import reconstruct_task
from company_performance.outputs import write_company_excel, write_company_word_report
from company_performance.ui import _preview_table_rows

def jira_xlsx(space='Engineering', status='Done'):
    wb=Workbook();s=wb.active;s.title='Jira_Data';s.append(['Issue key','Issue id','Summary','Project key','Project name','Status','Priority','Assignee','Created','Due date','Custom field (Start date)']);s.append(['ENG-1','1','Ship feature','ENG',space,status,'High','Maya','2026-09-01T08:00:00Z','2026-09-10','2026-09-02']);b=BytesIO();wb.save(b);return b.getvalue()
class JiraPrepared:
    def __init__(self,space='Engineering'):
        self.xlsx=jira_xlsx(space);self.space_name=space;self.collected_at='2026-09-30T20:00:00Z';self.cutoff='2026-09-30'
        self.history_json=json.dumps({'ENG-1':{'history_complete':True,'history_through':'2026-09-30T20:00:00Z','initial_status':'To Do','status_events':[{'changed_at':'2026-09-02T08:00:00Z','from_status':'To Do','to_status':'In Progress'},{'changed_at':'2026-09-03T08:00:00Z','from_status':'In Progress','to_status':'Done'}]}}).encode()
class ClickPrepared:
    def __init__(self,space='Growth', task_id='cu-1', status=None, priority='2', name='Review campaign'):
        self.space_name=space;self.collected_at='2026-09-30T20:00:00Z';self.cutoff='2026-09-30'
        self.history_json=json.dumps({'tasks':[{'id':task_id,'name':name,'status':status or {'status':'Review'},'priority':{'priority':priority},'assignees':[{'username':'Zaher'}],'date_created':'1788256800000','due_date':'1790683200000','list':{'name':'Marketing'}}],'time_in_status':{}}).encode()

class CompanySourceSelectionTests(unittest.TestCase):
    def test_company_ui_uses_independent_jira_clickup_space_selectors_without_radio(self):
        text=(ROOT/'src/company_performance/ui.py').read_text(encoding='utf-8')
        self.assertIn('collect_company_spaces', (ROOT/'src/company_performance/collection.py').read_text(encoding='utf-8'))
        self.assertIn('Complete Task Preview', text)
        self.assertIn('"Company": row.company_name', text)
        self.assertIn('Analysis Period From', text)
        self.assertIn('Analysis Period To', text)
        self.assertIn('Run Company Analysis', text)

    def test_app_exposes_all_companies_and_passes_selected_scope_to_collection(self):
        text=(ROOT/'app.py').read_text(encoding='utf-8')
        self.assertIn('ALL_COMPANIES_ID', text)
        self.assertIn('All Companies', text)
        self.assertIn('company_selection=selected_company', text)
    def test_combined_preview_contains_selected_jira_and_clickup(self):
        rows=build_company_preview(jira_prepared=(JiraPrepared(),),clickup_prepared=(ClickPrepared(),))
        self.assertEqual({r.source_tool for r in rows},{'Jira','ClickUp'})
        click=next(r for r in rows if r.source_tool=='ClickUp')
        self.assertEqual(click.space,'Growth');self.assertEqual(click.original_status,'Review');self.assertEqual(click.final_status,'In Review')

    def test_all_companies_preview_preserves_company_label(self):
        prepared = CompanyPreparedItem(
            prepared=JiraPrepared(),
            source_tool='Jira',
            source_space='Engineering',
            project_name='Najm Al-Shamal [TEST]',
            company_id='najm-al-shamal-test',
            company_name='Najm Al-Shamal [TEST]',
            source_id='NAST',
            source_kind='jira_project',
        )
        rows = build_company_preview(jira_prepared=(prepared,))
        self.assertEqual(rows[0].company_name, 'Najm Al-Shamal [TEST]')
        self.assertEqual(_preview_table_rows(rows)[0]['Company'], 'Najm Al-Shamal [TEST]')

    def test_jira_preview_exposes_issue_type_and_epic_relationship(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Jira_Data'
        sheet.append([
            'Issue key', 'Issue id', 'Summary', 'Issue Type', 'Parent',
            'Project key', 'Project name', 'Status', 'Created',
        ])
        sheet.append([
            'ENG-EPIC', '10', 'Website rollout', 'Epic', None,
            'ENG', 'Engineering', 'Done', '2026-09-01T08:00:00Z',
        ])
        sheet.append([
            'ENG-1', '11', 'Ship landing page', 'Task',
            '{"key":"ENG-EPIC"}', 'ENG', 'Engineering', 'Done',
            '2026-09-01T08:00:00Z',
        ])
        stream = BytesIO()
        workbook.save(stream)

        prepared = JiraPrepared()
        prepared.xlsx = stream.getvalue()
        prepared.history_json = json.dumps({
            key: {
                'history_complete': True,
                'history_through': '2026-09-30T20:00:00Z',
                'initial_status': 'To Do',
                'status_events': [],
            }
            for key in ('ENG-EPIC', 'ENG-1')
        }).encode()
        rows = build_company_preview(jira_prepared=(prepared,))
        epic = next(row for row in rows if row.task_name == 'Website rollout')
        child = next(row for row in rows if row.task_name == 'Ship landing page')
        self.assertEqual(epic.issue_type, 'Epic')
        self.assertEqual(epic.parent_classification, 'Container Parent')
        self.assertEqual(child.issue_type, 'Task')
        self.assertEqual(child.parent_id, 'ENG-EPIC')
        self.assertEqual(child.epic_name, 'Website rollout')
        table = _preview_table_rows(rows)
        child_table = next(row for row in table if row['Task Name'] == 'Ship landing page')
        self.assertEqual(child_table['Epic / Workstream'], 'Website rollout')
        self.assertEqual(child_table['Parent / Epic'], 'ENG-EPIC')
        self.assertEqual(child_table['Issue Type'], 'Task')
    def test_preview_filtering_covers_requested_fields(self):
        rows=build_company_preview(jira_prepared=(JiraPrepared(),),clickup_prepared=(ClickPrepared(),ClickPrepared(space='Ops',task_id='cu-2',status={'status':'In Progress'},priority='1',name='Launch API')))
        self.assertEqual(len(filter_company_preview(rows,source_tools=['ClickUp'])),2)
        self.assertEqual(len(filter_company_preview(rows,spaces=['Ops'])),1)
        self.assertEqual(len(filter_company_preview(rows,task_name='launch')),1)
        self.assertEqual(len(filter_company_preview(rows,original_statuses=['Review'])),1)
        self.assertEqual(len(filter_company_preview(rows,final_statuses=['In Execution'])),1)
        self.assertEqual(len(filter_company_preview(rows,assignees=['Zaher'])),2)
        self.assertEqual(len(filter_company_preview(rows,priorities=['Critical'])),1)
        self.assertEqual(len(filter_company_preview(rows,due_date_start=date(2026,9,1),due_date_end=date(2026,9,30))),3)

class ClickUpStatusTests(unittest.TestCase):
    def test_clickup_status_name_extraction_prefers_readable_fields_over_id(self):
        for key in ('status','name','status_name','label'):
            with self.subTest(key=key):
                source={'id':'p_INTERNAL',key:'In Progress'}
                result=adapt_clickup_collection([{'id':'x','name':'X','status':source,'list':{'name':'Ops'}}],unified_project='P',source_space='S')
                rec=result.records[0];self.assertEqual(rec.raw_status,'In Progress');self.assertNotIn('ClickUp Status Unresolved',rec.data_quality_flags)
    def test_clickup_status_id_fallback_is_original_unknown_and_quality_flagged(self):
        result=adapt_clickup_collection([{'id':'x','name':'X','status':{'id':'p_INTERNAL'},'list':{'name':'Ops'}}],unified_project='P',source_space='S',collection_timestamp='2026-09-30T10:00:00Z')
        rec=deduplicate_tasks(result.records)[0]
        self.assertEqual(rec.raw_status,'p_INTERNAL');self.assertIs(normalize_status('ClickUp',rec.raw_status),UnifiedStatus.UNKNOWN);self.assertIn('ClickUp Status Unresolved',rec.data_quality_flags);self.assertIn('Unmapped Status',rec.data_quality_flags)
    def test_all_approved_clickup_mappings(self):
        expected={'To Do':UnifiedStatus.NOT_STARTED,'Planning':UnifiedStatus.NOT_STARTED,'In Progress':UnifiedStatus.IN_EXECUTION,'At Risk':UnifiedStatus.AT_RISK,'Review':UnifiedStatus.IN_REVIEW,'On Hold':UnifiedStatus.ON_HOLD,'Complete':UnifiedStatus.COMPLETED,'Cancelled':UnifiedStatus.CANCELLED}
        for raw,final in expected.items():self.assertIs(normalize_status('ClickUp',raw),final)

class CompanyClickUpIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.result=build_company_analysis(period_start=date(2026,9,1),period_end=date(2026,9,30),jira_prepared=(JiraPrepared(),),clickup_prepared=(ClickPrepared(status={'id':'internal','status':'Review'}),ClickPrepared(space='Ops',task_id='cu-2',status={'status':'In Progress'},name='Launch API')),unified_project='Company Delivery')
    def test_original_and_final_status_exposed_and_final_filter_works(self):
        click=[d for d in self.result.model.task_details if d.source_tool=='ClickUp'];self.assertEqual({d.original_status for d in click},{'Review','In Progress'});self.assertEqual({d.final_status for d in click},{'Unknown'})
        filtered=build_company_dashboard(self.result.snapshots,coverages=self.result.collection.coverages,filters=DashboardFilters(statuses=(UnifiedStatus.UNKNOWN,)))
        self.assertEqual(len(filtered.task_details),2);self.assertTrue(all(item.source_tool == 'ClickUp' for item in filtered.task_details))
    def test_clickup_is_in_filters_kpi_cards_and_charts(self):
        model=self.result.model;self.assertIn('ClickUp',model.filter_options.source_tools);self.assertEqual(model.kpis.total_tasks,3);self.assertEqual(next(c for c in model.cards if c.key=='current-wip').value,'0');delivery=next(c for c in model.executive_charts if c.key=='delivery-outcome');labels={p.label:p.value for p in delivery.points};self.assertEqual(labels['Completed'],1)
    def test_excel_task_details_has_both_status_columns(self):
        with tempfile.TemporaryDirectory() as td:
            path=write_company_excel(self.result.model,Path(td)/'out.xlsx');ws=load_workbook(path,data_only=True)['Task Details'];headers=[c.value for c in ws[1]];self.assertIn('Original Status',headers);self.assertIn('Final Status',headers);rows=list(ws.iter_rows(min_row=2,values_only=True));click=[r for r in rows if r[headers.index('Source Tool')]=='ClickUp'];self.assertTrue(click);self.assertIn('Review',{r[headers.index('Original Status')] for r in click});self.assertEqual({r[headers.index('Final Status')] for r in click},{'Unknown'})
    def test_word_contains_clickup_and_normalized_delivery_status(self):
        with tempfile.TemporaryDirectory() as td:
            path=write_company_word_report(self.result.model,Path(td)/'out.docx');doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells);self.assertIn('ClickUp',text);self.assertIn('In Review',text);self.assertIn('In Execution',text)

    def test_dashboard_excel_and_word_share_the_same_kpis(self):
        with tempfile.TemporaryDirectory() as td:
            excel_path = write_company_excel(self.result.model, Path(td) / 'out.xlsx')
            word_path = write_company_word_report(self.result.model, Path(td) / 'out.docx')
            expected = {
                'Total Tasks': str(self.result.model.kpis.total_tasks),
                'Completion Rate': f"{self.result.model.kpis.completion_rate:.1f}%",
                'Current WIP': str(self.result.model.kpis.current_wip),
            }

            summary = load_workbook(excel_path, data_only=True)['Company_Executive_Dashboard']
            header_row = next(summary.iter_rows(min_row=3, max_row=3, values_only=True))
            value_row = next(summary.iter_rows(min_row=4, max_row=4, values_only=True))
            excel_values = {
                header_row[index]: value_row[index]
                for index in range(len(header_row))
                if header_row[index] in expected
            }
            self.assertEqual(excel_values, expected)

            document = Document(word_path)
            word_values = {}
            for table in document.tables:
                if not table.rows or table.rows[0].cells[0].text.strip() != "KPI":
                    continue
                for row in table.rows[1:]:
                    cells = [cell.text.strip() for cell in row.cells]
                    if len(cells) >= 2 and cells[0] == "Total Tasks":
                        word_values["Total Tasks"] = cells[1]
                    if len(cells) >= 2 and cells[0] == "Completion Rate":
                        word_values["Completion Rate"] = f"{float(cells[1]):.1f}%"
                    if len(cells) >= 2 and cells[0] == "WIP":
                        word_values["Current WIP"] = cells[1]
            self.assertEqual(word_values, expected)
    def test_unresolved_id_stays_out_of_delivery_chart_and_is_in_data_quality(self):
        unresolved=build_company_analysis(period_start=date(2026,9,1),period_end=date(2026,9,30),clickup_prepared=(ClickPrepared(status={'id':'p_INTERNAL'}),),unified_project='P')
        self.assertEqual(unresolved.model.task_details[0].original_status,'p_INTERNAL');self.assertEqual(unresolved.model.task_details[0].final_status,'Unknown');delivery=next(c for c in unresolved.model.executive_charts if c.key=='delivery-outcome');self.assertNotIn('Unknown',[p.label for p in delivery.points]);flags={i.flag for i in unresolved.model.data_quality};self.assertIn('ClickUp Status Unresolved',flags);self.assertIn('Unmapped Status',flags)

    def test_unknown_reason_details_are_exported_to_excel_and_word(self):
        unresolved=build_company_analysis(period_start=date(2026,9,1),period_end=date(2026,9,30),clickup_prepared=(ClickPrepared(status={'id':'p_INTERNAL'}),),unified_project='P')
        with tempfile.TemporaryDirectory() as td:
            excel_path=write_company_excel(unresolved.model,Path(td)/'quality.xlsx')
            quality=load_workbook(excel_path,data_only=True)['Data Quality']
            values=[[cell.value for cell in row] for row in quality.iter_rows(values_only=False)]
            flattened='\n'.join(str(value) for row in values for value in row if value is not None)
            self.assertIn('Task ID', flattened)
            self.assertIn('Unmapped Status', flattened)
            word_path=write_company_word_report(unresolved.model,Path(td)/'quality.docx')
            document=Document(word_path)
            text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
            self.assertIn('Task-level Unknown Status Details', text)
            self.assertIn('Unmapped Status:', text)

    def test_collection_reconciliation_explains_out_of_period_task(self):
        outside = TaskRecord(
            source_tool='Jira', task_id='OLD-1', task_name='Older task', raw_status='Done',
            initial_status='To Do', created_date=date(2026, 8, 1), history_complete=True,
            workflow_history=(StatusTransition(
                changed_at=datetime(2026, 8, 15, tzinfo=timezone.utc),
                from_status='To Do', to_status='Done',
            ),),
        )
        current = TaskRecord(
            source_tool='Jira', task_id='NOW-1', task_name='Current task', raw_status='To Do',
            initial_status='To Do', created_date=date(2026, 9, 1), history_complete=True,
        )
        snapshots = (
            reconstruct_task(outside, date(2026, 9, 1), date(2026, 9, 15)),
            reconstruct_task(current, date(2026, 9, 1), date(2026, 9, 15)),
        )
        model = build_company_dashboard(snapshots)
        self.assertEqual(model.in_period_task_count, 1)
        self.assertEqual(model.kpis.total_tasks, 1)
        self.assertEqual([item.task_id for item in model.task_details], ['NOW-1'])

if __name__=='__main__':unittest.main()
