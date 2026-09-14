import io
import unittest
from unittest.mock import Mock

from docx import Document
from openpyxl import load_workbook

from clickup_analysis import analyze_clickup
from clickup_export import collect_data
from department_analysis import build_department_result, department_excel, department_word


class DepartmentAnalysisTests(unittest.TestCase):
    def setUp(self):
        tasks = [
            {"id": "done", "name": "Delivered", "status": {"status": "COMPLETE", "type": "done"},
             "assignees": [{"username": "Maya"}], "date_created": "2026-09-01T09:00:00Z",
             "due_date": "2026-09-05T09:00:00Z", "date_closed": "2026-09-04T09:00:00Z"},
            {"id": "late", "name": "Needs attention", "status": {"status": "IN PROGRESS"},
             "priority": {"priority": "high"}, "assignees": [], "date_created": "2026-09-01T09:00:00Z",
             "due_date": "2026-09-03T09:00:00Z"},
        ]
        prepared = collect_data(
            Mock(), tasks, "Wujha", "department-test", time_status_data={},
            analysis_mode="department", department_name="Marketing", department_id="list-1",
        )
        prepared.cutoff = "2026-09-10T09:00:00+00:00"
        self.result = build_department_result(analyze_clickup(prepared), prepared)

    def test_kpis_use_department_scope_and_cancelled_safe_denominator(self):
        kpis = self.result["kpis"].set_index("KPI")
        self.assertEqual(kpis.loc["Total Tasks", "Value"], 2)
        self.assertEqual(kpis.loc["Task Completion Rate (%)", "Value"], 50.0)
        self.assertEqual(kpis.loc["Open Overdue Tasks", "Value"], 1)
        self.assertEqual(self.result["department_name"], "Marketing")
        self.assertEqual(len(self.result["attention"]), 1)

    def test_excel_contains_exactly_four_approved_sheets(self):
        workbook = load_workbook(io.BytesIO(department_excel(self.result)), data_only=True)
        self.assertEqual(workbook.sheetnames, [
            "Department Summary", "Employee Breakdown", "Task Details", "Exceptions & Data Quality",
        ])

    def test_word_is_department_report(self):
        document = Document(io.BytesIO(department_word(self.result)))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Department Performance Evaluation Report", text)
        self.assertIn("Department: Marketing", text)
        self.assertIn("Bottleneck Candidates", text)


if __name__ == "__main__":
    unittest.main()
