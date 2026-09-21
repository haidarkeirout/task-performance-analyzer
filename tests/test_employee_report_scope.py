"""Regression tests for Employee-scoped combined Jira + ClickUp reporting."""

from datetime import date, datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch
import sys
import tempfile

from openpyxl import load_workbook
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import employee_ui
from company_performance.dashboard import build_company_dashboard
from company_performance.models import StatusTransition, TaskRecord
from company_performance.outputs import write_company_excel, write_company_word_report
from company_performance.ui import _cached_output_bytes, _filters
from company_performance.workflow import reconstruct_task


UTC = timezone.utc
START = date(2026, 9, 1)
END = date(2026, 9, 30)


def _employee_snapshot():
    task = TaskRecord(
        source_tool="Jira",
        task_id="EMP-J-1",
        task_name="Employee Jira task",
        source_space="Software",
        unified_project="Employee: Test Employee",
        raw_status="Done",
        initial_status="To Do",
        created_date=START,
        workflow_history=(
            StatusTransition(
                datetime(2026, 9, 2, 9, tzinfo=UTC),
                "To Do",
                "Done",
            ),
        ),
        history_complete=True,
        history_through=datetime(2026, 10, 1, tzinfo=UTC),
        collection_timestamp=datetime(2026, 10, 1, tzinfo=UTC),
    )
    return reconstruct_task(task, START, END)


def _clickup_snapshot():
    task = TaskRecord(
        source_tool="ClickUp",
        task_id="EMP-C-1",
        task_name="Employee ClickUp task",
        source_space="Operations",
        unified_project="Employee: Test Employee",
        raw_status="Complete",
        initial_status="To Do",
        created_date=START,
        workflow_history=(
            StatusTransition(
                datetime(2026, 9, 3, 9, tzinfo=UTC),
                "To Do",
                "Complete",
            ),
        ),
        history_complete=True,
        history_through=datetime(2026, 10, 1, tzinfo=UTC),
        collection_timestamp=datetime(2026, 10, 1, tzinfo=UTC),
    )
    return reconstruct_task(task, START, END)


class _FakeColumn:
    def __init__(self, keys):
        self.keys = keys

    def multiselect(self, _label, _options, *, key):
        self.keys.append(key)
        return []


class _FakeStreamlit:
    def __init__(self):
        self.keys = []

    def columns(self, count):
        return tuple(_FakeColumn(self.keys) for _ in range(count))


class EmployeeScopeTests(TestCase):
    def test_employee_filters_use_employee_namespace(self):
        st = _FakeStreamlit()
        options = SimpleNamespace(
            source_tools=("Jira",),
            unified_projects=("Employee: Test Employee",),
            statuses=(),
            assignee_groups=("Test Employee",),
            priorities=("High",),
        )
        result = SimpleNamespace(snapshots=(), collection=SimpleNamespace(coverages=()))

        with patch(
            "company_performance.ui.build_company_dashboard",
            return_value=SimpleNamespace(filter_options=options),
        ):
            _filters(st, result, key_prefix="employee")

        self.assertEqual(
            st.keys,
            [
                "employee_filter_sources",
                "employee_filter_projects",
                "employee_filter_statuses",
                "employee_filter_assignees",
                "employee_filter_priorities",
            ],
        )

    def test_employee_output_cache_is_separate_from_company_cache(self):
        st = SimpleNamespace(session_state={})
        result = SimpleNamespace(snapshots=())
        model = SimpleNamespace(task_details=())

        with patch(
            "company_performance.ui._output_bytes",
            side_effect=[
                (b"employee-excel", b"employee-raw", b"employee-word"),
                (b"company-excel", b"company-raw", b"company-word"),
            ],
        ) as output:
            employee_bytes = _cached_output_bytes(
                st,
                result,
                model,
                cache_prefix="employee",
                output_stem="employee_performance",
                scope_label="Employee",
            )
            company_bytes = _cached_output_bytes(
                st,
                result,
                model,
                cache_prefix="company",
                output_stem="company_performance",
                scope_label="Company",
            )

        self.assertEqual(employee_bytes[0], b"employee-excel")
        self.assertEqual(company_bytes[0], b"company-excel")
        self.assertEqual(st.session_state["employee_output_cache"][0], b"employee-excel")
        self.assertEqual(st.session_state["company_output_cache"][0], b"company-excel")
        self.assertEqual(output.call_count, 2)

    def test_employee_exports_use_employee_labels_without_changing_task_content(self):
        jira_snapshot = _employee_snapshot()
        clickup_snapshot = _clickup_snapshot()
        snapshots = (jira_snapshot, clickup_snapshot)
        model = build_company_dashboard(snapshots, scope_label="Employee")

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            excel_path = write_company_excel(
                model,
                root / "employee_performance_analysis.xlsx",
                snapshots=snapshots,
                scope_label="Employee",
            )
            word_path = write_company_word_report(
                model,
                root / "employee_performance_report.docx",
                snapshots=snapshots,
                scope_label="Employee",
            )

            workbook = load_workbook(excel_path, read_only=True, data_only=True)
            document = Document(word_path)

        self.assertEqual(workbook.sheetnames[0], "Employee_Executive_Dashboard")
        self.assertEqual(
            workbook["Employee_Executive_Dashboard"]["A1"].value,
            "Employee Performance",
        )
        self.assertEqual(workbook["Analysis Context"]["B2"].value, "Employee")
        self.assertEqual(workbook["Task Details"]["A2"].value, "Jira")
        self.assertEqual(workbook["Task Details"]["D2"].value, "EMP-J-1")
        self.assertEqual(workbook["Task Details"]["D3"].value, "EMP-C-1")

        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Employee Performance Report", document_text)
        self.assertIn("Employee Executive Summary", document_text)
        self.assertNotIn("Company Performance Report", document_text)
        self.assertNotIn("Company Executive Summary", document_text)

        self.assertIn("Employee scope", model.cards[1].supporting_text)
        self.assertEqual(model.task_details[0].task_id, "EMP-J-1")


class _EmployeeSelectionStreamlit:
    def __init__(self):
        self.session_state = {
            "employee_selected_name": "New Employee",
            "employee_snapshot": object(),
            "employee_prepared": object(),
            "employee_company_analysis": object(),
            "task_metrics": object(),
            "process_data": object(),
            "employee_collection_id": "old-collection",
            "employee_collection_in_progress": True,
        }
        self.rerun_calls = 0

    def rerun(self):
        self.rerun_calls += 1


class EmployeeSelectionResetTests(TestCase):
    def test_changing_employee_clears_stale_result_and_requests_full_rerun(self):
        streamlit = _EmployeeSelectionStreamlit()

        with patch.object(employee_ui, "st", streamlit):
            employee_ui._on_employee_changed()

        self.assertEqual(streamlit.rerun_calls, 1)
        self.assertEqual(streamlit.session_state["employee_selected_name"], "New Employee")
        for key in (
            "employee_snapshot",
            "employee_prepared",
            "employee_company_analysis",
            "task_metrics",
            "process_data",
            "employee_collection_id",
            "employee_collection_in_progress",
        ):
            self.assertNotIn(key, streamlit.session_state)
