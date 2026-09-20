import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from clickup_export import _sheet
from excel_safety import write_excel_cell


class ExcelSafetyTests(unittest.TestCase):
    def test_source_strings_are_plain_text_and_illegal_characters_are_escaped(self):
        workbook = Workbook()
        sheet = workbook.active
        write_excel_cell(sheet, 1, 1, '=HYPERLINK("https://example.test","open")')
        write_excel_cell(sheet, 2, 1, "unsafe\x00name")
        output = BytesIO()
        workbook.save(output)

        loaded = load_workbook(BytesIO(output.getvalue()), data_only=False)
        self.assertEqual(loaded.active["A1"].data_type, "s")
        self.assertTrue(loaded.active["A1"].value.startswith("="))
        self.assertEqual(loaded.active["A2"].value, "unsafe\\u0000name")

    def test_clickup_export_uses_the_shared_guard(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        _sheet(workbook, "Tasks", ["Name"], [["=1+1"]])
        output = BytesIO()
        workbook.save(output)

        loaded = load_workbook(BytesIO(output.getvalue()), data_only=False)
        self.assertEqual(loaded["Tasks"]["A2"].data_type, "s")
        self.assertEqual(loaded["Tasks"]["A2"].value, "=1+1")


if __name__ == "__main__":
    unittest.main()
