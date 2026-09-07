"""Run: python -m unittest discover -s tests -v
Set JIRA_TEST_WORKBOOK to the original six-sheet workbook for reconciliation.
Tests do not call Jira or require credentials.
"""
import ast
import copy
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from zoneinfo import ZoneInfo
import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from metrics_engine import WorkCalendar, analyze_tasks, calculate_task, aggregate, parse_timestamp, business_hours_between
from process_analysis import workbook_histories, workbook_context, process_tables, stage_summary, weekly_flow_summary, excel_bytes
from jira_processor import choose_sheet, normalize_tasks
from report_builder import build_report

CUTOFF = pd.Timestamp('2026-09-03T20:59:59Z')
CAL = WorkCalendar()

def task(key='T-1', status='Idea', due='2026-09-02'):
    return dict(issue_key=key, task_name='Test task', created_at='2026-09-01T06:00:00Z',
                due_date=due, planned_start_date='2026-09-01', current_status=status,
                assignee_name='Unassigned', issue_type='Task')

def history(events=(), initial='Idea', complete=True):
    return {'history_complete': complete, 'history_through': CUTOFF.isoformat(),
            'initial_status': initial, 'status_events': [
                dict(from_status=a,to_status=b,changed_at=c) for a,b,c in events]}

class EngineTests(unittest.TestCase):
    def test_confirmed_no_transitions_and_missing_history(self):
        known = calculate_task(task(),history(),CUTOFF,CAL)
        self.assertTrue(known['history_complete'])
        self.assertEqual(known['status_at_cutoff'],'Idea')
        self.assertTrue(known['is_open'])
        self.assertEqual(known['rework_count'],0)
        missing = calculate_task(task(),None,CUTOFF,CAL)
        self.assertFalse(missing['history_complete'])
        self.assertEqual(missing['status_at_cutoff'],'Unavailable')
        self.assertIsNone(missing['rework_count'])
        self.assertIsNone(missing['execution_elapsed_hours'])

    def test_future_events_and_tasks_do_not_leak(self):
        h=history([('Idea','In Progress','2026-09-02T06:00Z'),('In Progress','Done','2026-09-03T06:00Z')])
        before=pd.Timestamp('2026-09-01T10:00Z')
        r=calculate_task(task(status='Done'),h,before,CAL)
        self.assertEqual(r['status_at_cutoff'],'Idea')
        self.assertIsNone(r['completed_at'])
        future=task('T-2');future['created_at']='2026-09-04T06:00Z'
        frame=analyze_tasks(pd.DataFrame([task(),future]),{'T-1':h},before,CAL)
        self.assertEqual(len(frame),1)
        self.assertEqual(frame.attrs['excluded_tasks'][0]['issue_key'],'T-2')

    def test_partial_broken_and_uncovered_histories_disable_metrics(self):
        events=[('Idea','In Progress','2026-09-01T07:00Z'),('In Review','Done','2026-09-01T08:00Z')]
        for h in [history(events),history(events[:1],complete=False),history([('Idea','Done','invalid')])]:
            r=calculate_task(task(status='Done'),h,CUTOFF,CAL)
            self.assertFalse(r['history_complete'])
            self.assertIsNone(r['actual_start_at'])
            self.assertIsNone(r['on_time_completion'])
        h=history();h['history_through']='2026-09-02T00:00Z'
        self.assertFalse(calculate_task(task(),h,CUTOFF,CAL)['history_complete'])

    def test_open_overdue_denominator_excludes_late_completed(self):
        done=history([('Idea','In Progress','2026-09-01T07:00Z'),('In Progress','Done','2026-09-03T08:00Z')])
        tasks=pd.DataFrame([task('T-1','Done'),task('T-2'),task('T-3',due='2026-09-20')])
        r=analyze_tasks(tasks,{'T-1':done,'T-2':history(),'T-3':history()},CUTOFF,CAL)
        summary=aggregate(r,[]).iloc[0]
        self.assertEqual(summary['overdue_open_tasks'],1)
        self.assertEqual(summary['overdue_valid_tasks'],2)
        self.assertEqual(summary['open_overdue_rate'],50)

    def test_repeated_visits_and_union_and_zero_denominator(self):
        events=[('Idea','In Progress','2026-09-01T07:00Z'),('In Progress','In Review','2026-09-01T08:00Z'),
                ('In Review','In Progress','2026-09-01T09:00Z'),('In Progress','In Review','2026-09-01T10:00Z'),
                ('In Review','To Do','2026-09-01T11:00Z')]
        r=analyze_tasks(pd.DataFrame([task()]),{'T-1':history(events)},CUTOFF,CAL)
        summary=aggregate(r,[]).iloc[0]
        self.assertEqual(summary['review_exception_tasks'],1)
        self.assertEqual(summary['review_exception_rate'],100)
        stages=stage_summary(r).set_index('status')
        self.assertEqual(stages.loc['In Progress','elapsed_total_hours'],2)
        self.assertEqual(stages.loc['In Review','tasks_visited'],1)
        empty=analyze_tasks(pd.DataFrame([task()]),{'T-1':history()},CUTOFF,CAL)
        self.assertIsNone(aggregate(empty,[]).iloc[0]['rework_rate'])

    def test_reopen_cutoff_and_calendar(self):
        h=history([('Idea','In Progress','2026-09-01T07:00Z'),('In Progress','Done','2026-09-01T08:00Z'),
                   ('Done','In Progress','2026-09-02T07:00Z')])
        r=calculate_task(task(),h,CUTOFF,CAL)
        self.assertFalse(r['is_completed']);self.assertIsNone(r['completed_at']);self.assertEqual(r['reopen_count'],1)
        self.assertEqual(business_hours_between(pd.Timestamp('2026-09-03T13:00Z'),pd.Timestamp('2026-09-06T07:00Z'),CAL),2)

    def test_export_round_trip_and_word(self):
        h={'T-1':history()};r=analyze_tasks(pd.DataFrame([task()]),h,CUTOFF,CAL)
        tables=process_tables(r,h,{'Process Name':'Test'})
        data=excel_bytes(r,tables)
        sheets=pd.read_excel(io.BytesIO(data),sheet_name=None)
        self.assertTrue({'workflow_events','stage_summary','metric_definitions','process_context','data_quality'}.issubset(sheets))
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'report.docx'
            build_report(r,p,title='Test',evaluation_cutoff=CUTOFF.isoformat(),timezone_name='Asia/Damascus',work_days='Sunday-Thursday',work_window='09:00-17:00',process_data=tables)
            doc=Document(p);text='\n'.join(x.text for x in doc.paragraphs)
            self.assertIn('Rework: Unavailable (0/0 tasks)',text)
            self.assertIn('Stage Residence',text)
        self.assertNotIn('Individual Achievement Profile',text)

    def test_weekly_flow_counts_created_and_completed_tasks(self):
        flow_cutoff = "2026-09-10T20:59:59Z"
        frame = pd.DataFrame([
            {"issue_key": "T-1", "created_at": "2026-09-01T08:00:00Z", "completed_at": "2026-09-03T08:00:00Z",
             "evaluation_cutoff": flow_cutoff, "work_calendar_timezone": "Asia/Damascus"},
            {"issue_key": "T-2", "created_at": "2026-09-08T08:00:00Z", "completed_at": None,
             "evaluation_cutoff": flow_cutoff, "work_calendar_timezone": "Asia/Damascus"},
            {"issue_key": "T-3", "created_at": "2026-09-08T08:00:00Z", "completed_at": "2026-09-09T08:00:00Z",
             "evaluation_cutoff": flow_cutoff, "work_calendar_timezone": "Asia/Damascus"},
        ])
        weekly = weekly_flow_summary(frame)
        self.assertEqual(len(weekly), 2)
        self.assertEqual(weekly.iloc[0]["tasks_opened"], 1)
        self.assertEqual(weekly.iloc[0]["tasks_completed"], 1)
        self.assertEqual(weekly.iloc[1]["tasks_opened"], 2)
        self.assertEqual(weekly.iloc[1]["tasks_completed"], 1)
        self.assertEqual(weekly.iloc[1]["net_flow"], 1)

class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        path=os.environ.get('JIRA_TEST_WORKBOOK')
        if not path: raise unittest.SkipTest('Set JIRA_TEST_WORKBOOK for original workbook reconciliation.')
        cls.workbook=pd.read_excel(path,sheet_name=None,dtype=object)
        config=json.loads((ROOT/'configs/jira_input_config.json').read_text())
        _,source=choose_sheet(cls.workbook,config,None)
        cls.tasks,cls.log=normalize_tasks(source,input_config=config,source_timezone=ZoneInfo('Asia/Damascus'),labels_delimiter=None,explicit_date_format=None)
        cls.histories=workbook_histories(cls.workbook,cls.tasks,'Asia/Damascus',CUTOFF.isoformat(),True)
        cls.frame=analyze_tasks(cls.tasks,cls.histories,CUTOFF,CAL)

    def test_reconciles_original_workbook(self):
        s=aggregate(self.frame,[]).iloc[0]
        for k,v in {'total_tasks':20,'completed_tasks':10,'open_tasks':8,'rejected_tasks':2,'wip_tasks':4,
                    'history_complete_tasks':20,'reviewed_valid_tasks':13,'total_rework_count':2,
                    'total_replanning_count':2,'total_re_evaluation_count':1,'review_exception_tasks':3}.items():
            self.assertEqual(s[k],v,k)
        original=dict(self.workbook['KPIs_Summary'][['KPI Name','Value']].itertuples(index=False,name=None))
        for col,label in [('mean_execution_elapsed_hours','Average Actual Execution Duration'),('mean_lead_time_elapsed_hours','Average Lead Time'),('mean_time_to_start_elapsed_hours','Average Time to Start')]:
            self.assertAlmostEqual(s[col],original[label],places=7)
        for col,label in [('rework_rate','Rework Rate'),('replanning_rate','Replanning Rate'),('re_evaluation_rate','Re-evaluation Rate'),('review_exception_rate','Review Exception Rate')]:
            self.assertAlmostEqual(s[col],100*original[label],places=7)
        tables=process_tables(self.frame,self.histories,workbook_context(self.workbook),self.log)
        self.assertEqual(len(tables['workflow_events']),87)
        self.assertEqual(len(tables['open_tasks']),8)

    def test_duplicate_and_malformed_events_flag_history(self):
        w={**self.workbook}
        w['Workflow_Events']=pd.concat([w['Workflow_Events'],w['Workflow_Events'].iloc[:1]],ignore_index=True)
        h=workbook_histories(w,self.tasks,'Asia/Damascus',CUTOFF.isoformat(),True)
        self.assertFalse(h['SCRUM-14']['history_complete'])

    def test_app_pipeline_and_both_downloads(self):
        # Execute actual app functions without launching Streamlit or requiring its UI package.
        source=ast.parse((ROOT/'app.py').read_text())
        constants={'ROOT_DIR','SRC_DIR','CONFIG_DIR','INPUT_CONFIG_PATH','METRICS_CONFIG_PATH'}
        nodes=[]
        for n in source.body:
            if isinstance(n,(ast.Import,ast.ImportFrom)):
                if isinstance(n,ast.Import) and any(a.name=='streamlit' for a in n.names):continue
                nodes.append(n)
            elif isinstance(n,ast.FunctionDef):nodes.append(n)
            elif isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id in constants for t in n.targets):nodes.append(n)
        ns={'__file__':str(ROOT/'app.py')}
        exec(compile(ast.Module(body=nodes,type_ignores=[]),'app.py','exec'),ns)
        class Upload:
            def getvalue(self):return Path(os.environ['JIRA_TEST_WORKBOOK']).read_bytes()
        r,log,tables=ns['run_analysis'](Upload(),None,'','','',CUTOFF.isoformat(),'Asia/Damascus',history_source='Workbook transitions',coverage_through=CUTOFF.isoformat(),coverage_confirmed=True)
        self.assertEqual(ns['calculate_dashboard_values'](r)['open'],8)
        exported=ns['create_excel_download'](r,None,None,None,tables=tables)
        self.assertIn('stage_summary',pd.read_excel(io.BytesIO(exported),sheet_name=None))
        word=ns['create_word_report_download'](r,CUTOFF.isoformat(),tables=tables)
        self.assertGreater(len(Document(io.BytesIO(word)).tables),20)

if __name__=='__main__':unittest.main()
